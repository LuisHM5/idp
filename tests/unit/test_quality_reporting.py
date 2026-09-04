from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "scripts"))
spec = importlib.util.spec_from_file_location("reporte_quality", REPO / "scripts" / "reporte.py")
report = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(report)


def quality() -> dict:
    free = {"suite": "quality-free", "status": "PASSED", "summary": {"tests": 508, "failures": 0, "lineCoverage": 68.3}}
    real = {"suite": "real-e2e", "status": "PASSED", "summary": {"fieldAccuracy": 1.0, "e2ePassRate": 1.0}}
    return {"runs": 2, "latest": real, "latestFree": free, "latestReal": real,
            "trend": [{"runId": "x", "at": "2026-09-04", "suite": "real-e2e", "status": "PASSED", "tests": 4, "failures": 0, "fieldAccuracy": 1.0}]}


def test_quality_html_distinguishes_accuracy_from_approval() -> None:
    html = report.render_quality_html(quality())
    assert "Exactitud contra valores humanos" in html
    assert "100.0%" in html
    assert "expected" not in html.lower()


def test_join_tags_test_traffic_from_metadata() -> None:
    rows = report.join_rows([], [{
        "document_id": "d1", "document_type": "PEDIMENTO_ADUANAL", "document_type_name": "",
        "document_name": "x.pdf", "app_origin": "qa", "processed_at": "2026-09-04",
        "workflow_status": "complete", "validation_status": "APROBADO", "model": "m",
        "processing_seconds": 1, "avg_confidence": .9, "warnings": [], "field_stats": {},
        "sub_document_count": 1, "metadata": {"test-run-id": "r1"},
    }])
    assert rows[0]["trafficKind"] == "test"


def test_excel_contains_executive_and_quality_sheets(tmp_path: Path) -> None:
    target = tmp_path / "report.xlsx"
    report.write_excel(target, [], quality())
    from openpyxl import load_workbook
    wb = load_workbook(target, read_only=True)
    assert "Resumen Ejecutivo" in wb.sheetnames
    assert "CalidadPruebas" in wb.sheetnames
