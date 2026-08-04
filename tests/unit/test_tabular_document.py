"""Tests for the deterministic tabular extraction pipeline
(pipeline/tabular_document.py).

Unlike test_xml_document.py (which validates a real catalog file,
FACTURA_HONORARIOS_XML.json, against a synthetic fixture of a real business
document), there is no concrete XLSX/CSV catalog type yet — these tests build
a throwaway type definition inline via `_parse_entry` to exercise the
mechanism (cell/column extraction, sheet/CSV parity, row-range detection)
independently of any specific business schema.
"""

import sys
from datetime import date
from pathlib import Path

LAMBDA_DIR = (
    Path(__file__).resolve().parents[2]
    / "src/document-processing/resources/enrichment-lambda"
)
sys.path.insert(0, str(LAMBDA_DIR))

import document_catalog  # noqa: E402
from document_catalog import _parse_entry  # noqa: E402
from pipeline import tabular_document  # noqa: E402

# A small "gastos" report: a title block (fixed cells) above a table (header
# row + data rows matched by column name) — the two shapes a real XLSX report
# combines in practice.
_CATALOG_RAW = {
    "REPORTE_GASTOS_XLSX": {
        "schema_version": "1.0.0",
        "display_name": "Reporte de Gastos (XLSX)",
        "fields": {
            "FOLIO": {"type": "string", "required": True, "cell": "B1", "validation": []},
            "FECHA": {"type": "string", "required": True, "cell": "B2", "validation": []},
            "TOTAL": {
                "type": "number",
                "required": True,
                "cell": "B3",
                "validation": [
                    {
                        "rule_type": "sum_array",
                        "source": "GASTOS.IMPORTE",
                        "tolerance": 0.02,
                        "message": "El total no coincide con la suma de gastos",
                    }
                ],
            },
            "GASTOS": {
                "type": "array",
                "required": True,
                "header_row": 5,  # title block occupies rows 1-4; table header starts at row 5
                "validation": [],
                "item_fields": {
                    "CONCEPTO": {"type": "string", "required": True, "column": "Concepto", "validation": []},
                    "IMPORTE": {"type": "number", "required": True, "column": "Importe", "validation": []},
                    "FECHA_GASTO": {"type": "date", "required": False, "column": "Fecha", "validation": []},
                },
            },
        },
        "cross_field_validation": [],
    }
}


def _load_entry():
    return _parse_entry("REPORTE_GASTOS_XLSX", _CATALOG_RAW["REPORTE_GASTOS_XLSX"])


def _patch_catalog(monkeypatch, type_id="REPORTE_GASTOS_XLSX"):
    entry = _load_entry()
    monkeypatch.setattr(document_catalog, "_CATALOG_CACHE", {type_id: entry})
    return entry


# GASTOS.IMPORTE sums to 1500.50 + 320.75 + 1200.00 = 3021.25 = TOTAL (B3),
# so the sum_array check passes end to end for both fixtures below.
def _build_xlsx_bytes() -> bytes:
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Folio:", "G-2026-001"
    ws["A2"], ws["B2"] = "Fecha:", "2026-08-04"
    ws["A3"], ws["B3"] = "Total:", 3021.25
    ws["A5"], ws["B5"], ws["C5"] = "Concepto", "Importe", "Fecha"
    ws["A6"], ws["B6"], ws["C6"] = "Combustible", 1500.50, date(2026, 8, 1)
    ws["A7"], ws["B7"] = "Casetas", 320.75
    ws["A8"], ws["B8"] = "Hospedaje", 1200.00

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_CSV_TEXT = (
    "Folio:,G-2026-001,\n"
    "Fecha:,2026-08-04,\n"
    "Total:,3021.25,\n"
    "\n"
    "Concepto,Importe,Fecha\n"
    "Combustible,1500.50,2026-08-01\n"
    "Casetas,320.75,\n"
    "Hospedaje,1200.00,\n"
)


def test_extracts_fixed_cells_and_table_rows_from_xlsx(monkeypatch):
    _patch_catalog(monkeypatch)
    result = tabular_document.process_tabular(
        _build_xlsx_bytes(), source_format="xlsx", declared_type="REPORTE_GASTOS_XLSX", source_name="reporte.xlsx",
    )
    doc = result["documents"][0]

    assert result["sourceFormat"] == "xlsx"
    assert doc["status"] == "APPROVED"
    hard_failures = [c for c in doc["checks"] if not c["valid"] and not c["skipped"]]
    assert hard_failures == []

    entities = doc["extractedEntities"]
    assert entities["FOLIO"]["value"] == "G-2026-001"
    assert entities["TOTAL"]["value"] == 3021.25
    rows = entities["GASTOS"]["items"]
    assert len(rows) == 3  # header_row=5 correctly skips the title block above it (rows 1-4)
    assert [r["fields"]["CONCEPTO"]["value"] for r in rows] == ["Combustible", "Casetas", "Hospedaje"]
    assert rows[0]["fields"]["IMPORTE"]["value"] == 1500.50
    assert rows[0]["fields"]["FECHA_GASTO"]["value"] == "2026-08-01"  # openpyxl date cell -> ISO string


def test_tags_values_as_tabular_extraction_with_full_confidence(monkeypatch):
    _patch_catalog(monkeypatch)
    result = tabular_document.process_tabular(
        _build_xlsx_bytes(), source_format="xlsx", declared_type="REPORTE_GASTOS_XLSX", source_name="reporte.xlsx",
    )
    entities = result["documents"][0]["extractedEntities"]
    assert entities["FOLIO"]["source"] == "TABULAR_EXTRACTION"
    assert entities["FOLIO"]["confidence"] == 1.0


def test_csv_is_a_workbook_with_one_implicit_sheet(monkeypatch):
    """Same catalog, same cell/column mapping, a CSV export instead of XLSX —
    the whole point of the shared grid abstraction is that this needs zero
    format-specific catalog changes."""
    _patch_catalog(monkeypatch)
    result = tabular_document.process_tabular(
        _CSV_TEXT.encode("utf-8"), source_format="csv", declared_type="REPORTE_GASTOS_XLSX", source_name="reporte.csv",
    )
    doc = result["documents"][0]
    assert result["sourceFormat"] == "csv"
    assert doc["status"] == "APPROVED"
    entities = doc["extractedEntities"]
    assert entities["FOLIO"]["value"] == "G-2026-001"
    assert len(entities["GASTOS"]["items"]) == 3


def test_requires_a_declared_type(monkeypatch):
    _patch_catalog(monkeypatch)
    result = tabular_document.process_tabular(
        _build_xlsx_bytes(), source_format="xlsx", declared_type=None,
    )
    assert result["documents"] == []
    assert result["warnings"]
    assert "declared type" in result["warnings"][0]


def test_invalid_workbook_fails_gracefully(monkeypatch):
    _patch_catalog(monkeypatch)
    result = tabular_document.process_tabular(
        b"not a real xlsx file", source_format="xlsx", declared_type="REPORTE_GASTOS_XLSX",
    )
    assert result["documents"] == []
    assert result["warnings"]


def test_cell_reference_parsing_handles_multi_letter_columns():
    assert tabular_document._cell_to_rowcol("A1") == (0, 0)
    assert tabular_document._cell_to_rowcol("B2") == (1, 1)
    assert tabular_document._cell_to_rowcol("Z1") == (0, 25)
    assert tabular_document._cell_to_rowcol("AA1") == (0, 26)


def test_data_rows_stops_at_first_blank_row_without_explicit_end_row():
    rows = [
        ["Concepto", "Importe"],   # header, row 1
        ["Casetas", 100],          # row 2
        ["Hospedaje", 200],        # row 3
        [None, None],              # row 4 — blank, should stop reading here
        ["Ignored", 999],          # row 5 — must NOT be picked up
    ]
    header, data = tabular_document._data_rows(rows, {"header_row": 1})
    assert header == ["Concepto", "Importe"]
    assert data == [["Casetas", 100], ["Hospedaje", 200]]


def test_data_rows_honors_explicit_end_row_past_a_blank_row():
    rows = [
        ["Concepto", "Importe"],
        ["Casetas", 100],
        [None, None],       # a legitimately blank row inside the range
        ["Hospedaje", 200],
    ]
    header, data = tabular_document._data_rows(rows, {"header_row": 1, "end_row": 4})
    assert data == [["Casetas", 100], ["Hospedaje", 200]]


# ------------------------------------------------------------- multi-sheet

# Three sheets, same table layout starting at row 3 (rows 1-2 hold a small
# metadata block instead — mirrors a real "one tab per month" workbook where
# each tab also carries its own header cells). FOLIO is deliberately blank on
# "Enero" and only set on "Febrero", to exercise the scalar fallback-across-
# sheets behavior without colliding with the table's own header row.
_MULTI_SHEET_CATALOG_RAW = {
    "REPORTE_MULTIHOJA_XLSX": {
        "schema_version": "1.0.0",
        "display_name": "Reporte multi-hoja (XLSX)",
        "fields": {
            "FOLIO": {
                "type": "string", "required": False, "sheet": ["Enero", "Febrero"], "cell": "B1", "validation": [],
            },
            "GASTOS_DOS_HOJAS": {
                "type": "array", "required": False, "item_sheet": ["Enero", "Febrero"], "header_row": 3,
                "validation": [],
                "item_fields": {
                    "CONCEPTO": {"type": "string", "required": True, "column": "Concepto", "validation": []},
                    "IMPORTE": {"type": "number", "required": True, "column": "Importe", "validation": []},
                },
            },
            "GASTOS_TODAS_HOJAS": {
                "type": "array", "required": False, "item_sheet": "*", "header_row": 3, "validation": [],
                "item_fields": {
                    "CONCEPTO": {"type": "string", "required": True, "column": "Concepto", "validation": []},
                    "IMPORTE": {"type": "number", "required": True, "column": "Importe", "validation": []},
                },
            },
            "GASTOS_HOJA_FALTANTE": {
                "type": "array", "required": False, "item_sheet": ["Enero", "NO_EXISTE"], "header_row": 3,
                "validation": [],
                "item_fields": {
                    "CONCEPTO": {"type": "string", "required": True, "column": "Concepto", "validation": []},
                },
            },
        },
        "cross_field_validation": [],
    }
}


def _build_multi_sheet_xlsx_bytes() -> bytes:
    import io

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Enero"
    ws["A1"], ws["B1"] = "Folio:", None  # deliberately blank on this sheet
    ws["A3"], ws["B3"] = "Concepto", "Importe"
    ws["A4"], ws["B4"] = "Renta", 1000

    febrero = wb.create_sheet("Febrero")
    febrero["A1"], febrero["B1"] = "Folio:", "F-002"
    febrero["A3"], febrero["B3"] = "Concepto", "Importe"
    febrero["A4"], febrero["B4"] = "Renta", 1100
    febrero["A5"], febrero["B5"] = "Luz", 200

    marzo = wb.create_sheet("Marzo")
    marzo["A1"], marzo["B1"] = "Folio:", "F-003"
    marzo["A3"], marzo["B3"] = "Concepto", "Importe"
    marzo["A4"], marzo["B4"] = "Renta", 1200

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _process_multi_sheet(monkeypatch):
    entry = _parse_entry("REPORTE_MULTIHOJA_XLSX", _MULTI_SHEET_CATALOG_RAW["REPORTE_MULTIHOJA_XLSX"])
    monkeypatch.setattr(document_catalog, "_CATALOG_CACHE", {"REPORTE_MULTIHOJA_XLSX": entry})
    return tabular_document.process_tabular(
        _build_multi_sheet_xlsx_bytes(), source_format="xlsx",
        declared_type="REPORTE_MULTIHOJA_XLSX", source_name="reporte-multihoja.xlsx",
    )


def test_item_sheet_as_list_concatenates_rows_in_order(monkeypatch):
    result = _process_multi_sheet(monkeypatch)
    entities = result["documents"][0]["extractedEntities"]
    rows = entities["GASTOS_DOS_HOJAS"]["items"]
    assert [(r["fields"]["CONCEPTO"]["value"], r["fields"]["IMPORTE"]["value"]) for r in rows] == [
        ("Renta", 1000.0),  # Enero
        ("Renta", 1100.0),  # Febrero
        ("Luz", 200.0),     # Febrero
    ]


def test_item_sheet_wildcard_reads_every_sheet(monkeypatch):
    result = _process_multi_sheet(monkeypatch)
    entities = result["documents"][0]["extractedEntities"]
    rows = entities["GASTOS_TODAS_HOJAS"]["items"]
    # Enero (1 row) + Febrero (2 rows) + Marzo (1 row), not just the sheets
    # explicitly named elsewhere in the catalog.
    assert len(rows) == 4
    assert sum(r["fields"]["IMPORTE"]["value"] for r in rows) == 1000.0 + 1100.0 + 200.0 + 1200.0


def test_item_sheet_skips_missing_names_without_failing(monkeypatch):
    result = _process_multi_sheet(monkeypatch)
    doc = result["documents"][0]
    entities = doc["extractedEntities"]
    rows = entities["GASTOS_HOJA_FALTANTE"]["items"]
    assert [r["fields"]["CONCEPTO"]["value"] for r in rows] == ["Renta"]  # only Enero's row
    hard_failures = [c for c in doc["checks"] if not c["valid"] and not c["skipped"]]
    assert hard_failures == []  # a missing sheet name is a warning-level log, not a validation failure


def test_scalar_sheet_list_tries_each_sheet_until_a_value_is_found(monkeypatch):
    result = _process_multi_sheet(monkeypatch)
    entities = result["documents"][0]["extractedEntities"]
    # B1 is blank on "Enero" (first in the list) — must fall through to
    # "Febrero" rather than reporting FOLIO as empty.
    assert entities["FOLIO"]["value"] == "F-002"


def test_resolve_sheet_names_covers_none_star_and_missing():
    grid = {"Enero": [], "Febrero": [], "Marzo": []}
    assert tabular_document._resolve_sheet_names(grid, None) == ["Enero"]
    assert tabular_document._resolve_sheet_names(grid, "*") == ["Enero", "Febrero", "Marzo"]
    assert tabular_document._resolve_sheet_names(grid, "Febrero") == ["Febrero"]
    assert tabular_document._resolve_sheet_names(grid, ["Marzo", "NO_EXISTE", "Enero"]) == ["Marzo", "Enero"]
    assert tabular_document._resolve_sheet_names({}, None) == []
