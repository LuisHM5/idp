#!/usr/bin/env python3
"""Reporte ÚNICO e interactivo: une la tabla de costos (`cost_audit.py`) y la
tabla de tracking de documentos (`indexing.py`) por `DocumentId` en un solo
dashboard — costo, tokens, confianza, estado de negocio y salud operativa,
todo en un archivo. `pnpm reporte:costos` y `pnpm reporte:documentos` se
quedan como reportes independientes (mismo dato, cada uno enfocado en su
tema); este es el que junta todo para no tener que abrir dos.

También separa ejecuciones con y sin IA usando la forma del catálogo y el uso
real de Bedrock, ordena los tipos por ruta de procesamiento y agrega el costo
de plataforma que no es atribuible a un documento: piso mensual WAF/KMS desde
el inventario CloudFormation y gasto MTD por CostStack o, si la etiqueta no
responde, por un cruce manual entre servicios facturados y recursos del stack.

El HTML es una página autocontenida (sin librerías externas, sin servidor):
los datos van embebidos como JSON y todo el filtrado/orden/detalle corre en
el navegador con JavaScript plano. Los filtros (fecha, tipo de documento,
sistema/`AppOrigin`, modelo, estado de ejecución y estado de negocio) viven
en un panel colapsado que se despliega con el botón "Filtros" — la barra
pegajosa se queda como un renglón delgado con el resumen del corte, y al
cambiar cualquier filtro se recalculan KPIs y tablas al vuelo. Click en un
documento abre el detalle completo (costo por servicio + calidad/campos).
Las secciones con mucho contenido van en `<details>` colapsables para no
forzar scroll infinito.

Además del costo y la calidad, analiza el **tiempo de proceso**
(`Tracking.ProcessingSeconds`: promedio, p50, p95, por página, distribución,
por tipo/sistema/servicio y los 10 más lentos) y contrasta el costo del corte
contra las mediciones reales de `documentacion/COSTOS.md` (ver
`COST_REFERENCE` abajo) — $/documento y $/página contra lo medido y contra
el presupuesto acordado.

Por default lee **dev**. Usa `--env prod` para el ambiente productivo.

Salida en `reportes/<env>/<timestamp>/`:
    reporte.md          resumen estático en texto/tablas
    reporte.html          dashboard interactivo (ábrelo en el navegador)
    datos_raw.xlsx         una fila por documento (costo + calidad) + hojas agregadas
    reporte_ejecutivo.xlsx copia compatible, formateada, con KPIs e histórico de pruebas
    snapshot.json          corte reproducible de datos normalizados

Uso:
    pnpm reporte                              # dev, año calendario actual completo
    pnpm reporte -- --env prod
    pnpm reporte -- --year 2025
    pnpm reporte -- --desde 2026-08-01 --hasta 2026-08-17
    pnpm reporte -- --sistema CEDIS
    pnpm reporte -- --open
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
import webbrowser
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional
from quality_history import DEFAULT_HISTORY, load_runs, quality_snapshot

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _report_lib import (  # noqa: E402
    CYAN,
    DIM,
    RED,
    RESET,
    YELLOW,
    cost_table_name,
    esc,
    normalize_cost_item,
    normalize_doc_item,
    resolve_processing_table,
    scan_all,
    stack_id,
)
from infrastructure_costs import (  # noqa: E402
    collect_infrastructure_costs,
    empty_infrastructure_snapshot,
)

REPO = Path(__file__).resolve().parent.parent
DEFAULT_OUT_ROOT = REPO / "reportes"


def resolve_report_period(year: int, desde: Optional[str], hasta: Optional[str]) -> tuple[date, date]:
    """Resolve the loaded document window; default is one full calendar year."""
    explicit_start = date.fromisoformat(desde) if desde else None
    explicit_end = date.fromisoformat(hasta) if hasta else None
    range_year = explicit_start.year if explicit_start else explicit_end.year if explicit_end else year
    report_start = explicit_start or date(range_year, 1, 1)
    report_end = explicit_end or date(range_year, 12, 31)
    if report_end < report_start:
        raise ValueError("--hasta no puede ser anterior a --desde")
    return report_start, report_end


def _has_structured_selector(value: Any) -> bool:
    if isinstance(value, dict):
        if any(key in value for key in ("xpath", "item_xpath", "cell", "column", "column_letter")):
            return True
        return any(_has_structured_selector(child) for child in value.values())
    if isinstance(value, list):
        return any(_has_structured_selector(child) for child in value)
    return False


def active_catalogs() -> Dict[str, Dict[str, Any]]:
    """Load active types for dynamic filters and processing classification.

    Mirrors the DSL's shape-based routing rule and recursively ignores every
    ``save/`` subtree plus ``_*.json`` templates, just like the catalog build.
    """
    types: Dict[str, Dict[str, Any]] = {}
    for path in (REPO / "catalogs").rglob("*.json"):
        if "save" in path.relative_to(REPO / "catalogs").parts or path.name.startswith("_"):
            continue
        try:
            raw = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        for type_id, config in (raw.items() if isinstance(raw, dict) else []):
            if isinstance(config, dict):
                types[type_id] = config
    return types


ACTIVE_CATALOGS = active_catalogs()
ACTIVE_CATALOG_TYPES = set(ACTIVE_CATALOGS)
DETERMINISTIC_TYPES = {
    type_id for type_id, config in ACTIVE_CATALOGS.items()
    if _has_structured_selector(config.get("fields") or {})
}


# ---------------------------------------------------------------------------
# Referencias de costo — NO son estimaciones: son las mediciones reales de
# `documentacion/COSTOS.md` (S3 `results/*.json` + `usage` real de Bedrock,
# cruzados contra Cost Explorer y la Pricing API). Sirven para contrastar lo
# que este reporte está midiendo hoy contra lo que se midió/acordó, en vez de
# tener que abrir el .md aparte. Si se re-mide el costo, se actualiza aquí
# **y** en COSTOS.md — el `source` de cada entrada dice de qué sección salió.
# ---------------------------------------------------------------------------

COST_REFERENCE: Dict[str, Any] = {
    "doc": "documentacion/COSTOS.md",
    "measuredOn": "2026-08-04",
    "volume": {"docsPerMonth": 2430, "pagesPerMonth": 7290, "budgetPerMonth": 60.52},
    # Líneas horizontales contra las que se compara el $/página del corte.
    "perPage": [
        {"label": "Colchón presupuestario ampliado", "value": 0.0131, "source": "§6", "role": "cushion"},
        {"label": "Presupuesto recomendado (IVA + 15% contingencia)", "value": 0.0083, "source": "§6", "role": "budget"},
        {"label": "Esperado con IVA (todo incluido, prod)", "value": 0.00743, "source": "§0.1", "role": "expected"},
        {"label": "Subtotal técnico antes de IVA", "value": 0.00640, "source": "§0.1", "role": "expected"},
        {"label": "Facturado real blended (Cost Explorer)", "value": 0.0049, "source": "§0", "role": "measured"},
        {"label": "Pedimento nativo 3 pág (caso puro)", "value": 0.0041, "source": "§0", "role": "measured"},
        {"label": "Pedimento escaneado 27 pág, 100 % OCR", "value": 0.0037, "source": "§0", "role": "measured"},
    ],
    # Costo real por documento medido por tipo (Bedrock + Textract atribuidos
    # documento por documento, ventana 29-jul → 04-ago 2026).
    "perDocByType": {
        "PEDIMENTO_ADUANAL": {"costPerDoc": 0.0421, "costPerPage": 0.0038, "n": 46,
                              "note": "mezcla 3 pág nativo + 27/50 pág escaneado"},
        "PEDIMENTO_ADUANAL_V1": {"costPerDoc": 0.0714, "costPerPage": 0.0357, "n": 1},
        "CONSTANCIA_SITUACION_FISCAL": {"costPerDoc": 0.0175, "costPerPage": 0.0087, "n": 4},
        "FACTURA_HONORARIOS": {"costPerDoc": 0.0138, "costPerPage": 0.0138, "n": 1},
        "INE": {"costPerDoc": 0.0080, "costPerPage": 0.0080, "n": 3, "note": "imagen, 100 % OCR"},
        "OTRO": {"costPerDoc": 0.0, "costPerPage": 0.0, "n": 28,
                 "note": "sin catálogo activo durante la migración del DSL — $0 temporal"},
        "CFDI_GASTO_TERCERO": {"costPerDoc": 0.0, "costPerPage": 0.0, "n": 7, "note": "catálogo archivado — $0 temporal"},
        "FACTURA_LOGISTICA": {"costPerDoc": 0.0, "costPerPage": 0.0, "n": 5, "note": "catálogo archivado — $0 temporal"},
        "PERMISO_CERTIFICADO_SANITARIO": {"costPerDoc": 0.0, "costPerPage": 0.0, "n": 4, "note": "catálogo archivado — $0 temporal"},
        "FACTURA_AGENCIA_ADUANAL": {"costPerDoc": 0.0, "costPerPage": 0.0, "n": 3, "note": "catálogo archivado — $0 temporal"},
        "FACTURA_HONORARIOS_XML": {"costPerDoc": 0.0, "costPerPage": 0.0, "n": 1, "note": "pipeline determinístico, nunca llama a Bedrock"},
    },
    # El caso de referencia vigente para presupuestar por documento (§6).
    "perDoc": {"label": "Pedimento de 3 páginas (Bedrock + Textract)", "value": 0.01365, "source": "§6"},
}


# ---------------------------------------------------------------------------
# Join
# ---------------------------------------------------------------------------

def _model_from_services(services: Dict[str, Any]) -> Optional[str]:
    """Last-resort model attribution when the tracking table has none —
    walks the cost breakdown's per-service `modelId` in the same phase
    priority the tracking side uses. Covers documents processed before
    `Tracking.ModelIds` existed at all (an older bug, since fixed, meant the
    key it did have often didn't match anyway — see `normalize_doc_item`)."""
    for name in ("BedrockExtraction", "BedrockSegmentation", "BedrockProcessingUpstream", "BedrockClassificationUpstream"):
        svc = services.get(name)
        if svc and svc.get("model_id"):
            return svc["model_id"]
    return None


def _ai_cost(services: Dict[str, Any]) -> float:
    """Costo de las fases Bedrock que realmente corrieron.

    La ruta se decide por evidencia de ejecución, no por TYPE_ID: cualquier
    tipo futuro con xpath/cell/column queda como ``Sin IA`` automáticamente y
    un mismo tipo puede aparecer como mixto si cambió de ruta en el periodo.
    """
    return sum(
        float(service.get("cost") or 0)
        for name, service in (services or {}).items()
        if name.startswith("Bedrock")
    )


def _uses_ai(has_cost: bool, services: Dict[str, Any]) -> Optional[bool]:
    if not has_cost:
        return None
    for name, service in (services or {}).items():
        if not name.startswith("Bedrock"):
            continue
        if (service.get("cost") or service.get("input_tokens") or service.get("output_tokens")):
            return True
    return False


def processing_mode_label(uses_ai: Optional[bool], *, expected_no_ai: bool = False) -> str:
    if uses_ai is True:
        return "Con IA"
    if uses_ai is False or expected_no_ai:
        return "Sin IA"
    return "Sin datos"


def join_rows(cost_rows: List[Dict[str, Any]], doc_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Full outer join por DocumentId — un documento puede tener costo sin
    tracking indexado (raro) o tracking sin costo (tabla de costos vacía o
    escritura best-effort que falló); ambos casos se marcan explícitos
    (`hasCost`/`hasQuality`) en vez de fingir un cero."""
    cost_by_id = {r["document_id"]: r for r in cost_rows if r["document_id"]}
    doc_by_id = {r["document_id"]: r for r in doc_rows if r["document_id"]}
    ids = set(cost_by_id) | set(doc_by_id)

    unified: List[Dict[str, Any]] = []
    for doc_id in ids:
        c = cost_by_id.get(doc_id)
        d = doc_by_id.get(doc_id)
        base = d or c  # identity fields — prefer tracking table, fall back to cost table
        doc_model = (d or {}).get("model")
        if not doc_model or doc_model == "(desconocido)":
            doc_model = _model_from_services((c or {}).get("services") or {}) or doc_model or "(desconocido)"
        services = (c or {}).get("services") or {}
        uses_ai = _uses_ai(c is not None, services)
        expected_no_ai = base["document_type"] in DETERMINISTIC_TYPES
        ai_cost = _ai_cost(services)
        total_cost = (c or {}).get("total_cost")
        unified.append({
            "id": doc_id,
            "type": base["document_type"],
            "typeName": base["document_type_name"],
            "name": (d or {}).get("document_name") or (c or {}).get("source_file") or "",
            "appOrigin": (d or {}).get("app_origin") or (c or {}).get("app_origin") or "",
            "metadata": (d or {}).get("metadata") or {},
            "trafficKind": "test" if ((d or {}).get("metadata") or {}).get("test-run-id") else "operational",
            "processedAt": (d or {}).get("processed_at") or (c or {}).get("processed_at") or "",
            "hasQuality": d is not None,
            "hasCost": c is not None,
            "workflowStatus": (d or {}).get("workflow_status") or "",
            "validationStatus": (d or {}).get("validation_status") or "",
            "model": doc_model,
            "processingSeconds": (d or {}).get("processing_seconds"),
            "avgConfidence": (d or {}).get("avg_confidence"),
            "subDocCount": (d or {}).get("sub_document_count") or 0,
            "warnings": (d or {}).get("warnings") or [],
            "fieldStats": (d or {}).get("field_stats") or {},
            "pricingVersion": (c or {}).get("pricing_version") or "",
            "pagesTotal": (c or {}).get("pages_total") or 0,
            "pagesNative": (c or {}).get("pages_native") or 0,
            "pagesOcr": (c or {}).get("pages_ocr") or 0,
            "pagesEmpty": (c or {}).get("pages_empty") or 0,
            "totalCost": total_cost,
            "costPerPage": (c or {}).get("cost_per_page"),
            "totalInputTokens": (c or {}).get("total_input_tokens"),
            "totalOutputTokens": (c or {}).get("total_output_tokens"),
            "services": services,
            "usesAi": uses_ai,
            "expectedNoAi": expected_no_ai,
            "processingMode": processing_mode_label(uses_ai, expected_no_ai=expected_no_ai),
            "processingEvidence": (
                "Catálogo estructurado (xpath/cell/column)" if expected_no_ai
                else "Servicios usados en la ejecución" if uses_ai is not None
                else "Sin auditoría de costo"
            ),
            "aiCost": ai_cost if c is not None else None,
            "nonAiCost": max(0.0, float(total_cost or 0) - ai_cost) if c is not None else None,
            # null cuando la selección de páginas no estuvo activa para ese
            # documento; el JS filtra por eso para no contar ceros.
            "pageSelection": (c or {}).get("page_selection"),
        })
    return unified


def hydrate_processing_fields(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Upgrade rows from schema-v1 snapshots without changing their source data."""
    for row in rows:
        services = row.get("services") or {}
        uses_ai = _uses_ai(bool(row.get("hasCost")), services)
        expected_no_ai = row.get("type") in DETERMINISTIC_TYPES
        ai_cost = _ai_cost(services)
        row["usesAi"] = uses_ai
        row["expectedNoAi"] = expected_no_ai
        row["processingMode"] = processing_mode_label(uses_ai, expected_no_ai=expected_no_ai)
        row["processingEvidence"] = (
            "Catálogo estructurado (xpath/cell/column)" if expected_no_ai
            else "Servicios usados en la ejecución" if uses_ai is not None
            else "Sin auditoría de costo"
        )
        row["aiCost"] = ai_cost if row.get("hasCost") else None
        row["nonAiCost"] = max(0.0, float(row.get("totalCost") or 0) - ai_cost) if row.get("hasCost") else None
    return rows


# ---------------------------------------------------------------------------
# Agregación server-side — solo para el .md estático y el .xlsx crudo. El
# HTML recalcula todo en el navegador (ver JS embebido) porque tiene que
# reaccionar a los filtros; esto de aquí es una foto fija sin filtrar.
# ---------------------------------------------------------------------------

def money(x: Optional[float]) -> str:
    if x is None:
        return "—"
    return f"${x:,.6f}" if 0 < abs(x) < 0.0001 else f"${x:,.4f}"


def num(x: Optional[float]) -> str:
    return f"{x:,.0f}" if x is not None else "—"


def pct(x: Optional[float]) -> str:
    return f"{x:.0%}" if x is not None else "—"


def money5(x: Optional[float]) -> str:
    """El costo por página vive en el quinto decimal ($0.00455) — `money`, con
    sus 4 decimales, lo redondearía a algo que no se puede comparar contra las
    referencias de COSTOS.md."""
    return f"${x:,.5f}" if x is not None else "—"


def secs(x: Optional[float]) -> str:
    """Mismo formato que `secsFmt` del lado JS — sub-segundo en ms, minutos
    cuando pasa del minuto, para que el .md y el HTML no se lean distinto."""
    if x is None:
        return "—"
    if x < 1:
        return f"{x * 1000:.0f}ms"
    if x < 60:
        return f"{x:.1f}s"
    if x < 3600:
        return f"{int(x // 60)}m {x % 60:.0f}s"
    return f"{int(x // 3600)}h {int(x % 3600 // 60)}m"


def percentile(sorted_values: List[float], q: float) -> Optional[float]:
    """Percentil con interpolación lineal (mismo método que numpy.percentile,
    sin la dependencia). La lista **ya viene ordenada**."""
    if not sorted_values:
        return None
    if len(sorted_values) == 1:
        return sorted_values[0]
    pos = (len(sorted_values) - 1) * q
    lo = int(pos)
    hi = min(lo + 1, len(sorted_values) - 1)
    return sorted_values[lo] + (sorted_values[hi] - sorted_values[lo]) * (pos - lo)


def app_of(row: Dict[str, Any]) -> str:
    """El sistema/app que subió el documento. Los registros viejos (o los que
    entraron sin API key en dev) no lo traen; se agrupan explícitos en vez de
    desaparecer del corte."""
    return row["appOrigin"] or "(sin sistema)"


def overall_snapshot(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    n = len(rows)
    completed = [r for r in rows if r["workflowStatus"] == "complete"]
    val_counts = Counter(r["validationStatus"] for r in completed if r["validationStatus"])
    with_cost = [r for r in rows if r["hasCost"]]
    confs = [r["avgConfidence"] for r in rows if r["avgConfidence"] is not None]
    durs = sorted(r["processingSeconds"] for r in rows if r["processingSeconds"] is not None)
    # segundos/página solo sobre los documentos que traen ambos datos —
    # dividir el tiempo total entre todas las páginas mezclaría documentos
    # sin tiempo registrado y bajaría el número artificialmente.
    dur_page = [r for r in rows if r["processingSeconds"] is not None and r["pagesTotal"]]
    total_cost = sum(r["totalCost"] or 0 for r in with_cost)
    total_pages = sum(r["pagesTotal"] or 0 for r in rows)
    dates = sorted(r["processedAt"] for r in rows if r["processedAt"])
    ai_rows = [r for r in rows if r.get("processingMode") == "Con IA"]
    no_ai_rows = [r for r in rows if r.get("processingMode") == "Sin IA"]
    ai_unknown = [r for r in rows if r.get("processingMode") == "Sin datos"]
    return {
        "n": n,
        "n_completed": len(completed),
        "approval_rate": (val_counts.get("APROBADO", 0) / len(completed)) if completed else None,
        "avg_confidence": (sum(confs) / len(confs)) if confs else None,
        "n_with_duration": len(durs),
        "avg_seconds": (sum(durs) / len(durs)) if durs else None,
        "p50_seconds": percentile(durs, 0.5),
        "p95_seconds": percentile(durs, 0.95),
        "min_seconds": durs[0] if durs else None,
        "max_seconds": durs[-1] if durs else None,
        "total_seconds": sum(durs),
        "seconds_per_page": (
            sum(r["processingSeconds"] for r in dur_page) / sum(r["pagesTotal"] for r in dur_page)
        ) if dur_page else None,
        "total_cost": total_cost,
        "avg_cost_doc": (total_cost / len(with_cost)) if with_cost else None,
        "total_input_tokens": sum(r["totalInputTokens"] or 0 for r in with_cost),
        "total_output_tokens": sum(r["totalOutputTokens"] or 0 for r in with_cost),
        "total_pages": total_pages,
        "first_date": dates[0] if dates else None,
        "last_date": dates[-1] if dates else None,
        "n_with_cost": len(with_cost),
        "n_ai": len(ai_rows),
        "n_no_ai": len(no_ai_rows),
        "n_no_ai_with_cost": sum(bool(r.get("hasCost")) for r in no_ai_rows),
        "n_ai_unknown": len(ai_unknown),
        "ai_cost": sum(r.get("aiCost") or 0 for r in with_cost),
        "non_ai_cost": sum(r.get("nonAiCost") or 0 for r in with_cost),
        # Total completo de las ejecuciones que no invocaron Bedrock. No es lo
        # mismo que non_ai_cost, que también contiene Textract/Lambda de las
        # rutas que sí usaron IA.
        "no_ai_docs_cost": sum(r.get("totalCost") or 0 for r in no_ai_rows),
        # Ahorro por selección de páginas. Se mantiene FUERA de total_cost:
        # `cost_avoided` valora una llamada que nunca se hizo, así que sumarlo
        # inflaría el gasto real del corte.
        "page_selection": _page_sel_snapshot(rows),
    }


def monthly_spend_snapshot(
    rows: List[Dict[str, Any]], infrastructure: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Actual known report spend by month; never mixes future projection into actual."""
    fixed_history = (infrastructure or {}).get("fixedMonthlyHistory") or {}
    fixed_by_month = {item.get("month"): item for item in fixed_history.get("rows") or []}
    buckets: Dict[str, Dict[str, Any]] = defaultdict(
        lambda: {"documents": 0, "documentsWithCost": 0, "processingActualUsd": 0.0}
    )
    for row in rows:
        month = (row.get("processedAt") or "")[:7]
        if not month:
            continue
        bucket = buckets[month]
        bucket["documents"] += 1
        if row.get("hasCost"):
            bucket["documentsWithCost"] += 1
            bucket["processingActualUsd"] += row.get("totalCost") or 0

    months = sorted(set(buckets) | {month for month in fixed_by_month if month})
    output_rows: List[Dict[str, Any]] = []
    for month in months:
        bucket = buckets[month]
        fixed = fixed_by_month.get(month) or {}
        status = fixed.get("status") or "closed"
        fixed_actual = fixed.get("actualUsd")
        processing_actual = round(bucket["processingActualUsd"], 12)
        # Futuro se representa como ausencia, nunca como gasto real cero. Si el
        # fijo no es atribuible, el total conocido conserva procesamiento pero
        # se marca como parcial mediante fixedActualAvailable.
        known_actual = None if status == "future" else round(processing_actual + (fixed_actual or 0), 12)
        output_rows.append({
            "month": month,
            "status": status,
            "documents": bucket["documents"],
            "documentsWithCost": bucket["documentsWithCost"],
            "processingActualUsd": processing_actual,
            "fixedActualUsd": fixed_actual,
            "fixedActualAvailable": fixed_actual is not None,
            "knownActualUsd": known_actual,
            "fixedProjectionUsd": fixed.get("projectedUsd"),
        })

    elapsed = [row for row in output_rows if row["status"] != "future"]
    return {
        "rows": output_rows,
        "documents": sum(row["documents"] for row in elapsed),
        "documentsWithCost": sum(row["documentsWithCost"] for row in elapsed),
        "processingActualUsd": round(sum(row["processingActualUsd"] for row in elapsed), 12),
        "fixedActualUsd": (
            round(sum(row["fixedActualUsd"] or 0 for row in elapsed), 12)
            if all(row["fixedActualAvailable"] for row in elapsed) else None
        ),
        "knownActualUsd": round(sum(row["knownActualUsd"] or 0 for row in elapsed), 12),
        "note": "Real conocido = procesamiento con auditoría de costo + fijo atribuible; no incluye proyecciones ni consumo variable compartido.",
    }


def _page_sel_snapshot(rows: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    active = [r["pageSelection"] for r in rows if r.get("pageSelection")]
    if not active:
        return None
    in_saved = sum(p["input_tokens_saved"] for p in active)
    in_billed = sum(p["billed_input_tokens"] for p in active)
    cost_avoided = sum(p["estimated_cost_avoided"] for p in active)
    billed_cost = sum(r["totalCost"] or 0 for r in rows if r.get("pageSelection"))
    return {
        "docs": len(active),
        "pages_considered": sum(p["pages_considered"] for p in active),
        "pages_extracted": sum(p["pages_extracted"] for p in active),
        "pages_skipped": sum(p["pages_skipped"] for p in active),
        "billed_input_tokens": in_billed,
        "billed_output_tokens": sum(p["billed_output_tokens"] for p in active),
        "input_tokens_saved": in_saved,
        "output_tokens_saved": sum(p["output_tokens_saved"] for p in active),
        "input_saving_pct": (in_saved / (in_billed + in_saved) * 100) if (in_billed + in_saved) else None,
        "cost_avoided": cost_avoided,
        "billed_cost": billed_cost,
        "cost_without_selection": billed_cost + cost_avoided,
    }


def snapshot_by(rows: List[Dict[str, Any]], key_fn) -> Dict[str, Dict[str, Any]]:
    buckets: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for r in rows:
        buckets[key_fn(r)].append(r)
    return {k: overall_snapshot(rs) for k, rs in buckets.items()}


def by_key_snapshot(rows: List[Dict[str, Any]], key: str) -> Dict[str, Dict[str, Any]]:
    return snapshot_by(rows, lambda r: r[key])


def snapshot_mode(m: Dict[str, Any]) -> str:
    present = sum(bool(m.get(key)) for key in ("n_ai", "n_no_ai", "n_ai_unknown"))
    if present > 1:
        return "Mixto"
    if m.get("n_no_ai"):
        return "Sin IA"
    if m.get("n_ai"):
        return "Con IA"
    return "Sin datos"


def sorted_types(by_type: Dict[str, Dict[str, Any]]) -> List[tuple[str, Dict[str, Any]]]:
    """Determinísticos primero y luego nombre canónico, no volumen accidental."""
    rank = {"Sin IA": 0, "Mixto": 1, "Con IA": 2, "Sin datos": 3}
    return sorted(by_type.items(), key=lambda kv: (rank[snapshot_mode(kv[1])], kv[0].casefold()))


def render_markdown(
    rows: List[Dict[str, Any]], env: str, quality: Optional[Dict[str, Any]] = None,
    infrastructure: Optional[Dict[str, Any]] = None,
) -> str:
    overall = overall_snapshot(rows)
    by_type = by_key_snapshot(rows, "type")
    by_model = by_key_snapshot(rows, "model")
    by_app = snapshot_by(rows, app_of)
    monthly_spend = monthly_spend_snapshot(rows, infrastructure)
    report_period = (infrastructure or {}).get("reportPeriod") or {}
    report_period_text = (
        f"{report_period.get('start')} a {report_period.get('end')}"
        if report_period.get("start") and report_period.get("end") else "sin rango explícito"
    )

    lines = [
        f"# Reporte — {env}",
        "",
        f"Generado {datetime.now():%Y-%m-%d %H:%M} · {overall['n']} documento(s)"
        + (f" · {overall['first_date'][:10]} a {overall['last_date'][:10]}" if overall["first_date"] else ""),
        f"Periodo cargado por defecto: **{report_period_text}**.",
        "",
        "Este resumen es una foto fija sin filtrar. Para filtrar por fecha/tipo/sistema/modelo/"
        "estado y ver el detalle de un documento, abre `reporte.html`.",
        "",
        "## Totales",
        "",
        "| Métrica | Valor |",
        "| --- | --- |",
        f"| Documentos | {overall['n']} |",
        f"| Ejecuciones completas | {overall['n_completed']} |",
        f"| Tasa de aprobación | {pct(overall['approval_rate'])} |",
        f"| Confianza promedio | {pct(overall['avg_confidence'])} |",
        f"| Documentos con datos de costo | {overall['n_with_cost']} |",
        f"| Costo de procesamiento atribuible | {money(overall['total_cost'])} |",
        f"| Costo promedio / documento | {money(overall['avg_cost_doc'])} |",
        f"| Documentos procesados sin IA | {overall['n_no_ai']} |",
        f"| Costo total de documentos sin IA | {money(overall['no_ai_docs_cost'])} |",
        f"| Páginas procesadas | {num(overall['total_pages'])} |",
        f"| Tokens de entrada (total) | {num(overall['total_input_tokens'])} |",
        f"| Tokens de salida (total) | {num(overall['total_output_tokens'])} |",
        "",
    ]

    ps = overall.get("page_selection")
    if ps:
        lines += [
            "## Ahorro por selección de páginas",
            "",
            f"Activa en {ps['docs']} de {overall['n']} documento(s): se extrajeron "
            f"{num(ps['pages_extracted'])} de {num(ps['pages_considered'])} páginas "
            f"({num(ps['pages_skipped'])} nunca llegaron al modelo).",
            "",
            "| | Facturado | Evitado (est.) | Sin selección |",
            "| --- | --- | --- | --- |",
            f"| Tokens de entrada | {num(ps['billed_input_tokens'])} | {num(ps['input_tokens_saved'])} | "
            f"{num(ps['billed_input_tokens'] + ps['input_tokens_saved'])} |",
            f"| Tokens de salida | {num(ps['billed_output_tokens'])} | {num(ps['output_tokens_saved'])} | "
            f"{num(ps['billed_output_tokens'] + ps['output_tokens_saved'])} |",
            f"| Costo | {money(ps['billed_cost'])} | {money(ps['cost_avoided'])} | "
            f"{money(ps['cost_without_selection'])} |",
            "",
            "> **Evitado** es una estimación de la llamada que nunca se hizo — no está sumado "
            "al costo total de arriba.",
            "",
        ]

    lines += [
        "## Procesamiento con y sin IA",
        "",
        "La ruta se determina por la forma del catálogo (`xpath`/`cell`/`column` = determinística) y se "
        "contrasta con los servicios realmente usados: cualquier cargo o token Bedrock significa **Con IA**. "
        "Un tipo sin fuente estructurada y sin auditoría de costo queda como **Sin datos**.",
        "",
        "| Ruta | Documentos | Costo total de esos docs | Componente IA | Resto AWS |",
        "| --- | ---: | ---: | ---: | ---: |",
        f"| Con IA | {overall['n_ai']} | {money(sum(r['totalCost'] or 0 for r in rows if r.get('processingMode') == 'Con IA'))} | "
        f"{money(overall['ai_cost'])} | {money(sum(r['nonAiCost'] or 0 for r in rows if r.get('processingMode') == 'Con IA'))} |",
        f"| **Sin IA** | **{overall['n_no_ai']}** ({overall['n_no_ai_with_cost']} con costo) | **{money(overall['no_ai_docs_cost'])}** | $0.0000 | "
        f"{money(overall['no_ai_docs_cost'])} |",
        f"| Sin datos suficientes | {overall['n_ai_unknown']} | — | — | — |",
        "",
        "## Tiempo de proceso",
        "",
        f"Medido de punta a punta (`Tracking.ProcessingSeconds`): del inicio del enriquecimiento "
        f"al resultado guardado. {overall['n_with_duration']} de {overall['n']} documento(s) lo traen registrado.",
        "",
        "| Métrica | Valor |",
        "| --- | --- |",
        f"| Promedio | {secs(overall['avg_seconds'])} |",
        f"| Mediana (p50) | {secs(overall['p50_seconds'])} |",
        f"| p95 | {secs(overall['p95_seconds'])} |",
        f"| Mínimo | {secs(overall['min_seconds'])} |",
        f"| Máximo | {secs(overall['max_seconds'])} |",
        f"| Por página | {secs(overall['seconds_per_page'])} |",
        f"| Tiempo acumulado | {secs(overall['total_seconds'])} |",
        "",
        "## Por tipo de documento",
        "",
        "| Tipo | Ruta | Docs | Docs con IA | Aprobación | Tiempo prom. | Costo total | Costo IA | $/doc |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    for t, m in sorted_types(by_type):
        lines.append(
            f"| {t} | **{snapshot_mode(m)}** | {m['n']} | {m['n_ai']} | {pct(m['approval_rate'])} | "
            f"{secs(m['avg_seconds'])} | {money(m['total_cost'])} | {money(m['ai_cost'])} | {money(m['avg_cost_doc'])} |"
        )

    lines += ["", "## Por sistema", "",
              "Sistema/app que subió el documento (`AppOrigin`, resuelto de la API key del request).",
              "",
              "| Sistema | Docs | Aprobación | Confianza | Tiempo prom. | p95 | Costo total | $/doc |",
              "| --- | --- | --- | --- | --- | --- | --- | --- |"]
    for app, m in sorted(by_app.items(), key=lambda kv: kv[1]["n"], reverse=True):
        lines.append(f"| {app} | {m['n']} | {pct(m['approval_rate'])} | {pct(m['avg_confidence'])} | {secs(m['avg_seconds'])} | {secs(m['p95_seconds'])} | {money(m['total_cost'])} | {money(m['avg_cost_doc'])} |")

    lines += ["", "## Por modelo", "", "| Modelo | Docs | Aprobación | Confianza | Tiempo prom. | p95 | Costo total | $/doc |", "| --- | --- | --- | --- | --- | --- | --- | --- |"]
    for mdl, m in sorted(by_model.items(), key=lambda kv: kv[1]["n"], reverse=True):
        lines.append(f"| {mdl} | {m['n']} | {pct(m['approval_rate'])} | {pct(m['avg_confidence'])} | {secs(m['avg_seconds'])} | {secs(m['p95_seconds'])} | {money(m['total_cost'])} | {money(m['avg_cost_doc'])} |")

    budget = next(x for x in COST_REFERENCE["perPage"] if x["role"] == "budget")
    cost_page_now = (overall["total_cost"] / overall["total_pages"]) if overall["total_pages"] else None
    lines += [
        "", "## Costo contra las referencias medidas", "",
        f"Referencias de `{COST_REFERENCE['doc']}` (medición real de {COST_REFERENCE['measuredOn']}: "
        "`usage` de Bedrock y páginas reales, cruzados contra Cost Explorer).",
        "",
        "| Concepto | $/página |",
        "| --- | --- |",
        f"| **Este corte** | {money5(cost_page_now)} |",
    ]
    for ref in COST_REFERENCE["perPage"]:
        lines.append(f"| {ref['label']} ({COST_REFERENCE['doc']} {ref['source']}) | {money5(ref['value'])} |")
    projected_processing = (cost_page_now * COST_REFERENCE["volume"]["pagesPerMonth"]) if cost_page_now is not None else None
    fixed_monthly = (infrastructure or {}).get("fixedMonthlyUsd")
    projected = (projected_processing + (fixed_monthly or 0)) if projected_processing is not None else None
    lines += [
        "",
        f"Costo por documento de este corte: {money(overall['avg_cost_doc'])} "
        f"(referencia {money(COST_REFERENCE['perDoc']['value'])} — {COST_REFERENCE['perDoc']['label']}).",
        "",
        f"A {COST_REFERENCE['volume']['pagesPerMonth']:,} páginas/mes el corte actual proyecta "
        + (f"**${projected:,.2f}/mes**" if projected is not None else "—")
        + (f" (procesamiento ${projected_processing:,.2f} + fijo {money(fixed_monthly)})" if projected_processing is not None else "")
        + f" contra el presupuesto acordado de ${COST_REFERENCE['volume']['budgetPerMonth']:,.2f}/mes "
        f"(${budget['value']:.4f}/página).",
        "",
        "| Tipo | Docs | $/doc hoy | $/doc ref. | $/pág hoy | $/pág ref. |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for t, m in sorted_types(by_type):
        ref = COST_REFERENCE["perDocByType"].get(t)
        page_now = (m["total_cost"] / m["total_pages"]) if m["total_pages"] else None
        lines.append(
            f"| {t} | {m['n']} | {money(m['avg_cost_doc'])} | {money(ref['costPerDoc']) if ref else '—'} | "
            f"{money5(page_now)} | {money5(ref['costPerPage']) if ref else '—'} |"
        )

    infra = infrastructure or {}
    lines += ["", "## Infraestructura del stack", ""]
    if infra.get("fixedMonthlyUsd") is None:
        lines.append("No se pudo obtener un inventario AWS atribuible al stack.")
    else:
        tagged = infra.get("taggedMtd") or {}
        fixed_mtd = infra.get("fixedMtdActualUsd")
        lines += [
            f"Inventario en vivo de `{infra.get('stack', '—')}` / `{infra.get('region', '—')}`, "
            f"capturado {str(infra.get('capturedAt', ''))[:19]}. El costo fijo no se reparte entre "
            "documentos ni cambia con los filtros del dashboard.",
            "",
            f"**Ambiente seleccionado:** {infra.get('environmentExplanation', '—')}",
            "",
            "| Métrica | Valor |",
            "| --- | ---: |",
            f"| Piso fijo mensual actual | **{money(infra.get('fixedMonthlyUsd'))}** |",
            f"| Fijo real MTD en Cost Explorer | {money(fixed_mtd) if fixed_mtd is not None else '—'} |",
            f"| Gasto MTD etiquetado `CostStack` | {money(tagged.get('totalUsd')) if tagged.get('available') else 'No disponible'} |",
            "",
            "| Componente fijo | Cantidad | Tarifa / mes | Total / mes |",
            "| --- | ---: | ---: | ---: |",
        ]
        for item in infra.get("fixedItems") or []:
            lines.append(
                f"| {item['service']} · {item['component']} | {item['quantity']} | "
                f"{money(item['unitMonthlyUsd'])} | {money(item['monthlyUsd'])} |"
            )
        if not infra.get("fixedItems"):
            lines.append("| Ninguno en este ambiente | 0 | $0.0000 | $0.0000 |")
        fixed_history = infra.get("fixedMonthlyHistory") or {}
        if fixed_history.get("rows"):
            lines += [
                "", "### Gasto fijo mes a mes", "",
                f"Real acumulado: **{money(fixed_history.get('actualToDateUsd'))}** · "
                f"proyección hasta {fixed_history.get('end', '—')}: "
                f"**{money(fixed_history.get('projectedPeriodUsd'))}**.", "",
                "| Mes | Estado | Real facturado | Piso actual de referencia | Usado en proyección |",
                "| --- | --- | ---: | ---: | ---: |",
            ]
            status_labels = {"closed": "Cerrado", "mtd": "Mes actual · MTD", "future": "Futuro"}
            for item in fixed_history["rows"]:
                actual = item.get("actualUsd")
                lines.append(
                    f"| {item.get('month')} | {status_labels.get(item.get('status'), item.get('status'))} | "
                    f"{money(actual) if actual is not None else '—'} | {money(item.get('estimatedFullMonthUsd'))} | "
                    f"{money(item.get('projectedUsd'))} |"
                )
        if monthly_spend.get("rows"):
            lines += [
                "", "### Gasto real del reporte, mes a mes", "",
                f"Gastado real conocido: **{money(monthly_spend.get('knownActualUsd'))}** = procesamiento "
                f"{money(monthly_spend.get('processingActualUsd'))} + fijo "
                f"{money(monthly_spend.get('fixedActualUsd'))}.",
                "La proyección anterior permanece separada y nunca se suma en esta tabla.", "",
                "| Mes | Estado | Procesamiento real | Fijo real | Total real conocido | Cobertura de costos |",
                "| --- | --- | ---: | ---: | ---: | ---: |",
            ]
            for item in monthly_spend["rows"]:
                if item.get("status") == "future":
                    continue
                fixed_actual = item.get("fixedActualUsd")
                lines.append(
                    f"| {item.get('month')} | {'Mes actual · MTD' if item.get('status') == 'mtd' else 'Cerrado'} | "
                    f"{money(item.get('processingActualUsd'))} | {money(fixed_actual) if fixed_actual is not None else '—'} | "
                    f"{money(item.get('knownActualUsd'))} | {item.get('documentsWithCost')}/{item.get('documents')} docs |"
                )
        lines += ["", "S3, DynamoDB, API Gateway/API keys y los demás recursos serverless no tienen renta fija; "
                  "se cobran por uso. Si `CostStack` no devuelve datos, el reporte cruza manualmente los servicios "
                  "presentes en CloudFormation con los cargos reales de la cuenta.", ""]
        manual = infra.get("manualMtd") or {}
        if manual.get("services"):
            lines += [
                "### Fallback manual por servicio", "",
                "`Cuenta MTD` es contexto real de facturación. Sólo `Atribuido al stack` entra como costo scopeado; "
                "un cargo de cuenta compartido no se reparte artificialmente.", "",
                "| Servicio presente en el stack | Recursos | Cuenta MTD | Atribuido al stack | Alcance |",
                "| --- | ---: | ---: | ---: | --- |",
            ]
            for item in manual["services"]:
                scoped = item.get("scopedMtdUsd")
                lines.append(
                    f"| {item['service']} | {item['resources']} | {money(item.get('accountMtdUsd'))} | "
                    f"{money(scoped) if scoped is not None else '—'} | {item.get('scope', '')} |"
                )
            lines.append("")
        for warning in infra.get("warnings") or []:
            lines.append(f"> {warning}")
            lines.append("")

    if quality:
        latest = quality.get("latestFree") or quality.get("latest") or {}
        real = quality.get("latestReal") or {}
        qs, rs = latest.get("summary") or {}, real.get("summary") or {}
        lines += [
            "", "## Calidad comprobada", "",
            f"Corridas históricas: **{quality.get('runs', 0)}**.", "",
            "| Métrica | Último valor |", "| --- | --- |",
            f"| Última suite | {latest.get('suite', '—')} · {latest.get('status', '—')} |",
            f"| Pruebas gratuitas | {qs.get('tests', '—')} ({qs.get('failures', '—')} fallas) |",
            f"| Cobertura de líneas Python | {qs.get('lineCoverage', 0):.1f}% |" if qs.get("lineCoverage") is not None else "| Cobertura de líneas Python | — |",
            f"| Cobertura de líneas Node | {qs.get('nodeLineCoverage', 0):.1f}% |" if qs.get("nodeLineCoverage") is not None else "| Cobertura de líneas Node | — |",
            f"| Último E2E completo | {real.get('status', '—')} |",
            f"| Exactitud de campos etiquetados | {pct(rs.get('fieldAccuracy'))} |",
            f"| Éxito técnico E2E | {pct(rs.get('e2ePassRate'))} |",
        ]
    lines += ["", "_Generado por `pnpm reporte` (scripts/reporte.py) — dashboard interactivo: `reporte.html` · datos crudos: `datos_raw.xlsx`._"]
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Excel (datos crudos)
# ---------------------------------------------------------------------------

def write_excel(
    out_path: Path, rows: List[Dict[str, Any]], quality: Optional[Dict[str, Any]] = None,
    infrastructure: Optional[Dict[str, Any]] = None,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.chart import LineChart, Reference

    wb = Workbook()
    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="17365D")
    stripe_fill = PatternFill("solid", fgColor="EAF2F8")

    def write_sheet(ws, headers: List[str], data_rows: List[List[Any]]) -> None:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in data_rows:
            ws.append(row)
        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(12, min(40, len(header) + 4))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row_idx in range(2, ws.max_row + 1):
            if row_idx % 2 == 0:
                for cell in ws[row_idx]:
                    cell.fill = stripe_fill

    service_names = sorted({name for r in rows for name in r["services"]})
    ws_exec = wb.active
    ws_exec.title = "Resumen Ejecutivo"
    ws_docs = wb.create_sheet("Documentos")
    ws_docs.title = "Documentos"
    headers = [
        "DocumentId", "DocumentType", "DocumentTypeName", "DocumentName", "AppOrigin", "ProcessedAt",
        "WorkflowStatus", "ValidationStatus", "ProcessingMode", "ProcessingEvidence", "UsesAi", "Model", "ProcessingSeconds", "AvgConfidence", "SubDocumentCount",
        "PagesTotal", "PagesNative", "PagesOcr", "PagesEmpty",
        "TotalCostUsd", "AiCostUsd", "NonAiCostUsd", "CostPerPageUsd", "TotalInputTokens", "TotalOutputTokens", "PricingVersion",
        "HasCostData", "HasQualityData",
    ]
    for svc in service_names:
        headers += [f"{svc}.CostUsd", f"{svc}.InputTokens", f"{svc}.OutputTokens", f"{svc}.ModelId"]
    data_rows = []
    for r in rows:
        row = [
            r["id"], r["type"], r["typeName"], r["name"], r["appOrigin"], r["processedAt"],
            r["workflowStatus"], r["validationStatus"], r["processingMode"], r["processingEvidence"], r["usesAi"], r["model"], r["processingSeconds"], r["avgConfidence"],
            r["subDocCount"], r["pagesTotal"], r["pagesNative"], r["pagesOcr"], r["pagesEmpty"],
            r["totalCost"], r["aiCost"], r["nonAiCost"], r["costPerPage"], r["totalInputTokens"], r["totalOutputTokens"], r["pricingVersion"],
            r["hasCost"], r["hasQuality"],
        ]
        for svc in service_names:
            s = r["services"].get(svc)
            row += [s["cost"], s["input_tokens"], s["output_tokens"], s["model_id"]] if s else [None, None, None, None]
        data_rows.append(row)
    write_sheet(ws_docs, headers, data_rows)

    operational = [r for r in rows if r.get("trafficKind") != "test"]
    tests_traffic = [r for r in rows if r.get("trafficKind") == "test"]
    all_docs = overall_snapshot(rows)
    write_sheet(ws_exec, ["Indicador", "Valor", "Lectura"], [
        ["Ambiente operativo", len(operational), "Documentos sin test-run-id"],
        ["Tráfico de prueba", len(tests_traffic), "Incluido por defecto; se puede filtrar en el HTML"],
        ["Éxito técnico", all_docs["n_completed"] / all_docs["n"] if all_docs["n"] else None, "Todas las ejecuciones cargadas en el periodo"],
        ["Aprobación de reglas", all_docs["approval_rate"], "Todos los documentos; no equivale a exactitud humana"],
        ["Costo total de procesamiento USD", all_docs["total_cost"], "Incluye todo el tráfico cargado en el periodo"],
        ["Documentos sin IA", all_docs["n_no_ai"], "Todos los documentos; sin uso real de servicios Bedrock"],
        ["Costo de documentos sin IA USD", all_docs["no_ai_docs_cost"], "Lambda/otros servicios; componente IA = 0"],
        ["Piso fijo mensual del stack USD", (infrastructure or {}).get("fixedMonthlyUsd"), "No se reparte por documento"],
        ["Latencia p95 segundos", all_docs["p95_seconds"], "Punta a punta; todos los documentos del periodo"],
    ])

    agg_headers = [
        "Documentos", "TasaAprobacion", "ConfianzaPromedio",
        "DocsConTiempo", "SegundosPromedio", "SegundosP50", "SegundosP95", "SegundosMax", "SegundosPorPagina",
        "DocsConIA", "DocsSinIA", "DocsSinIAConCosto", "DocsRutaDesconocida", "CostoIAUsd", "CostoNoIAUsd", "CostoDocsSinIAUsd",
        "CostoTotalUsd", "CostoPorDocUsd", "CostoPorPaginaUsd",
    ]

    def agg_values(m: Dict[str, Any]) -> List[Any]:
        return [
            m["n"], m["approval_rate"], m["avg_confidence"],
            m["n_with_duration"], m["avg_seconds"], m["p50_seconds"], m["p95_seconds"], m["max_seconds"], m["seconds_per_page"],
            m["n_ai"], m["n_no_ai"], m["n_no_ai_with_cost"], m["n_ai_unknown"], m["ai_cost"], m["non_ai_cost"], m["no_ai_docs_cost"],
            m["total_cost"], m["avg_cost_doc"], (m["total_cost"] / m["total_pages"]) if m["total_pages"] else None,
        ]

    ws_type = wb.create_sheet("PorTipo")
    by_type = by_key_snapshot(rows, "type")
    write_sheet(
        ws_type, ["DocumentType", "ProcessingMode"] + agg_headers,
        [[t, snapshot_mode(m)] + agg_values(m) for t, m in sorted_types(by_type)],
    )

    ws_app = wb.create_sheet("PorSistema")
    by_app = snapshot_by(rows, app_of)
    write_sheet(ws_app, ["AppOrigin"] + agg_headers, [[a] + agg_values(m) for a, m in by_app.items()])

    ws_model = wb.create_sheet("PorModelo")
    by_model = by_key_snapshot(rows, "model")
    write_sheet(ws_model, ["Model"] + agg_headers, [[m_] + agg_values(s) for m_, s in by_model.items()])

    # Referencias medidas de COSTOS.md contra lo que trae este corte, para
    # poder comparar en Excel sin volver a teclear los números del documento.
    ws_ref = wb.create_sheet("Referencias")
    ref_rows: List[List[Any]] = [
        [f"{r['label']} ({COST_REFERENCE['doc']} {r['source']})", "$/pagina", r["value"], None, None]
        for r in COST_REFERENCE["perPage"]
    ]
    overall_ref = overall_snapshot(rows)
    ref_rows.append([
        "Este corte", "$/pagina",
        (overall_ref["total_cost"] / overall_ref["total_pages"]) if overall_ref["total_pages"] else None,
        None, None,
    ])
    for t, m in by_type.items():
        ref = COST_REFERENCE["perDocByType"].get(t)
        ref_rows.append([
            t, "$/doc", ref["costPerDoc"] if ref else None, m["avg_cost_doc"],
            (m["avg_cost_doc"] - ref["costPerDoc"]) if (ref and m["avg_cost_doc"] is not None) else None,
        ])
    write_sheet(ws_ref, ["Concepto", "Unidad", "Referencia", "Medido", "Diferencia"], ref_rows)

    ws_svc = wb.create_sheet("PorServicio")
    svc_agg: Dict[str, Dict[str, Any]] = defaultdict(lambda: {"cost": 0.0, "input_tokens": 0, "output_tokens": 0, "docs": 0})
    for r in rows:
        for name, s in r["services"].items():
            if not s["cost"] and not s["input_tokens"] and not s["output_tokens"]:
                continue
            agg = svc_agg[name]
            agg["cost"] += s["cost"]
            agg["input_tokens"] += s["input_tokens"]
            agg["output_tokens"] += s["output_tokens"]
            agg["docs"] += 1
    write_sheet(
        ws_svc,
        ["Servicio", "Documentos", "InputTokens", "OutputTokens", "CostoUsd"],
        [[name, a["docs"], a["input_tokens"], a["output_tokens"], a["cost"]] for name, a in svc_agg.items()],
    )

    ws_vol = wb.create_sheet("VolumenDiario")
    vol: Dict[str, Dict[str, Any]] = defaultdict(lambda: {"n": 0, "cost": 0.0, "val": Counter()})
    for r in rows:
        day = (r["processedAt"] or "")[:10]
        if not day:
            continue
        v = vol[day]
        v["n"] += 1
        v["cost"] += r["totalCost"] or 0
        if r["workflowStatus"] == "complete" and r["validationStatus"]:
            v["val"][r["validationStatus"]] += 1
    write_sheet(
        ws_vol,
        ["Dia", "Documentos", "CostoUsd", "Aprobado", "RequiereRevision", "Rechazado"],
        [
            [day, v["n"], v["cost"], v["val"].get("APROBADO", 0), v["val"].get("REQUIERE_REVISION", 0), v["val"].get("RECHAZADO", 0)]
            for day, v in sorted(vol.items())
        ],
    )

    overall = overall_snapshot(rows)
    ws_tot = wb.create_sheet("Totales")
    write_sheet(
        ws_tot,
        ["Metrica", "Valor"],
        [
            ["Documentos", overall["n"]],
            ["Ejecuciones completas", overall["n_completed"]],
            ["Tasa de aprobacion", overall["approval_rate"]],
            ["Confianza promedio", overall["avg_confidence"]],
            ["Documentos con tiempo registrado", overall["n_with_duration"]],
            ["Segundos promedio por documento", overall["avg_seconds"]],
            ["Segundos p50", overall["p50_seconds"]],
            ["Segundos p95", overall["p95_seconds"]],
            ["Segundos maximo", overall["max_seconds"]],
            ["Segundos por pagina", overall["seconds_per_page"]],
            ["Segundos acumulados", overall["total_seconds"]],
            ["Documentos con datos de costo", overall["n_with_cost"]],
            ["Documentos con IA", overall["n_ai"]],
            ["Documentos sin IA", overall["n_no_ai"]],
            ["Documentos sin IA con costo registrado", overall["n_no_ai_with_cost"]],
            ["Documentos sin dato de ruta IA", overall["n_ai_unknown"]],
            ["Costo de componentes IA USD", overall["ai_cost"]],
            ["Costo de componentes no IA USD", overall["non_ai_cost"]],
            ["Costo total de documentos sin IA USD", overall["no_ai_docs_cost"]],
            ["Costo total USD", overall["total_cost"]],
            ["Costo promedio por documento USD", overall["avg_cost_doc"]],
            ["Costo por pagina USD", (overall["total_cost"] / overall["total_pages"]) if overall["total_pages"] else None],
            ["Costo por documento de referencia USD (COSTOS.md)", COST_REFERENCE["perDoc"]["value"]],
            ["Presupuesto por pagina USD (COSTOS.md)", next(r["value"] for r in COST_REFERENCE["perPage"] if r["role"] == "budget")],
            ["Paginas procesadas", overall["total_pages"]],
            ["Tokens de entrada totales", overall["total_input_tokens"]],
            ["Tokens de salida totales", overall["total_output_tokens"]],
            ["Primer documento", overall["first_date"]],
            ["Ultimo documento", overall["last_date"]],
        ],
    )

    infra = infrastructure or {}
    ws_infra = wb.create_sheet("Infraestructura")
    infra_rows: List[List[Any]] = [
        ["Resumen", "Piso fijo mensual", None, None, infra.get("fixedMonthlyUsd"), "Inventario CloudFormation"],
        ["Resumen", "Fijo real MTD", None, None, infra.get("fixedMtdActualUsd"), "Cost Explorer, sólo con exclusividad comprobada"],
        ["Resumen", "Ambiente seleccionado", None, None, None, infra.get("environmentExplanation")],
        ["Resumen", "Gasto MTD etiquetado", None, None,
         (infra.get("taggedMtd") or {}).get("totalUsd") if (infra.get("taggedMtd") or {}).get("available") else None,
         "CostStack"],
    ]
    for item in infra.get("fixedItems") or []:
        infra_rows.append([
            item.get("service"), item.get("component"), item.get("quantity"), item.get("unitMonthlyUsd"),
            item.get("monthlyUsd"), item.get("source"),
        ])
    for item in infra.get("usageOnlyItems") or []:
        infra_rows.append([
            item.get("service"), item.get("billing"), item.get("resources"), None,
            item.get("fixedMonthlyUsd"), "Costo variable por consumo",
        ])
    for item in (infra.get("manualMtd") or {}).get("services") or []:
        infra_rows.append([
            item.get("service"), "Fallback manual; cuenta MTD " + str(item.get("accountMtdUsd")),
            item.get("resources"), None, item.get("scopedMtdUsd"), item.get("scope"),
        ])
    for warning in infra.get("warnings") or []:
        infra_rows.append(["Advertencia", warning, None, None, None, None])
    write_sheet(ws_infra, ["Servicio", "Componente/lectura", "Cantidad", "TarifaMensualUsd", "TotalUsd", "Fuente"], infra_rows)

    fixed_history = infra.get("fixedMonthlyHistory") or {}
    ws_fixed = wb.create_sheet("CostosFijosMensuales")
    write_sheet(
        ws_fixed,
        ["Mes", "Estado", "RealFacturadoUsd", "PisoActualReferenciaUsd", "UsadoEnProyeccionUsd", "Alcance"],
        [[
            item.get("month"), item.get("status"), item.get("actualUsd"),
            item.get("estimatedFullMonthUsd"), item.get("projectedUsd"), fixed_history.get("scope"),
        ] for item in fixed_history.get("rows") or []],
    )

    monthly_spend = monthly_spend_snapshot(rows, infrastructure)
    ws_monthly_spend = wb.create_sheet("GastoMensualReporte")
    write_sheet(
        ws_monthly_spend,
        ["Mes", "Estado", "Documentos", "DocsConCosto", "ProcesamientoRealUsd", "FijoRealUsd", "TotalRealConocidoUsd", "FijoUsadoEnProyeccionUsd"],
        [[
            item.get("month"), item.get("status"), item.get("documents"), item.get("documentsWithCost"),
            item.get("processingActualUsd"), item.get("fixedActualUsd"), item.get("knownActualUsd"),
            item.get("fixedProjectionUsd"),
        ] for item in monthly_spend.get("rows") or []],
    )

    if quality:
        trend = quality.get("trend") or []
        ws_q = wb.create_sheet("CalidadPruebas")
        write_sheet(ws_q, ["Corrida", "Fecha", "Suite", "Estado", "Pruebas", "Fallas", "PythonLineas", "PythonRamas", "NodeLineas", "NodeRamas", "ExactitudCampos", "ExitoE2E", "DuracionSeg"], [
            [r.get("runId"), r.get("at"), r.get("suite"), r.get("status"), r.get("tests"), r.get("failures"),
             r.get("lineCoverage"), r.get("branchCoverage"), r.get("nodeLineCoverage"), r.get("nodeBranchCoverage"),
             r.get("fieldAccuracy"), r.get("e2ePassRate"), r.get("durationSeconds")]
            for r in trend
        ])
        for row in range(2, ws_q.max_row + 1):
            for col in (11, 12):
                ws_q.cell(row, col).number_format = "0.0%"
        if len(trend) >= 2:
            chart = LineChart()
            chart.title = "Tendencia de cobertura"
            chart.y_axis.title = "%"
            chart.add_data(Reference(ws_q, min_col=7, max_col=8, min_row=1, max_row=ws_q.max_row), titles_from_data=True)
            chart.set_categories(Reference(ws_q, min_col=1, min_row=2, max_row=ws_q.max_row))
            ws_q.add_chart(chart, "M2")

    wb.save(out_path)


# ---------------------------------------------------------------------------
# HTML interactivo
# ---------------------------------------------------------------------------

EXTRA_CSS = '''
.filters-card { position: sticky; top: 12px; z-index: 5; padding: 16px 22px; }
.filters-head { display: flex; align-items: center; gap: 14px; flex-wrap: wrap; }
.filters-toggle { display: inline-flex; align-items: center; gap: 8px; background: var(--plane); border: 1px solid var(--border);
  color: var(--text-primary); border-radius: 999px; padding: 7px 16px; font-size: 13.5px; font-weight: 600; cursor: pointer; }
.filters-toggle:hover { border-color: var(--accent); }
.ft-caret { color: var(--text-muted); font-size: 11px; }
.ft-badge { background: var(--accent); color: #fff; border-radius: 999px; padding: 1px 8px; font-size: 11.5px; font-weight: 600; }
.filters { display: flex; flex-wrap: wrap; gap: 18px; align-items: flex-start;
  margin-top: 16px; padding-top: 16px; border-top: 1px solid var(--gridline); max-height: 52vh; overflow-y: auto; }
/* `display: flex` de arriba le gana a la regla `[hidden]` del user-agent, así
   que el panel colapsado hay que apagarlo explícitamente. */
.filters[hidden] { display: none; }
.filter-group { display: flex; flex-direction: column; gap: 6px; }
.filter-group .fg-label { font-size: 11.5px; color: var(--text-muted); text-transform: uppercase; letter-spacing: .03em; }
.date-inputs { display: flex; gap: 8px; align-items: center; }
.date-inputs input[type=date] { background: var(--plane); color: var(--text-primary); border: 1px solid var(--border); border-radius: 8px; padding: 6px 8px; font-size: 13px; }
.filter-search { width: min(100%, 360px); background: var(--plane); color: var(--text-primary); border: 1px solid var(--border); border-radius: 8px; padding: 7px 10px; font-size: 13px; outline: none; }
.filter-search:focus { border-color: var(--accent); }
.chip-row { display: flex; flex-wrap: wrap; gap: 6px; max-width: 480px; }
.chip { border: 1px solid var(--border); background: var(--plane); color: var(--text-secondary); border-radius: 999px; padding: 4px 12px; font-size: 12.5px; cursor: pointer; user-select: none; white-space: nowrap; }
.chip.active { background: var(--accent); border-color: var(--accent); color: #fff; }
.chip-more { border-style: dashed; color: var(--accent); font-weight: 600; }
#f-reset { background: var(--plane); border: 1px solid var(--border); color: var(--text-secondary); border-radius: 8px; padding: 7px 14px; font-size: 13px; cursor: pointer; }
#f-reset:hover { color: var(--text-primary); }
.filters-summary { color: var(--text-secondary); font-size: 12.5px; margin-left: auto; text-align: right; }
details.card > summary { cursor: pointer; font-size: 19px; font-weight: 600; margin: 0 0 4px; list-style: none; }
details.card > summary::-webkit-details-marker { display: none; }
details.card > summary::before { content: "▸ "; color: var(--text-muted); font-size: 15px; }
details.card[open] > summary::before { content: "▾ "; }
details.card[open] > summary { margin-bottom: 14px; }
th.sortable { cursor: pointer; }
th.sortable:hover { color: var(--text-primary); }
th.sortable .arrow { opacity: .5; margin-left: 3px; }
tbody tr.doc-row { cursor: pointer; }
tbody tr.doc-row:hover { background: var(--plane); }
dialog#detail-dialog { border: none; border-radius: 14px; padding: 0; background: var(--surface); color: var(--text-primary); max-width: 720px; width: 92vw; max-height: 86vh; }
dialog#detail-dialog::backdrop { background: rgba(0,0,0,0.5); }
.detail-inner { padding: 24px 26px; overflow-y: auto; max-height: 86vh; }
.detail-inner h3 { margin: 18px 0 8px; font-size: 15px; }
.detail-inner h3:first-child { margin-top: 0; }
.detail-close { position: sticky; top: 10px; float: right; background: var(--plane); border: 1px solid var(--border); color: var(--text-secondary); border-radius: 999px; width: 30px; height: 30px; cursor: pointer; font-size: 15px; }
.kv-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 6px 16px; font-size: 13px; }
.kv-grid .k { color: var(--text-muted); }
.warn-list { margin: 4px 0 0; padding-left: 18px; font-size: 13px; color: var(--text-secondary); }
.more-btn { margin-top: 12px; background: var(--plane); border: 1px solid var(--border); color: var(--text-secondary); border-radius: 8px; padding: 7px 14px; font-size: 13px; cursor: pointer; }
.empty-hint { color: var(--text-muted); font-size: 13px; }
.kpi-row.dense { grid-template-columns: repeat(auto-fit, minmax(128px, 1fr)); gap: 10px; margin-bottom: 18px; }
.kpi-row.dense .kpi-tile { padding: 12px 14px; }
.kpi-row.dense .kpi-value { font-size: 19px; }
.sub-head { font-size: 14px; font-weight: 600; margin: 22px 0 10px; }
.sub-head:first-of-type { margin-top: 8px; }
td.delta-good { color: #0ca30c; }
td.delta-warn { color: #b07d00; }
td.delta-bad { color: #d03b3b; }
.ref-line { display: grid; grid-template-columns: 300px 1fr 96px; align-items: center; gap: 10px; font-size: 13px; }
.ref-line .ref-label { color: var(--text-secondary); white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
.ref-line .ref-track { position: relative; background: var(--gridline); border-radius: 6px; height: 14px; }
.ref-line .ref-fill { position: absolute; left: 0; top: 0; bottom: 0; border-radius: 6px; }
.ref-line .ref-mark { position: absolute; top: -3px; bottom: -3px; width: 2px; background: var(--text-primary); }
.ref-value { text-align: right; font-variant-numeric: tabular-nums; }
.chart-point[data-tooltip]:focus-visible { outline: 2px solid var(--accent); outline-offset: 3px; }
.chart-legend { display: flex; flex-wrap: wrap; gap: 7px 16px; margin: 8px 0 12px; color: var(--text-secondary); font-size: 12px; }
.legend-item { display: inline-flex; align-items: center; gap: 6px; }
.legend-swatch, .series-swatch { display: inline-block; width: 11px; height: 11px; border-radius: 3px; flex: 0 0 11px; border: 1px solid rgba(0,0,0,.12); }
.series-label { display: inline-flex; align-items: center; gap: 7px; min-width: 0; max-width: 100%; }
.series-label-text { overflow: hidden; text-overflow: ellipsis; }
.report-tooltip { position: fixed; z-index: 1000; max-width: min(360px, calc(100vw - 24px)); padding: 9px 11px;
  border: 1px solid var(--border); border-radius: 8px; background: var(--text-primary); color: var(--surface);
  box-shadow: 0 8px 28px rgba(0,0,0,.24); font-size: 12.5px; line-height: 1.4; white-space: pre-line;
  pointer-events: none; opacity: 0; visibility: hidden; transform: translateY(3px); transition: opacity .08s, transform .08s; }
.report-tooltip.visible { opacity: 1; visibility: visible; transform: translateY(0); }
.volume-plot { display: grid; grid-template-columns: 48px minmax(0, 1fr); gap: 9px; align-items: stretch; }
.volume-y { height: 180px; display: flex; flex-direction: column; justify-content: space-between; align-items: flex-end;
  color: var(--text-muted); font-size: 11px; font-variant-numeric: tabular-nums; padding-bottom: 1px; }
.volume-main { min-width: 0; }
.volume-plot-area { height: 180px; border-left: 1px solid var(--gridline); border-bottom: 1px solid var(--gridline);
  background: repeating-linear-gradient(to bottom, transparent 0, transparent calc(25% - 1px), var(--gridline) 25%); }
.volume-bars { height: 100%; display: flex; align-items: flex-end; gap: 3px; }
.volume-axis-note { color: var(--text-muted); font-size: 11px; margin-top: 5px; }
.monthly-cost-chart { display: grid; gap: 7px; }
.monthly-cost-row { display: grid; grid-template-columns: 92px minmax(180px, 1fr) 230px; gap: 12px; align-items: center; padding: 7px 9px; border-radius: 8px; }
.monthly-cost-row.current { background: rgba(42,120,214,.08); outline: 1px solid rgba(42,120,214,.22); }
.monthly-name { font-size: 12.5px; font-weight: 600; text-transform: capitalize; }
.monthly-status { display: block; color: var(--text-muted); font-size: 10px; font-weight: 500; margin-top: 2px; }
.monthly-bars { display: grid; gap: 3px; }
.monthly-track { height: 7px; background: var(--gridline); border-radius: 999px; overflow: hidden; }
.monthly-fill { height: 100%; border-radius: 999px; }
.monthly-stack-track { display: flex; height: 11px; background: var(--gridline); border-radius: 999px; overflow: hidden; }
.monthly-stack-segment { height: 100%; }
.monthly-values { color: var(--text-secondary); font-size: 12px; font-variant-numeric: tabular-nums; }
.mode-badge { display: inline-block; border-radius: 999px; padding: 2px 9px; font-size: 11.5px; font-weight: 650; white-space: nowrap; }
.mode-no-ai { background: rgba(27,175,122,.14); color: #0b855a; }
.mode-ai { background: rgba(42,120,214,.14); color: #2367b5; }
.mode-mixed { background: rgba(237,161,0,.16); color: #946600; }
.mode-unknown { background: var(--plane); color: var(--text-muted); }
.infra-note { border-left: 3px solid var(--accent); padding-left: 12px; color: var(--text-secondary); font-size: 13px; }
@media (max-width: 680px) { .monthly-cost-row { grid-template-columns: 76px minmax(110px, 1fr); } .monthly-values { grid-column: 1 / -1; } }
@media (max-width: 560px) { .volume-plot { grid-template-columns: 38px minmax(0, 1fr); gap: 6px; } .volume-y { font-size: 10px; } }
'''

JS_TEMPLATE = r'''
const DATA = JSON.parse(document.getElementById('report-data').textContent);
const ROWS = DATA.rows;
const CATALOG_TYPES = DATA.catalogTypes || [];
const COST_REF = DATA.costRef;
const INFRA = DATA.infrastructure || {};
const MONTHLY_SPEND = DATA.monthlySpend || {rows:[]};
const DEFAULT_RANGE = INFRA.reportPeriod || {};
const ROWS_BY_ID = new Map(ROWS.map(r => [r.id, r]));
const CATEGORICAL = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"];
const STATUS_ROLE_COLOR = {good: "#0ca30c", warning: "#fab219", critical: "#d03b3b", neutral: "#898781"};
const VALIDATION_ROLE = {APROBADO: "good", REQUIERE_REVISION: "warning", RECHAZADO: "critical"};
const WORKFLOW_ROLE = {"complete":"good","pending":"neutral","classification-complete":"neutral","processing-complete":"neutral","classification-failure":"critical","processing-failure":"critical","enrichment-failure":"critical"};
const STATUS_LABEL = {
  complete: 'Completa', pending: 'Pendiente',
  'classification-complete': 'Clasificación completa', 'processing-complete': 'Procesamiento completo',
  'classification-failure': 'Falló clasificación', 'processing-failure': 'Falló procesamiento',
  'enrichment-failure': 'Falló enriquecimiento', APROBADO: 'Aprobado',
  REQUIERE_REVISION: 'Requiere revisión', RECHAZADO: 'Rechazado',
};
const SERVICE_LABEL = {
  BedrockClassificationUpstream: "Bedrock · Clasificación (upstream)",
  BedrockProcessingUpstream: "Bedrock · Procesamiento (upstream)",
  BedrockSegmentation: "Bedrock · Segmentación",
  BedrockExtraction: "Bedrock · Extracción",
  TextractOcr: "Textract · OCR",
  LambdaCompute: "Lambda · Cómputo (aprox.)",
};
// Cubetas del histograma de tiempo. El corte a 5 min es deliberado: arriba de
// eso ya no importa el detalle, importa que se salió del rango normal.
const DURATION_BUCKETS = [
  ["< 5s", 0, 5], ["5–10s", 5, 10], ["10–20s", 10, 20], ["20–30s", 20, 30],
  ["30–60s", 30, 60], ["1–2 min", 60, 120], ["2–5 min", 120, 300], ["> 5 min", 300, Infinity],
];

function esc(s){ const d=document.createElement('div'); d.textContent = (s===null||s===undefined) ? '' : String(s); return d.innerHTML; }
function money(x){
  if (x===null||x===undefined) return '—';
  const n = Number(x), digits = Math.abs(n) > 0 && Math.abs(n) < 0.0001 ? 6 : 4;
  return '$'+n.toLocaleString('en-US',{minimumFractionDigits:digits,maximumFractionDigits:digits});
}
function num(x){ return (x===null||x===undefined) ? '—' : Number(x).toLocaleString('en-US',{maximumFractionDigits:0}); }
function pctFmt(x){ return (x===null||x===undefined) ? '—' : (x*100).toFixed(0)+'%'; }
function exactNumber(x, maxDigits=12) {
  if (x===null || x===undefined || !Number.isFinite(Number(x))) return '—';
  return Number(x).toLocaleString('en-US', {useGrouping:true, maximumFractionDigits:maxDigits});
}
function exactMoney(x){ return x===null || x===undefined ? '—' : '$'+exactNumber(x, 12)+' USD'; }
function exactPct(x){ return x===null || x===undefined ? '—' : exactNumber(Number(x)*100, 8)+'%'; }
function exactSeconds(x){ return x===null || x===undefined ? '—' : exactNumber(x, 8)+' segundos'; }
function secsFmt(x){
  if (x===null||x===undefined) return '—';
  if (x < 1) return Math.round(x*1000)+'ms';
  if (x < 60) return x.toFixed(1)+'s';
  if (x < 3600) return Math.floor(x/60)+'m '+Math.round(x%60)+'s';
  return Math.floor(x/3600)+'h '+Math.floor((x%3600)/60)+'m';
}
function money5(x){ return (x===null||x===undefined) ? '—' : '$'+Number(x).toFixed(5); }
function tooltipAttrs(text, keyboard=true) {
  if (!text) return '';
  return `data-tooltip="${esc(text)}" aria-label="${esc(text)}"${keyboard ? ' tabindex="0"' : ''}`;
}
function tooltipText(title, ...lines) {
  return [title, ...lines.flat().filter(x => x!==null && x!==undefined && x!=='')].join('\n');
}
function pctl(sorted, q){
  if (!sorted.length) return null;
  if (sorted.length === 1) return sorted[0];
  const pos = (sorted.length-1)*q, lo = Math.floor(pos), hi = Math.min(lo+1, sorted.length-1);
  return sorted[lo] + (sorted[hi]-sorted[lo])*(pos-lo);
}
// El sistema/app que subió el documento. Los registros viejos (o los que
// entraron sin API key en dev) no lo traen: se agrupan explícitos en vez de
// desaparecer del corte.
function appOf(r){ return r.appOrigin || '(sin sistema)'; }

// ---------------------------------------------------------------------
// Estado de filtros
// ---------------------------------------------------------------------
const state = {
  desde: DEFAULT_RANGE.start || '', hasta: DEFAULT_RANGE.end || '',
  types: new Set(), modes: new Set(), apps: new Set(), models: new Set(), workflow: new Set(), validation: new Set(),
  traffic: new Set(),
  hiddenServices: new Set(),
  sortKey: 'processedAt', sortDir: -1,
  docLimit: 200,
  filtersOpen: false,   // el panel arranca cerrado — ver setFiltersOpen()
  typesExpanded: false,
};

// Los principales se calculan en cada generación por volumen real, no son una
// lista de TYPE_IDs. Un tipo nuevo aparece automáticamente en el universo y en
// el contador +N; si entra al top por volumen se vuelve principal solo.
const TYPE_FILTER_PRIMARY_LIMIT = 6;

function distinct(key) { return [...new Set(ROWS.map(r => r[key]).filter(Boolean))].sort(); }

function filteredRows() {
  return ROWS.filter(r => {
    if (state.desde && (r.processedAt||'').slice(0,10) < state.desde) return false;
    if (state.hasta && (r.processedAt||'').slice(0,10) > state.hasta) return false;
    if (state.types.size && !state.types.has(r.type)) return false;
    if (state.modes.size && !state.modes.has(r.processingMode)) return false;
    if (state.apps.size && !state.apps.has(appOf(r))) return false;
    if (state.models.size && !state.models.has(r.model)) return false;
    if (state.workflow.size && !state.workflow.has(r.workflowStatus)) return false;
    if (state.validation.size && !state.validation.has(r.validationStatus)) return false;
    if (state.traffic.size && !state.traffic.has(r.trafficKind)) return false;
    return true;
  });
}

// ---------------------------------------------------------------------
// Agregación (misma forma que el lado Python, recalculada en cliente)
// ---------------------------------------------------------------------
function aggOverall(rows) {
  const n = rows.length;
  const wfCounts = {};
  const completed = [];
  rows.forEach(r => { const k = r.workflowStatus || '(sin dato)'; wfCounts[k] = (wfCounts[k]||0)+1; if (r.workflowStatus === 'complete') completed.push(r); });
  const valCounts = {};
  completed.forEach(r => { if (r.validationStatus) valCounts[r.validationStatus] = (valCounts[r.validationStatus]||0)+1; });
  const approved = valCounts['APROBADO'] || 0;
  const confs = rows.filter(r => r.avgConfidence !== null && r.avgConfidence !== undefined).map(r => r.avgConfidence);
  const durs = rows.filter(r => r.processingSeconds !== null && r.processingSeconds !== undefined)
                   .map(r => r.processingSeconds).sort((a,b) => a-b);
  // segundos/página solo sobre documentos que traen ambos datos — repartir el
  // tiempo total entre TODAS las páginas metería documentos sin tiempo
  // registrado en el denominador y bajaría el número artificialmente.
  const durPage = rows.filter(r => r.processingSeconds !== null && r.processingSeconds !== undefined && r.pagesTotal);
  const withCost = rows.filter(r => r.hasCost);
  const totalCost = withCost.reduce((a,r) => a + (r.totalCost||0), 0);
  const totalPages = rows.reduce((a,r) => a + (r.pagesTotal||0), 0);
  const totalIn = withCost.reduce((a,r) => a + (r.totalInputTokens||0), 0);
  const totalOut = withCost.reduce((a,r) => a + (r.totalOutputTokens||0), 0);
  const dates = rows.map(r => r.processedAt).filter(Boolean).sort();
  return {
    n, wfCounts, nCompleted: completed.length, valCounts,
    approvalRate: completed.length ? approved/completed.length : null,
    avgConfidence: confs.length ? confs.reduce((a,b)=>a+b,0)/confs.length : null,
    nWithDuration: durs.length,
    avgDuration: durs.length ? durs.reduce((a,b)=>a+b,0)/durs.length : null,
    p50Duration: pctl(durs, 0.5), p95Duration: pctl(durs, 0.95),
    minDuration: durs.length ? durs[0] : null,
    maxDuration: durs.length ? durs[durs.length-1] : null,
    totalDuration: durs.reduce((a,b)=>a+b,0),
    secondsPerPage: durPage.length
      ? durPage.reduce((a,r)=>a+r.processingSeconds,0) / durPage.reduce((a,r)=>a+r.pagesTotal,0) : null,
    totalCost, nWithCost: withCost.length,
    avgCostDoc: withCost.length ? totalCost/withCost.length : null,
    avgCostPage: totalPages ? totalCost/totalPages : null,
    totalPages, totalInputTokens: totalIn, totalOutputTokens: totalOut,
    nAi: rows.filter(r => r.processingMode === 'Con IA').length,
    nNoAi: rows.filter(r => r.processingMode === 'Sin IA').length,
    nNoAiWithCost: rows.filter(r => r.processingMode === 'Sin IA' && r.hasCost).length,
    nAiUnknown: rows.filter(r => r.processingMode === 'Sin datos').length,
    aiCost: withCost.reduce((a,r) => a + (r.aiCost||0), 0),
    nonAiCost: withCost.reduce((a,r) => a + (r.nonAiCost||0), 0),
    noAiDocsCost: rows.filter(r => r.processingMode === 'Sin IA').reduce((a,r) => a + (r.totalCost||0), 0),
    firstDate: dates[0]||null, lastDate: dates[dates.length-1]||null,
    pageSel: pageSelAgg(rows),
  };
}

// Ahorro por selección de páginas, recalculado con el corte filtrado igual que
// todo lo demás. Devuelve null si ningún documento del corte la tuvo activa,
// para no pintar una sección de ceros. `costAvoided` NUNCA entra en totalCost:
// valora una llamada que no se hizo, y sumarla inflaría el gasto real.
// La tarjeta se esconde entera cuando el corte filtrado no tiene ningún
// documento con selección activa — mejor ausente que llena de ceros.
function renderPageSel(ps, nDocs) {
  const card = document.getElementById('page-sel-card');
  if (!ps) { card.hidden = true; return; }
  card.hidden = false;
  document.getElementById('page-sel-kpis').innerHTML = [
    kpiTile(money(ps.costAvoided), 'Costo evitado (est.)', `de ~${money(ps.costWithoutSelection)} a ${money(ps.billedCost)}`),
    kpiTile(ps.inputSavingPct === null ? '—' : pctFmt(ps.inputSavingPct/100), 'Input tokens ahorrados'),
    kpiTile(num(ps.pagesSkipped), 'Páginas no extraídas', `de ${num(ps.pagesConsidered)}`),
    kpiTile(`${ps.docs} de ${nDocs}`, 'Documentos con selección'),
  ].join('');
  const row = (label, billed, saved, fmt) =>
    `<tr><td>${label}</td><td class="num">${fmt(billed)}</td><td class="num">${fmt(saved)}</td><td class="num">${fmt(billed+saved)}</td></tr>`;
  document.getElementById('page-sel-table').innerHTML = `
    <div class="table-wrap"><table>
      <thead><tr><th></th><th class="num">Facturado</th><th class="num">Evitado (est.)</th><th class="num">Sin selección</th></tr></thead>
      <tbody>
        ${row('Tokens de entrada', ps.billedInputTokens, ps.inputTokensSaved, num)}
        ${row('Tokens de salida', ps.billedOutputTokens, ps.outputTokensSaved, num)}
        ${row('Páginas extraídas', ps.pagesExtracted, ps.pagesSkipped, num)}
        ${row('Costo', ps.billedCost, ps.costAvoided, money)}
      </tbody>
    </table></div>`;
}

function pageSelAgg(rows) {
  const act = rows.filter(r => r.pageSelection);
  if (!act.length) return null;
  const s = (f) => act.reduce((a,r) => a + (r.pageSelection[f]||0), 0);
  const inBilled = s('billed_input_tokens'), inSaved = s('input_tokens_saved');
  const billedCost = act.reduce((a,r) => a + (r.totalCost||0), 0);
  const costAvoided = s('estimated_cost_avoided');
  return {
    docs: act.length,
    pagesConsidered: s('pages_considered'), pagesExtracted: s('pages_extracted'),
    pagesSkipped: s('pages_skipped'),
    billedInputTokens: inBilled, inputTokensSaved: inSaved,
    billedOutputTokens: s('billed_output_tokens'), outputTokensSaved: s('output_tokens_saved'),
    inputSavingPct: (inBilled+inSaved) ? inSaved/(inBilled+inSaved)*100 : null,
    billedCost, costAvoided, costWithoutSelection: billedCost + costAvoided,
  };
}

function aggBy(rows, keyFn) {
  const buckets = new Map();
  rows.forEach(r => { const k = keyFn(r); if (!buckets.has(k)) buckets.set(k, []); buckets.get(k).push(r); });
  const out = [];
  for (const [k, rs] of buckets) out.push([k, aggOverall(rs)]);
  return out;
}

function snapshotMode(m) {
  const present = [m.nAi, m.nNoAi, m.nAiUnknown].filter(Boolean).length;
  if (present > 1) return 'Mixto';
  if (m.nNoAi) return 'Sin IA';
  if (m.nAi) return 'Con IA';
  return 'Sin datos';
}

function modeRank(mode) { return ({'Sin IA':0, 'Mixto':1, 'Con IA':2, 'Sin datos':3})[mode] ?? 4; }
function modeBadge(mode) {
  const cls = mode === 'Sin IA' ? 'mode-no-ai' : mode === 'Con IA' ? 'mode-ai' : mode === 'Mixto' ? 'mode-mixed' : 'mode-unknown';
  return `<span class="mode-badge ${cls}">${esc(mode)}</span>`;
}

function sortedTypeAgg(rows) {
  return aggBy(rows, r => r.type).sort((a,b) => {
    const ma = snapshotMode(a[1]), mb = snapshotMode(b[1]);
    return modeRank(ma) - modeRank(mb) || String(a[0]).localeCompare(String(b[0]), 'es');
  });
}

function aggByService(rows) {
  const svc = {};
  rows.forEach(r => {
    Object.entries(r.services || {}).forEach(([name, s]) => {
      if (!s.cost && !s.input_tokens && !s.output_tokens) return;
      if (!svc[name]) svc[name] = {cost:0, inputTokens:0, outputTokens:0, docs:0};
      svc[name].cost += s.cost || 0;
      svc[name].inputTokens += s.input_tokens || 0;
      svc[name].outputTokens += s.output_tokens || 0;
      svc[name].docs += 1;
    });
  });
  return svc;
}

function durationHistogram(rows) {
  const counts = DURATION_BUCKETS.map(() => 0);
  rows.forEach(r => {
    const s = r.processingSeconds;
    if (s === null || s === undefined) return;
    const i = DURATION_BUCKETS.findIndex(([, lo, hi]) => s >= lo && s < hi);
    if (i >= 0) counts[i] += 1;
  });
  return counts;
}

// Duración por servicio: `durationMs` lo escribe cost_audit.py por servicio,
// así que solo existe para los documentos que sí tienen registro de costo.
function aggServiceDurations(rows) {
  const svc = {};
  rows.forEach(r => Object.entries(r.services || {}).forEach(([name, s]) => {
    if (s.duration_ms === null || s.duration_ms === undefined) return;
    if (!svc[name]) svc[name] = {ms: 0, docs: 0};
    svc[name].ms += s.duration_ms;
    svc[name].docs += 1;
  }));
  return svc;
}

function dailyVolume(rows) {
  const byDay = new Map();
  rows.forEach(r => {
    const day = (r.processedAt || '').slice(0,10);
    if (!day) return;
    if (!byDay.has(day)) byDay.set(day, {day, n:0, cost:0, val:{}});
    const b = byDay.get(day);
    b.n += 1;
    b.cost += r.totalCost || 0;
    if (r.workflowStatus === 'complete' && r.validationStatus) b.val[r.validationStatus] = (b.val[r.validationStatus]||0)+1;
  });
  return [...byDay.values()].sort((a,b) => a.day.localeCompare(b.day));
}

function niceAxisMax(value) {
  if (!value || value <= 0) return 1;
  const magnitude = 10 ** Math.floor(Math.log10(value));
  const normalized = value / magnitude;
  const step = normalized <= 1 ? 1 : normalized <= 2 ? 2 : normalized <= 5 ? 5 : 10;
  return step * magnitude;
}

function axisNumber(value) {
  return Number.isInteger(value) ? num(value) : Number(value).toLocaleString('en-US', {maximumFractionDigits: 2});
}

function worstFields(rows, limit) {
  const agg = {};
  rows.forEach(r => {
    Object.entries(r.fieldStats || {}).forEach(([name, s]) => {
      if (!agg[name]) agg[name] = {ok:0, total:0};
      agg[name].ok += s.ok; agg[name].total += s.total;
    });
  });
  return Object.entries(agg)
    .filter(([,s]) => s.total && s.ok < s.total)
    .sort((a,b) => (a[1].ok/a[1].total) - (b[1].ok/b[1].total))
    .slice(0, limit);
}

// ---------------------------------------------------------------------
// Render: bloques reusables
// ---------------------------------------------------------------------
function kpiTile(value, label, sub) {
  return `<div class="kpi-tile"><div class="kpi-value">${esc(value)}</div><div class="kpi-label">${esc(label)}</div>${sub?`<div class="kpi-sub">${esc(sub)}</div>`:''}</div>`;
}
function chartLegend(items, label='Leyenda de colores') {
  if (!items.length) return '';
  return `<div class="chart-legend" role="list" aria-label="${esc(label)}">${items.map(item =>
    `<span class="legend-item" role="listitem"><span class="legend-swatch" style="background:${item.color}"></span><span>${esc(item.label)}${item.value ? ` · ${esc(item.value)}` : ''}</span></span>`
  ).join('')}</div>`;
}
function hBar(value, maxValue, color, label, valueText, detail) {
  const w = maxValue <= 0 ? 0 : Math.max(2, value/maxValue*100);
  const tip = tooltipText(label, detail || `Valor exacto: ${exactNumber(value)}`, `Valor mostrado: ${valueText}`);
  return `<div class="bar-row chart-point" ${tooltipAttrs(tip)} role="img"><div class="bar-label"><span class="series-label"><span class="series-swatch" style="background:${color}"></span><span class="series-label-text">${esc(label)}</span></span></div><div class="bar-track"><div class="bar-fill" style="width:${w.toFixed(2)}%;background:${color}"></div></div><div class="bar-value">${esc(valueText)}</div></div>`;
}
function stackBar(counts, n, roleMap) {
  if (!n) return '<p class="empty-hint">Sin datos con los filtros actuales.</p>';
  const entries = Object.entries(counts).sort((a,b) => b[1]-a[1]);
  const segs = entries.filter(([,c]) => c>0).map(([status,count]) => {
    const color = STATUS_ROLE_COLOR[roleMap[status] || 'neutral'];
    const w = count/n*100;
    const tip = tooltipText(status, `Valor exacto: ${exactNumber(count)} de ${exactNumber(n)} documentos`, `Participación: ${exactPct(count/n)}`);
    return `<div class="stack-seg chart-point" style="width:${w.toFixed(3)}%;background:${color}" ${tooltipAttrs(tip)} role="img"></div>`;
  }).join('');
  const legend = chartLegend(entries.filter(([,count]) => count > 0).map(([status,count]) => ({
    color: STATUS_ROLE_COLOR[roleMap[status] || 'neutral'],
    label: STATUS_LABEL[status] || status,
    value: `${count} · ${pctFmt(count/n)}`,
  })), 'Significado de los colores');
  return `<div class="stack-bar">${segs}</div>${legend}`;
}

function renderFilters() {
  const mk = (containerId, values, activeSet) => {
    const el = document.getElementById(containerId);
    el.innerHTML = values.map(v => `<span class="chip${activeSet.has(v)?' active':''}" data-value="${esc(v)}">${esc(v)}</span>`).join('');
  };
  const typeCounts = new Map();
  const typeNames = new Map();
  CATALOG_TYPES.forEach(type => typeCounts.set(type, 0));
  ROWS.forEach(r => {
    if (!r.type) return;
    typeCounts.set(r.type, (typeCounts.get(r.type)||0) + 1);
    if (r.typeName) typeNames.set(r.type, r.typeName);
  });
  const typeOptions = [...typeCounts].map(([type,count]) => [type,count,typeNames.get(type)||''])
    .sort((a,b) => b[1]-a[1] || a[0].localeCompare(b[0]));
  const primary = new Set(typeOptions.slice(0, TYPE_FILTER_PRIMARY_LIMIT).map(([type]) => type));
  const query = (document.getElementById('f-type-search').value || '').trim().toLocaleLowerCase('es');
  const matching = query
    ? typeOptions.filter(([type,,name]) => `${type} ${name}`.toLocaleLowerCase('es').includes(query))
    : typeOptions;
  // Un filtro activo nunca desaparece al contraer la lista.
  const visible = query || state.typesExpanded
    ? matching
    : matching.filter(([type]) => primary.has(type) || state.types.has(type));
  const hidden = query ? 0 : typeOptions.filter(([type]) => !visible.some(([shown]) => shown === type)).length;
  const typeEl = document.getElementById('f-types');
  typeEl.innerHTML = visible.map(([type,count,name]) => `<span class="chip${state.types.has(type)?' active':''}" data-value="${esc(type)}" title="${esc(name || type)} · ${count} documento(s)">${esc(type)} · ${num(count)}</span>`).join('')
    + (!visible.length ? '<span class="muted">Sin tipos coincidentes.</span>' : '')
    + (hidden ? `<button type="button" class="chip chip-more" id="f-types-more">+${hidden} ver todos</button>`
      : (!query && typeOptions.length > TYPE_FILTER_PRIMARY_LIMIT ? '<button type="button" class="chip chip-more" id="f-types-more">− ver principales</button>' : ''));
  mk('f-modes', distinct('processingMode'), state.modes);
  mk('f-apps', [...new Set(ROWS.map(appOf))].sort(), state.apps);
  mk('f-models', distinct('model'), state.models);
  mk('f-workflow', distinct('workflowStatus'), state.workflow);
  mk('f-validation', distinct('validationStatus'), state.validation);
  mk('f-traffic', distinct('trafficKind'), state.traffic);
  mk('f-services', Object.keys(SERVICE_LABEL).filter(s => ROWS.some(r => r.services[s])), state.hiddenServices);
  // hiddenServices marca "oculto", no "activo" — invertimos la clase abajo en el listener.
}

function activeFilterCount() {
  return (state.desde ? 1 : 0) + (state.hasta ? 1 : 0)
    + state.types.size + state.modes.size + state.apps.size + state.models.size + state.workflow.size + state.validation.size + state.traffic.size;
}

// Los filtros viven en un panel colapsado: la barra pegajosa de arriba se
// queda como un renglón delgado (botón + resumen) y solo se despliega cuando
// se pide, para no comerse la pantalla en cada scroll.
function setFiltersOpen(open) {
  state.filtersOpen = open;
  const body = document.getElementById('filters-body');
  const btn = document.getElementById('f-toggle');
  body.hidden = !open;
  btn.setAttribute('aria-expanded', String(open));
  btn.querySelector('.ft-caret').textContent = open ? '▾' : '▸';
}

function attachChipHandlers() {
  document.querySelectorAll('#f-types .chip[data-value]').forEach(el => el.addEventListener('click', () => { toggle(state.types, el.dataset.value); render(); }));
  const typeMore = document.getElementById('f-types-more');
  if (typeMore) typeMore.onclick = () => { state.typesExpanded = !state.typesExpanded; render(); };
  document.querySelectorAll('#f-modes .chip').forEach(el => el.addEventListener('click', () => { toggle(state.modes, el.dataset.value); render(); }));
  document.querySelectorAll('#f-apps .chip').forEach(el => el.addEventListener('click', () => { toggle(state.apps, el.dataset.value); render(); }));
  document.querySelectorAll('#f-models .chip').forEach(el => el.addEventListener('click', () => { toggle(state.models, el.dataset.value); render(); }));
  document.querySelectorAll('#f-workflow .chip').forEach(el => el.addEventListener('click', () => { toggle(state.workflow, el.dataset.value); render(); }));
  document.querySelectorAll('#f-validation .chip').forEach(el => el.addEventListener('click', () => { toggle(state.validation, el.dataset.value); render(); }));
  document.querySelectorAll('#f-traffic .chip').forEach(el => el.addEventListener('click', () => { toggle(state.traffic, el.dataset.value); render(); }));
  document.querySelectorAll('#f-services .chip').forEach(el => {
    if (!state.hiddenServices.has(el.dataset.value)) el.classList.add('active');
    el.addEventListener('click', () => { toggle(state.hiddenServices, el.dataset.value); renderServiceSection(filteredRows()); attachServiceChipHandlers(); });
  });
}
function attachServiceChipHandlers() {
  document.querySelectorAll('#f-services .chip').forEach(el => {
    el.classList.toggle('active', !state.hiddenServices.has(el.dataset.value));
    el.onclick = () => { toggle(state.hiddenServices, el.dataset.value); renderServiceSection(filteredRows()); attachServiceChipHandlers(); };
  });
}
function toggle(set, value) { set.has(value) ? set.delete(value) : set.add(value); }

function renderServiceSection(rows) {
  const bySvc = aggByService(rows);
  const visible = Object.entries(bySvc).filter(([name]) => !state.hiddenServices.has(name));
  const maxCost = Math.max(0, ...visible.map(([,a]) => a.cost));
  const bars = visible.sort((a,b) => b[1].cost - a[1].cost)
    .map(([name,a],i) => hBar(a.cost, maxCost, CATEGORICAL[i % CATEGORICAL.length], SERVICE_LABEL[name]||name, money(a.cost),
      tooltipText(`Costo exacto: ${exactMoney(a.cost)}`, `Documentos: ${exactNumber(a.docs)}`, `Tokens entrada: ${exactNumber(a.inputTokens)}`, `Tokens salida: ${exactNumber(a.outputTokens)}`)))
    .join('');
  document.getElementById('by-service-bars').innerHTML = bars || '<p class="empty-hint">Sin datos de costo con los filtros actuales.</p>';
}

function renderAiSection(rows) {
  const a = aggOverall(rows);
  const aiDocsTotal = rows.filter(r => r.processingMode === 'Con IA').reduce((sum,r) => sum + (r.totalCost||0), 0);
  const aiDocsOther = rows.filter(r => r.processingMode === 'Con IA').reduce((sum,r) => sum + (r.nonAiCost||0), 0);
  document.getElementById('ai-kpis').innerHTML = [
    kpiTile(num(a.nNoAi), 'Documentos sin IA', `${a.nAi} con IA · ${a.nAiUnknown} sin dato`),
    kpiTile(money(a.noAiDocsCost), 'Costo de docs sin IA', 'componente Bedrock $0'),
    kpiTile(money(a.aiCost), 'Costo IA', 'sólo fases Bedrock'),
    kpiTile(money(a.nonAiCost), 'Otros servicios', 'Lambda, Textract y runtime'),
  ].join('');
  document.getElementById('ai-table').innerHTML = `<div class="table-wrap"><table>
    <thead><tr><th>Ruta</th><th class="num">Documentos</th><th class="num">Costo de esos docs</th><th class="num">Componente IA</th><th class="num">Resto AWS</th></tr></thead>
    <tbody>
      <tr><td>${modeBadge('Sin IA')}</td><td class="num">${a.nNoAi}<div class="muted">${a.nNoAiWithCost} con costo</div></td><td class="num">${money(a.noAiDocsCost)}</td><td class="num">$0.0000</td><td class="num">${money(a.noAiDocsCost)}</td></tr>
      <tr><td>${modeBadge('Con IA')}</td><td class="num">${a.nAi}</td><td class="num">${money(aiDocsTotal)}</td><td class="num">${money(a.aiCost)}</td><td class="num">${money(aiDocsOther)}</td></tr>
      <tr><td>${modeBadge('Sin datos')}</td><td class="num">${a.nAiUnknown}</td><td class="num">—</td><td class="num">—</td><td class="num">—</td></tr>
    </tbody></table></div>`;
}

function renderInfrastructureSection() {
  const fixed = INFRA.fixedMonthlyUsd;
  const tagged = INFRA.taggedMtd || {};
  const manual = INFRA.manualMtd || {};
  const fixedHistory = INFRA.fixedMonthlyHistory || {};
  document.getElementById('infra-kpis').innerHTML = [
    kpiTile(fixed === null || fixed === undefined ? '—' : money(fixed), 'Piso fijo / mes', `${INFRA.environment || 'ambiente desconocido'} · inventario del stack`),
    kpiTile(INFRA.fixedMtdActualUsd === null || INFRA.fixedMtdActualUsd === undefined ? '—' : money(INFRA.fixedMtdActualUsd), 'Fijo facturado MTD', INFRA.fixedMtdEstimated ? 'Cost Explorer · estimado del mes abierto' : 'Cost Explorer'),
    kpiTile(tagged.available ? money(tagged.totalUsd) : 'No disponible', 'CostStack MTD', tagged.available ? 'etiqueta de asignación' : 'se usa fallback manual'),
    kpiTile(num(INFRA.resourceCount), 'Recursos del stack', `${INFRA.stack || 'sin inventario'} · ${INFRA.region || ''}`),
  ].join('');
  document.getElementById('infra-env-note').innerHTML = INFRA.environmentExplanation
    ? `<p class="infra-note"><strong>Ambiente seleccionado:</strong> ${esc(INFRA.environmentExplanation)}</p>` : '';
  const currentFixed = (fixedHistory.rows || []).find(row => row.status === 'mtd');
  document.getElementById('infra-history-kpis').innerHTML = [
    kpiTile(fixedHistory.actualToDateUsd === null || fixedHistory.actualToDateUsd === undefined ? '—' : money(fixedHistory.actualToDateUsd), 'Fijo real acumulado', `${fixedHistory.start || '—'} a hoy`),
    kpiTile(fixedHistory.projectedPeriodUsd === null || fixedHistory.projectedPeriodUsd === undefined ? '—' : money(fixedHistory.projectedPeriodUsd), 'Proyección fija del periodo', `hasta ${fixedHistory.end || '—'}`),
    kpiTile(!currentFixed || currentFixed.actualUsd === null ? '—' : money(currentFixed.actualUsd), 'Mes actual · gastado MTD', currentFixed ? currentFixed.month : 'fuera del rango'),
    kpiTile(!currentFixed ? '—' : money(currentFixed.estimatedFullMonthUsd), 'Mes actual · cierre estimado', 'run-rate del inventario actual'),
  ].join('');
  const historyRows = fixedHistory.rows || [];
  const historyMax = Math.max(1, ...historyRows.flatMap(row => [row.actualUsd || 0, row.projectedUsd || 0]));
  const monthName = value => new Intl.DateTimeFormat('es-MX', {month:'short', year:'numeric', timeZone:'UTC'}).format(new Date(`${value}-01T00:00:00Z`));
  const statusName = status => status === 'mtd' ? 'Mes actual · MTD' : status === 'future' ? 'Proyección futura' : 'Mes cerrado';
  const monthlyRows = historyRows.map(row => {
    const actualWidth = row.actualUsd === null || row.actualUsd === undefined ? 0 : row.actualUsd/historyMax*100;
    const projectionWidth = (row.projectedUsd || 0)/historyMax*100;
    const actualText = row.actualUsd === null || row.actualUsd === undefined ? '—' : money(row.actualUsd);
    const projectionText = row.projectedUsd === null || row.projectedUsd === undefined ? '—' : money(row.projectedUsd);
    return `<div class="monthly-cost-row${row.status === 'mtd' ? ' current' : ''}">
      <div class="monthly-name">${esc(monthName(row.month))}<span class="monthly-status">${esc(statusName(row.status))}</span></div>
      <div class="monthly-bars" aria-label="${esc(row.month)}: real ${actualText}, usado en proyección ${projectionText}">
        <div class="monthly-track"><div class="monthly-fill" style="width:${actualWidth.toFixed(2)}%;background:${STATUS_ROLE_COLOR.good}"></div></div>
        <div class="monthly-track"><div class="monthly-fill" style="width:${projectionWidth.toFixed(2)}%;background:#4a3aa7"></div></div>
      </div>
      <div class="monthly-values">Real ${actualText} · proyección ${projectionText}</div>
    </div>`;
  }).join('');
  document.getElementById('infra-fixed-history').innerHTML = historyRows.length
    ? chartLegend([
        {color:STATUS_ROLE_COLOR.good, label:'Real facturado (MTD en el mes actual)'},
        {color:'#4a3aa7', label:'Valor usado en la proyección (real cerrado o estimado abierto/futuro)'},
        {color:CATEGORICAL[0], label:'Fondo resaltado = mes en curso'},
      ], 'Significado de los colores del gasto fijo mensual') + `<div class="monthly-cost-chart">${monthlyRows}</div><p class="section-sub">${esc(fixedHistory.scope || '')}</p>`
    : '<p class="empty-hint">Sin histórico mensual de costo fijo disponible.</p>';
  const spendRows = (MONTHLY_SPEND.rows || []).filter(row => row.status !== 'future');
  const currentSpend = spendRows.find(row => row.status === 'mtd');
  document.getElementById('monthly-spend-kpis').innerHTML = [
    kpiTile(money(MONTHLY_SPEND.knownActualUsd), 'Gastado real conocido', 'sin mezclar ninguna proyección'),
    kpiTile(money(MONTHLY_SPEND.processingActualUsd), 'Procesamiento real', `${MONTHLY_SPEND.documentsWithCost || 0} de ${MONTHLY_SPEND.documents || 0} docs con costo`),
    kpiTile(MONTHLY_SPEND.fixedActualUsd === null || MONTHLY_SPEND.fixedActualUsd === undefined ? '—' : money(MONTHLY_SPEND.fixedActualUsd), 'Fijo real', 'facturado y atribuible al stack'),
    kpiTile(!currentSpend || currentSpend.knownActualUsd === null ? '—' : money(currentSpend.knownActualUsd), 'Mes en curso · real MTD', currentSpend ? currentSpend.month : 'fuera del periodo'),
  ].join('');
  const spendMax = Math.max(1, ...spendRows.map(row => row.knownActualUsd || 0));
  const reportSpendRows = spendRows.map(row => {
    const processingWidth = (row.processingActualUsd || 0)/spendMax*100;
    const fixedWidth = (row.fixedActualUsd || 0)/spendMax*100;
    const coverage = `${row.documentsWithCost || 0}/${row.documents || 0} docs con costo`;
    return `<div class="monthly-cost-row${row.status === 'mtd' ? ' current' : ''}">
      <div class="monthly-name">${esc(monthName(row.month))}<span class="monthly-status">${row.status === 'mtd' ? 'Mes actual · MTD' : 'Mes cerrado'}</span></div>
      <div class="monthly-stack-track" aria-label="${esc(row.month)}: procesamiento ${money(row.processingActualUsd)}, fijo ${row.fixedActualUsd === null ? 'no atribuible' : money(row.fixedActualUsd)}">
        <div class="monthly-stack-segment" style="width:${processingWidth.toFixed(2)}%;background:${CATEGORICAL[0]}"></div>
        <div class="monthly-stack-segment" style="width:${fixedWidth.toFixed(2)}%;background:${STATUS_ROLE_COLOR.good}"></div>
      </div>
      <div class="monthly-values">Real ${money(row.knownActualUsd)} · proceso ${money(row.processingActualUsd)} + fijo ${row.fixedActualUsd === null ? '—' : money(row.fixedActualUsd)} · ${coverage}</div>
    </div>`;
  }).join('');
  document.getElementById('monthly-spend-chart').innerHTML = spendRows.length
    ? chartLegend([
        {color:CATEGORICAL[0], label:'Procesamiento real registrado'},
        {color:STATUS_ROLE_COLOR.good, label:'Costo fijo real atribuible'},
        {color:'#4a3aa7', label:'La proyección se muestra aparte; nunca dentro de “real”'},
      ], 'Composición del gasto real') + `<div class="monthly-cost-chart">${reportSpendRows}</div><p class="section-sub">${esc(MONTHLY_SPEND.note || '')}</p>`
    : '<p class="empty-hint">Sin gasto real mensual dentro del periodo.</p>';
  const fixedRows = (INFRA.fixedItems || []).map(item => `<tr><td>${esc(item.service)}</td><td>${esc(item.component)}</td><td class="num">${num(item.quantity)}</td><td class="num">${money(item.unitMonthlyUsd)}</td><td class="num">${money(item.monthlyUsd)}</td></tr>`).join('')
    || '<tr><td colspan="5">Este ambiente no tiene recursos con renta fija.</td></tr>';
  document.getElementById('infra-fixed-table').innerHTML = `<div class="table-wrap"><table><thead><tr><th>Servicio</th><th>Componente</th><th class="num">Cantidad</th><th class="num">Tarifa / mes</th><th class="num">Total / mes</th></tr></thead><tbody>${fixedRows}</tbody></table></div>`;
  const usageRows = (INFRA.usageOnlyItems || []).map(item => `<tr><td>${esc(item.service)}</td><td class="num">${num(item.resources)}</td><td class="num">${money(item.fixedMonthlyUsd)}</td><td>${esc(item.billing)}</td></tr>`).join('');
  document.getElementById('infra-usage-table').innerHTML = usageRows ? `<div class="table-wrap"><table><thead><tr><th>Servicio</th><th class="num">Recursos</th><th class="num">Renta fija</th><th>Qué se cobra</th></tr></thead><tbody>${usageRows}</tbody></table></div>` : '';
  const manualRows = (manual.services || []).map(item => `<tr><td>${esc(item.service)}</td><td class="num">${num(item.resources)}</td><td class="num">${money(item.accountMtdUsd)}</td><td class="num">${item.scopedMtdUsd === null || item.scopedMtdUsd === undefined ? '—' : money(item.scopedMtdUsd)}</td><td>${esc(item.scope)}</td></tr>`).join('');
  document.getElementById('infra-manual-table').innerHTML = manualRows
    ? `<p class="section-sub">Cruce automático entre recursos físicos de CloudFormation y cargos reales por servicio. “Cuenta MTD” sólo da contexto; no se suma al stack si AWS no permite demostrar el recurso que lo originó.</p><div class="table-wrap"><table><thead><tr><th>Servicio presente en el stack</th><th class="num">Recursos</th><th class="num">Cuenta MTD</th><th class="num">Atribuido al stack</th><th>Alcance</th></tr></thead><tbody>${manualRows}</tbody></table></div>`
    : '<p class="empty-hint">No fue posible ejecutar el cruce manual de servicios.</p>';
  const taggedRows = (tagged.services || []).map(item => `<tr><td>${esc(item.service)}</td><td class="num">${money(item.cost)}</td></tr>`).join('');
  document.getElementById('infra-tagged-table').innerHTML = taggedRows ? `<div class="table-wrap"><table><thead><tr><th>Servicio etiquetado</th><th class="num">Costo MTD</th></tr></thead><tbody>${taggedRows}</tbody></table></div>` : '<p class="empty-hint">CostStack no devolvió datos; el reporte continuó con el cruce manual de arriba.</p>';
  document.getElementById('infra-warnings').innerHTML = (INFRA.warnings || []).map(w => `<p class="infra-note">${esc(w)}</p>`).join('');
}

function renderDurationSection(rows) {
  const a = aggOverall(rows);
  document.getElementById('duration-kpis').innerHTML = [
    kpiTile(secsFmt(a.avgDuration), 'Promedio', `${a.nWithDuration} de ${a.n} con dato`),
    kpiTile(secsFmt(a.p50Duration), 'Mediana (p50)'),
    kpiTile(secsFmt(a.p95Duration), 'p95'),
    kpiTile(secsFmt(a.minDuration), 'Mínimo'),
    kpiTile(secsFmt(a.maxDuration), 'Máximo'),
    kpiTile(secsFmt(a.secondsPerPage), 'Por página'),
  ].join('');

  const counts = durationHistogram(rows);
  const maxC = Math.max(0, ...counts);
  document.getElementById('duration-hist').innerHTML = a.nWithDuration
    ? chartLegend([{color:CATEGORICAL[0], label:'Cantidad de documentos en cada rango'}], 'Significado del color de duración')
      + counts.map((c, i) => hBar(c, maxC, CATEGORICAL[0], DURATION_BUCKETS[i][0], `${c} doc${c === 1 ? '' : 's'}`,
      tooltipText(`Documentos: ${exactNumber(c)}`, `Proporción del corte con tiempo: ${exactPct(c/a.nWithDuration)}`))).join('')
    : '<p class="empty-hint">Ningún documento filtrado trae tiempo de proceso registrado (<code>Tracking.ProcessingSeconds</code>).</p>';

  const byType = aggBy(rows, r => r.type).filter(([, m]) => m.nWithDuration).sort((x, y) => y[1].avgDuration - x[1].avgDuration);
  document.getElementById('duration-by-type').innerHTML = byType.length ? `<div class="table-wrap"><table>
    <thead><tr><th>Tipo</th><th class="num">Docs con dato</th><th class="num">Promedio</th><th class="num">p50</th><th class="num">p95</th><th class="num">Máximo</th><th class="num">Por página</th></tr></thead>
    <tbody>${byType.map(([t, m]) => `<tr><td>${esc(t)}</td><td class="num">${m.nWithDuration}</td><td class="num">${secsFmt(m.avgDuration)}</td><td class="num">${secsFmt(m.p50Duration)}</td><td class="num">${secsFmt(m.p95Duration)}</td><td class="num">${secsFmt(m.maxDuration)}</td><td class="num">${secsFmt(m.secondsPerPage)}</td></tr>`).join('')}</tbody>
  </table></div>` : '';

  const byApp = aggBy(rows, appOf).filter(([, m]) => m.nWithDuration).sort((x, y) => y[1].avgDuration - x[1].avgDuration);
  document.getElementById('duration-by-app').innerHTML = byApp.length ? `<div class="table-wrap"><table>
    <thead><tr><th>Sistema</th><th class="num">Docs con dato</th><th class="num">Promedio</th><th class="num">p50</th><th class="num">p95</th><th class="num">Máximo</th></tr></thead>
    <tbody>${byApp.map(([t, m]) => `<tr><td>${esc(t)}</td><td class="num">${m.nWithDuration}</td><td class="num">${secsFmt(m.avgDuration)}</td><td class="num">${secsFmt(m.p50Duration)}</td><td class="num">${secsFmt(m.p95Duration)}</td><td class="num">${secsFmt(m.maxDuration)}</td></tr>`).join('')}</tbody>
  </table></div>` : '';

  const svcDur = Object.entries(aggServiceDurations(rows)).map(([n, s]) => [n, s.ms / s.docs / 1000, s.docs])
    .sort((x, y) => y[1] - x[1]);
  const maxSvc = Math.max(0, ...svcDur.map(e => e[1]));
  document.getElementById('duration-by-service').innerHTML = svcDur.length
    ? svcDur.map(([n, v, docs], i) => hBar(v, maxSvc, CATEGORICAL[i % CATEGORICAL.length], `${SERVICE_LABEL[n] || n} (${docs} docs)`, secsFmt(v),
      tooltipText(`Promedio exacto: ${exactSeconds(v)}`, `Documentos con dato: ${exactNumber(docs)}`))).join('')
    : '<p class="empty-hint">Los registros de costo filtrados no traen duración por servicio.</p>';

  const slow = rows.filter(r => r.processingSeconds !== null && r.processingSeconds !== undefined)
    .sort((x, y) => y.processingSeconds - x.processingSeconds).slice(0, 10);
  document.getElementById('duration-slowest').innerHTML = slow.length ? `<div class="table-wrap"><table>
    <thead><tr><th>Documento</th><th>Tipo</th><th>Sistema</th><th class="num">Páginas</th><th class="num">Costo</th><th class="num">Tiempo</th></tr></thead>
    <tbody>${slow.map(r => `<tr class="doc-row" data-id="${esc(r.id)}"><td class="mono">${esc(r.id)}</td><td>${esc(r.typeName || r.type)}</td><td class="muted">${esc(appOf(r))}</td><td class="num">${r.pagesTotal || '—'}</td><td class="num">${r.hasCost ? money(r.totalCost) : '—'}</td><td class="num">${secsFmt(r.processingSeconds)}</td></tr>`).join('')}</tbody>
  </table></div>` : '';
  document.querySelectorAll('#duration-slowest tr.doc-row').forEach(tr => tr.addEventListener('click', () => openDetail(tr.dataset.id)));
}

// Contraste contra las mediciones reales de documentacion/COSTOS.md: el corte
// filtrado vs lo que se midió/presupuestó, para ver de un vistazo si el
// sistema sigue costando lo que dice el documento.
function deltaCell(actual, ref) {
  if (actual === null || actual === undefined || ref === null || ref === undefined) return '<td class="num">—</td>';
  if (!ref) return `<td class="num">${actual > 0 ? 'ref. $0' : '='}</td>`;
  const d = (actual - ref) / ref;
  const cls = d <= 0.05 ? 'delta-good' : (d <= 0.30 ? 'delta-warn' : 'delta-bad');
  return `<td class="num ${cls}">${d >= 0 ? '+' : ''}${(d * 100).toFixed(0)}%</td>`;
}

function renderReferenceSection(rows) {
  const a = aggOverall(rows);
  const budget = COST_REF.perPage.find(x => x.role === 'budget');
  const projectedProcessing = a.avgCostPage === null ? null : a.avgCostPage * COST_REF.volume.pagesPerMonth;
  const fixedMonthly = INFRA.fixedMonthlyUsd;
  const projected = projectedProcessing === null ? null : projectedProcessing + (fixedMonthly || 0);
  document.getElementById('ref-kpis').innerHTML = [
    kpiTile(a.avgCostPage === null ? '—' : money5(a.avgCostPage), 'Costo / página (corte actual)', `${a.nWithCost} doc(s) con costo · ${num(a.totalPages)} pág`),
    kpiTile(money(a.avgCostDoc), 'Costo / documento (corte actual)', `referencia ${money(COST_REF.perDoc.value)} — ${COST_REF.perDoc.label}`),
    kpiTile(money5(budget.value), 'Presupuesto / página', `${COST_REF.doc} ${budget.source}`),
    kpiTile(projected === null ? '—' : '$' + projected.toFixed(2), `Proyección / mes + fijo`, `procesamiento ${projectedProcessing===null?'—':'$'+projectedProcessing.toFixed(2)} + fijo ${fixedMonthly===null||fixedMonthly===undefined?'—':money(fixedMonthly)}`),
  ].join('');

  const maxRef = Math.max(a.avgCostPage || 0, ...COST_REF.perPage.map(x => x.value));
  const line = (label, value, color, sub) => {
    const w = maxRef <= 0 ? 0 : Math.max(1, value / maxRef * 100);
    const tip = tooltipText(label, `Valor exacto: ${exactMoney(value)} por página`, sub ? `Fuente: ${sub}` : 'Datos del corte filtrado');
    return `<div class="ref-line chart-point" ${tooltipAttrs(tip)} role="img"><div class="ref-label"><span class="series-label"><span class="series-swatch" style="background:${color}"></span><span class="series-label-text">${esc(label)}${sub ? ` <span class="muted">${esc(sub)}</span>` : ''}</span></span></div>
      <div class="ref-track"><div class="ref-fill" style="width:${w.toFixed(2)}%;background:${color}"></div></div>
      <div class="ref-value">${money5(value)}</div></div>`;
  };
  const rolesColor = {budget: STATUS_ROLE_COLOR.warning, cushion: STATUS_ROLE_COLOR.neutral, expected: '#4a3aa7', measured: STATUS_ROLE_COLOR.good};
  document.getElementById('ref-bars').innerHTML = chartLegend([
    {color:CATEGORICAL[0], label:'Costo del corte actual'},
    {color:rolesColor.budget, label:'Presupuesto'},
    {color:rolesColor.cushion, label:'Colchón presupuestario'},
    {color:rolesColor.expected, label:'Costo esperado'},
    {color:rolesColor.measured, label:'Costo medido históricamente'},
  ], 'Significado de los colores de referencia') + [
    a.avgCostPage === null ? '' : line('Este corte', a.avgCostPage, CATEGORICAL[0]),
    ...COST_REF.perPage.map(x => line(x.label, x.value, rolesColor[x.role], x.source)),
  ].join('');
  document.getElementById('ref-delta-legend').innerHTML = chartLegend([
    {color:STATUS_ROLE_COLOR.good, label:'Dentro de referencia (≤ 5 %)'},
    {color:STATUS_ROLE_COLOR.warning, label:'Atención (≤ 30 % arriba)'},
    {color:STATUS_ROLE_COLOR.critical, label:'Por encima (> 30 %)'},
    {color:STATUS_ROLE_COLOR.neutral, label:'Sin referencia comparable'},
  ], 'Significado de los colores de variación');

  const byType = sortedTypeAgg(rows);
  const refRows = byType.map(([t, m]) => {
    const ref = COST_REF.perDocByType[t];
    const actualPage = m.totalPages ? m.totalCost / m.totalPages : null;
    return `<tr>
      <td>${esc(t)}</td><td>${modeBadge(snapshotMode(m))}</td>
      <td class="num">${m.n}</td>
      <td class="num">${money(m.avgCostDoc)}</td>
      <td class="num muted">${ref ? money(ref.costPerDoc) : '—'}</td>
      ${deltaCell(m.avgCostDoc, ref ? ref.costPerDoc : null)}
      <td class="num">${actualPage === null ? '—' : money5(actualPage)}</td>
      <td class="num muted">${ref ? money5(ref.costPerPage) : '—'}</td>
      ${deltaCell(actualPage, ref ? ref.costPerPage : null)}
    </tr>`;
  }).join('');
  document.getElementById('ref-by-type').innerHTML = byType.length ? `<div class="table-wrap"><table>
    <thead><tr><th>Tipo</th><th>Ruta</th><th class="num">Docs</th><th class="num">$/doc hoy</th><th class="num">$/doc ref.</th><th class="num">Δ</th><th class="num">$/pág hoy</th><th class="num">$/pág ref.</th><th class="num">Δ</th></tr></thead>
    <tbody>${refRows}</tbody></table></div>` : '<p class="empty-hint">Sin datos con los filtros actuales.</p>';
}

function renderDocumentsTable(rows) {
  const sorted = [...rows].sort((a,b) => {
    const ka = a[state.sortKey], kb = b[state.sortKey];
    if (ka === null || ka === undefined) return 1;
    if (kb === null || kb === undefined) return -1;
    if (ka < kb) return -1 * state.sortDir;
    if (ka > kb) return 1 * state.sortDir;
    return 0;
  });
  const shown = sorted.slice(0, state.docLimit);
  document.getElementById('doc-count').textContent = rows.length;
  const cols = [
    ['id','Documento'], ['typeName','Tipo'], ['processingMode','Procesamiento'], ['appOrigin','Sistema'], ['workflowStatus','Ejecución'], ['validationStatus','Negocio'],
    ['model','Modelo'], ['avgConfidence','Confianza'], ['processingSeconds','Tiempo'], ['totalCost','Costo'], ['processedAt','Procesado'],
  ];
  const head = cols.map(([key,label]) => {
    const arrow = state.sortKey === key ? (state.sortDir === 1 ? '↑' : '↓') : '';
    return `<th class="sortable" data-key="${key}">${esc(label)} <span class="arrow">${arrow}</span></th>`;
  }).join('');
  const body = shown.map(r => `<tr class="doc-row" data-id="${esc(r.id)}">
      <td class="mono">${esc(r.id)}</td>
      <td>${esc(r.typeName || r.type)}</td>
      <td>${modeBadge(r.processingMode)}</td>
      <td class="muted">${esc(r.appOrigin || '—')}</td>
      <td>${esc(r.workflowStatus || '—')}</td>
      <td>${esc(r.validationStatus || '—')}</td>
      <td class="muted">${esc(r.model)}</td>
      <td class="num">${pctFmt(r.avgConfidence)}</td>
      <td class="num">${secsFmt(r.processingSeconds)}</td>
      <td class="num">${r.hasCost ? money(r.totalCost) : '—'}</td>
      <td class="muted">${esc((r.processedAt||'').slice(0,19))}</td>
    </tr>`).join('');
  document.getElementById('documents-table').innerHTML =
    `<table><thead><tr>${head}</tr></thead><tbody>${body}</tbody></table>`;
  const more = document.getElementById('documents-more');
  more.innerHTML = rows.length > state.docLimit
    ? `<button class="more-btn" id="show-more">Mostrar ${Math.min(200, rows.length-state.docLimit)} más (${rows.length-state.docLimit} restantes)</button>`
    : '';
  const btn = document.getElementById('show-more');
  if (btn) btn.addEventListener('click', () => { state.docLimit += 200; renderDocumentsTable(filteredRows()); });
  document.querySelectorAll('#documents-table th.sortable').forEach(th => th.addEventListener('click', () => {
    const key = th.dataset.key;
    state.sortDir = (state.sortKey === key) ? -state.sortDir : -1;
    state.sortKey = key;
    renderDocumentsTable(filteredRows());
  }));
  document.querySelectorAll('#documents-table tbody tr.doc-row').forEach(tr => tr.addEventListener('click', () => openDetail(tr.dataset.id)));
}

function openDetail(id) {
  const r = ROWS_BY_ID.get(id);
  if (!r) return;
  const fieldRows = Object.entries(r.fieldStats||{}).map(([name,s]) =>
    `<tr><td>${esc(name)}</td><td class="num">${s.ok}/${s.total}</td></tr>`).join('');
  const svcRows = Object.entries(r.services||{}).filter(([,s]) => s.cost || s.input_tokens || s.output_tokens || s.pages || s.duration_ms).map(([name,s]) =>
    `<tr><td>${esc(SERVICE_LABEL[name]||name)}</td><td class="num">${money(s.cost)}</td><td class="num">${num(s.input_tokens)}</td><td class="num">${num(s.output_tokens)}</td><td class="num">${s.duration_ms===null||s.duration_ms===undefined?'—':secsFmt(s.duration_ms/1000)}</td><td class="muted">${esc(s.model_id||'—')}</td></tr>`).join('');
  const ref = COST_REF.perDocByType[r.type];
  const warnings = (r.warnings||[]).length ? `<ul class="warn-list">${r.warnings.map(w=>`<li>${esc(w)}</li>`).join('')}</ul>` : '<p class="empty-hint">Sin advertencias.</p>';

  document.getElementById('detail-content').innerHTML = `
    <button class="detail-close" id="detail-close-btn">✕</button>
    <h2 style="margin:0 0 4px">${esc(r.typeName || r.type)}</h2>
    <p class="section-sub" style="margin-bottom:14px">${esc(r.id)} · ${esc(r.name || 'sin nombre de archivo')}</p>

    <h3>Calidad</h3>
    ${r.hasQuality ? `<div class="kv-grid">
      <div class="k">Estado de ejecución</div><div>${esc(r.workflowStatus||'—')}</div>
      <div class="k">Estado de negocio</div><div>${esc(r.validationStatus||'—')}</div>
      <div class="k">Sistema (app de origen)</div><div>${esc(appOf(r))}</div>
      <div class="k">Modelo</div><div>${esc(r.model)}</div>
      <div class="k">Ruta de procesamiento</div><div>${modeBadge(r.processingMode)}</div>
      <div class="k">Evidencia de ruta</div><div>${esc(r.processingEvidence||'—')}</div>
      <div class="k">Confianza promedio</div><div>${pctFmt(r.avgConfidence)}</div>
      <div class="k">Tiempo de proceso</div><div>${secsFmt(r.processingSeconds)}</div>
      <div class="k">Documentos en el expediente</div><div>${r.subDocCount}</div>
      <div class="k">Procesado</div><div>${esc((r.processedAt||'').slice(0,19))}</div>
    </div>
    <h3>Advertencias</h3>${warnings}
    ${fieldRows ? `<h3>Campos (cobertura parcial en expedientes multi-documento)</h3>
      <div class="table-wrap"><table><thead><tr><th>Campo</th><th class="num">Válidos</th></tr></thead><tbody>${fieldRows}</tbody></table></div>` : ''}
    ` : '<p class="empty-hint">Sin datos de calidad para este documento (no está en la tabla de tracking).</p>'}

    <h3>Costo</h3>
    ${r.hasCost ? `<div class="kv-grid">
      <div class="k">Costo total</div><div>${money(r.totalCost)}</div>
      <div class="k">Componente IA</div><div>${money(r.aiCost)}</div>
      <div class="k">Componente no IA</div><div>${money(r.nonAiCost)}</div>
      <div class="k">Costo por página</div><div>${r.costPerPage!==null&&r.costPerPage!==undefined?money(r.costPerPage):'—'}</div>
      <div class="k">Páginas</div><div>${r.pagesTotal} (${r.pagesNative} nativas, ${r.pagesOcr} OCR, ${r.pagesEmpty} vacías)</div>
      <div class="k">Tokens de entrada</div><div>${num(r.totalInputTokens)}</div>
      <div class="k">Tokens de salida</div><div>${num(r.totalOutputTokens)}</div>
      <div class="k">Versión de tarifa</div><div>${esc(r.pricingVersion||'—')}</div>
      <div class="k">Referencia medida (${esc(COST_REF.doc)})</div><div>${ref ? `${money(ref.costPerDoc)}/doc · ${money5(ref.costPerPage)}/pág${ref.note ? ` <span class="muted">(${esc(ref.note)})</span>` : ''}` : '<span class="muted">sin referencia para este tipo</span>'}</div>
    </div>
    <div class="table-wrap"><table><thead><tr><th>Servicio</th><th class="num">Costo</th><th class="num">Tokens in</th><th class="num">Tokens out</th><th class="num">Duración</th><th>Modelo</th></tr></thead><tbody>${svcRows}</tbody></table></div>
    ` : '<p class="empty-hint">Sin datos de costo para este documento (tabla de costos vacía, o el registro no se llegó a escribir).</p>'}
  `;
  document.getElementById('detail-close-btn').addEventListener('click', () => document.getElementById('detail-dialog').close());
  document.getElementById('detail-dialog').showModal();
}

// Un solo tooltip flotante evita depender del `title` nativo (que tarda en
// aparecer, no admite varias líneas de forma consistente y no se puede leer
// con teclado). Los elementos dinámicos solo declaran `data-tooltip`.
const REPORT_TOOLTIP = document.getElementById('report-tooltip');
let activeTooltipTarget = null;

function placeTooltip(clientX, clientY) {
  const gap = 14, edge = 12;
  REPORT_TOOLTIP.style.left = `${clientX + gap}px`;
  REPORT_TOOLTIP.style.top = `${clientY + gap}px`;
  const box = REPORT_TOOLTIP.getBoundingClientRect();
  let left = clientX + gap, top = clientY + gap;
  if (left + box.width > window.innerWidth - edge) left = clientX - box.width - gap;
  if (top + box.height > window.innerHeight - edge) top = clientY - box.height - gap;
  REPORT_TOOLTIP.style.left = `${Math.max(edge, left)}px`;
  REPORT_TOOLTIP.style.top = `${Math.max(edge, top)}px`;
}

function showTooltip(target, event) {
  const detail = target && target.dataset ? target.dataset.tooltip : '';
  if (!detail) return;
  activeTooltipTarget = target;
  REPORT_TOOLTIP.textContent = detail;
  REPORT_TOOLTIP.classList.add('visible');
  REPORT_TOOLTIP.setAttribute('aria-hidden', 'false');
  if (event && Number.isFinite(event.clientX)) {
    placeTooltip(event.clientX, event.clientY);
  } else {
    const rect = target.getBoundingClientRect();
    placeTooltip(rect.left + rect.width/2, rect.bottom);
  }
}

function hideTooltip() {
  activeTooltipTarget = null;
  REPORT_TOOLTIP.classList.remove('visible');
  REPORT_TOOLTIP.setAttribute('aria-hidden', 'true');
}

document.addEventListener('pointerover', event => {
  const target = event.target.closest?.('.chart-point[data-tooltip]');
  if (target && target !== activeTooltipTarget) showTooltip(target, event);
});
document.addEventListener('pointermove', event => {
  const target = event.target.closest?.('.chart-point[data-tooltip]');
  if (target && target !== activeTooltipTarget) showTooltip(target, event);
  else if (activeTooltipTarget && activeTooltipTarget.matches(':hover')) placeTooltip(event.clientX, event.clientY);
});
document.addEventListener('pointerout', event => {
  const target = event.target.closest?.('.chart-point[data-tooltip]');
  if (target === activeTooltipTarget && (!event.relatedTarget || !target.contains(event.relatedTarget))) hideTooltip();
});
document.addEventListener('focusin', event => {
  const target = event.target.closest?.('.chart-point[data-tooltip]');
  if (target) showTooltip(target);
});
document.addEventListener('focusout', event => {
  if (event.target === activeTooltipTarget) hideTooltip();
});
document.addEventListener('keydown', event => { if (event.key === 'Escape') hideTooltip(); });
window.addEventListener('scroll', hideTooltip, true);
window.addEventListener('resize', hideTooltip);

function render() {
  hideTooltip();
  const rows = filteredRows();
  renderFilters();
  attachChipHandlers();

  const nActive = activeFilterCount();
  const badge = document.getElementById('f-badge');
  badge.hidden = nActive === 0;
  badge.textContent = nActive;
  document.getElementById('f-summary').textContent = nActive
    ? `Mostrando ${rows.length} de ${ROWS.length} documento(s) · ${nActive} filtro(s) activo(s)`
    : `Mostrando los ${ROWS.length} documento(s) — sin filtros`;

  const overall = aggOverall(rows);
  document.getElementById('kpis').innerHTML = [
    kpiTile(num(overall.n), 'Documentos', 'todos los procesados en el periodo'),
    kpiTile(num(overall.totalPages), 'Páginas procesadas'),
    kpiTile(num(overall.nCompleted), 'Ejecuciones completas'),
    kpiTile(pctFmt(overall.approvalRate), 'Tasa de aprobación'),
    kpiTile(pctFmt(overall.avgConfidence), 'Confianza promedio'),
    kpiTile(secsFmt(overall.avgDuration), 'Tiempo promedio', overall.nWithDuration ? `p95 ${secsFmt(overall.p95Duration)}` : 'sin dato de tiempo'),
    kpiTile(money(overall.totalCost), 'Costo de procesamiento', overall.nWithCost < overall.n ? `${overall.nWithCost} de ${overall.n} con dato de costo` : 'atribuible por documento'),
    kpiTile(num(overall.nNoAi), 'Documentos sin IA', `${money(overall.noAiDocsCost)} en total`),
    kpiTile(money(overall.avgCostDoc), 'Costo promedio / doc', `referencia ${money(COST_REF.perDoc.value)}`),
    kpiTile(overall.avgCostPage === null ? '—' : money5(overall.avgCostPage), 'Costo / página', `presupuesto ${money5(0.0083)}`),
    kpiTile(num(overall.totalInputTokens), 'Tokens de entrada'),
    kpiTile(num(overall.totalOutputTokens), 'Tokens de salida'),
  ].join('');

  renderPageSel(overall.pageSel, overall.n);

  document.getElementById('wf-stack').innerHTML = stackBar(overall.wfCounts, overall.n, WORKFLOW_ROLE);
  document.getElementById('validation-stack').innerHTML = stackBar(overall.valCounts, overall.nCompleted, VALIDATION_ROLE);

  const volume = dailyVolume(rows);
  if (volume.length) {
    const maxN = Math.max(...volume.map(v => v.n));
    const axisMax = niceAxisMax(maxN);
    const yTicks = [axisMax, axisMax * .75, axisMax * .5, axisMax * .25, 0];
    const bars = volume.map(v => {
      const statuses = Object.entries(v.val).sort((a,b)=>b[1]-a[1]).map(([s,n]) => `${s}: ${exactNumber(n)}`);
      const tip = tooltipText(v.day, `Documentos: ${exactNumber(v.n)}`, `Costo exacto: ${exactMoney(v.cost)}`, statuses.length ? ['Resultado de negocio:', ...statuses] : 'Sin resultados de negocio completos');
      return `<div class="vol-bar chart-point" style="height:${Math.max(2, v.n/axisMax*100).toFixed(1)}%;background:${CATEGORICAL[0]}" ${tooltipAttrs(tip)} role="img"></div>`;
    }).join('');
    document.getElementById('volume-chart').innerHTML =
      chartLegend([{color:CATEGORICAL[0], label:'Altura = documentos procesados en el día'}], 'Significado del color de volumen') +
      `<div class="volume-plot" role="img" aria-label="Volumen diario de documentos, máximo ${axisNumber(maxN)} documentos">
        <div class="volume-y" aria-hidden="true">${yTicks.map(v => `<span>${axisNumber(v)}</span>`).join('')}</div>
        <div class="volume-main"><div class="volume-plot-area"><div class="volume-bars">${bars}</div></div>
          <div class="vol-labels"><span>${esc(volume[0].day)}</span><span>${esc(volume[volume.length-1].day)}</span></div>
          <div class="volume-axis-note">Documentos procesados por día</div>
        </div>
      </div>`;
  } else {
    document.getElementById('volume-chart').innerHTML = '<p class="empty-hint">Sin datos con los filtros actuales.</p>';
  }

  const byType = sortedTypeAgg(rows);
  document.getElementById('by-type-table').innerHTML = byType.length ? `<div class="table-wrap"><table>
    <thead><tr><th>Tipo</th><th>Ruta</th><th class="num">Docs</th><th class="num">Con IA</th><th class="num">Aprobación</th><th class="num">Tiempo prom.</th><th class="num">Costo total</th><th class="num">Costo IA</th><th class="num">$/doc</th></tr></thead>
    <tbody>${byType.map(([t,m]) => `<tr><td>${esc(t)}</td><td>${modeBadge(snapshotMode(m))}</td><td class="num">${m.n}</td><td class="num">${m.nAi}</td><td class="num">${pctFmt(m.approvalRate)}</td><td class="num">${secsFmt(m.avgDuration)}</td><td class="num">${money(m.totalCost)}</td><td class="num">${money(m.aiCost)}</td><td class="num">${money(m.avgCostDoc)}</td></tr>`).join('')}</tbody>
  </table></div>` : '<p class="empty-hint">Sin datos con los filtros actuales.</p>';

  const byApp = aggBy(rows, appOf).sort((a,b) => b[1].n - a[1].n);
  const maxAppN = Math.max(0, ...byApp.map(([,m]) => m.n));
  document.getElementById('by-app-bars').innerHTML = byApp.map(([name,m],i) =>
    hBar(m.n, maxAppN, CATEGORICAL[i % CATEGORICAL.length], `${name} (${m.n} docs)`, money(m.totalCost),
      tooltipText(`Documentos: ${exactNumber(m.n)}`, `Costo exacto: ${exactMoney(m.totalCost)}`, `Aprobación: ${exactPct(m.approvalRate)}`, `Tiempo promedio: ${exactSeconds(m.avgDuration)}`))).join('');
  document.getElementById('by-app-table').innerHTML = byApp.length ? `<div class="table-wrap"><table>
    <thead><tr><th>Sistema</th><th class="num">Docs</th><th class="num">Aprobación</th><th class="num">Confianza</th><th class="num">Tiempo prom.</th><th class="num">p95</th><th class="num">Costo total</th><th class="num">$/doc</th></tr></thead>
    <tbody>${byApp.map(([t,m]) => `<tr><td>${esc(t)}</td><td class="num">${m.n}</td><td class="num">${pctFmt(m.approvalRate)}</td><td class="num">${pctFmt(m.avgConfidence)}</td><td class="num">${secsFmt(m.avgDuration)}</td><td class="num">${secsFmt(m.p95Duration)}</td><td class="num">${money(m.totalCost)}</td><td class="num">${money(m.avgCostDoc)}</td></tr>`).join('')}</tbody>
  </table></div>` : '<p class="empty-hint">Sin datos con los filtros actuales.</p>';

  const byModel = aggBy(rows, r => r.model).sort((a,b) => b[1].n - a[1].n);
  const maxModelN = Math.max(0, ...byModel.map(([,m]) => m.n));
  document.getElementById('by-model-bars').innerHTML = byModel.map(([name,m],i) =>
    hBar(m.n, maxModelN, CATEGORICAL[i % CATEGORICAL.length], `${name} (${m.n} docs)`, pctFmt(m.approvalRate),
      tooltipText(`Documentos: ${exactNumber(m.n)}`, `Aprobación exacta: ${exactPct(m.approvalRate)}`, `Confianza promedio: ${exactPct(m.avgConfidence)}`, `Costo exacto: ${exactMoney(m.totalCost)}`))).join('');
  document.getElementById('by-model-table').innerHTML = byModel.length ? `<div class="table-wrap"><table>
    <thead><tr><th>Modelo</th><th class="num">Docs</th><th class="num">Aprobación</th><th class="num">Confianza</th><th class="num">Tiempo prom.</th><th class="num">Costo total</th></tr></thead>
    <tbody>${byModel.map(([name,m]) => `<tr><td>${esc(name)}</td><td class="num">${m.n}</td><td class="num">${pctFmt(m.approvalRate)}</td><td class="num">${pctFmt(m.avgConfidence)}</td><td class="num">${secsFmt(m.avgDuration)}</td><td class="num">${money(m.totalCost)}</td></tr>`).join('')}</tbody>
  </table></div>` : '';

  renderDurationSection(rows);
  renderAiSection(rows);
  renderReferenceSection(rows);
  renderServiceSection(rows);
  renderInfrastructureSection();
  attachServiceChipHandlers();

  const wf = worstFields(rows, 25);
  document.getElementById('worst-fields-table').innerHTML = wf.length ? `<div class="table-wrap"><table>
    <thead><tr><th>Campo</th><th class="num">Válidos</th><th class="num">Tasa</th></tr></thead>
    <tbody>${wf.map(([name,s]) => `<tr><td>${esc(name)}</td><td class="num">${s.ok}/${s.total}</td><td class="num">${(s.ok/s.total*100).toFixed(0)}%</td></tr>`).join('')}</tbody>
  </table></div>` : '<p class="empty-hint">Sin campos con fallas registradas en los documentos filtrados.</p>';

  renderDocumentsTable(rows);
}

document.getElementById('f-desde').addEventListener('change', e => { state.desde = e.target.value; render(); });
document.getElementById('f-hasta').addEventListener('change', e => { state.hasta = e.target.value; render(); });
document.getElementById('f-type-search').addEventListener('input', () => { renderFilters(); attachChipHandlers(); });
document.getElementById('f-toggle').addEventListener('click', () => setFiltersOpen(!state.filtersOpen));
document.getElementById('f-reset').addEventListener('click', () => {
  state.desde = DEFAULT_RANGE.start || ''; state.hasta = DEFAULT_RANGE.end || '';
  state.typesExpanded = false;
  state.types.clear(); state.modes.clear(); state.apps.clear(); state.models.clear(); state.workflow.clear(); state.validation.clear(); state.traffic.clear();
  document.getElementById('f-desde').value = state.desde; document.getElementById('f-hasta').value = state.hasta; document.getElementById('f-type-search').value = '';
  render();
});

setFiltersOpen(state.filtersOpen);
document.getElementById('f-desde').value = state.desde;
document.getElementById('f-hasta').value = state.hasta;
render();
'''


def _json_default(o: Any) -> Any:
    """Safety net for the browser-embedded JSON: `_report_lib.normalize_*`
    already casts every known DynamoDB numeric field away from Decimal, but a
    future field added there (or upstream in cost_audit.py/indexing.py)
    could reintroduce one — better a correct float here than a crashed
    report generation."""
    if isinstance(o, Decimal):
        return int(o) if o == o.to_integral_value() else float(o)
    raise TypeError(f"Object of type {o.__class__.__name__} is not JSON serializable")


def render_quality_html(quality: Dict[str, Any]) -> str:
    """Bloque ejecutivo de precisión real; nunca imprime expected/actual."""
    latest_free = quality.get("latestFree") or {}
    latest_real = quality.get("latestReal") or {}
    free = latest_free.get("summary") or {}
    real = latest_real.get("summary") or {}
    def percentage(value: Any) -> str:
        return "—" if value is None else f"{float(value):.1%}"

    def coverage(value: Any) -> str:
        return "—" if value is None else f"{float(value):.1f}%"
    insights = []
    if not latest_real:
        insights.append("Aún no existe una corrida E2E completa; aprobación y confianza no demuestran exactitud.")
    elif latest_real.get("status") != "PASSED":
        insights.append("La última validación E2E falló: revisar los casos antes de promover cambios.")
    if real.get("fieldAccuracy") is not None and real["fieldAccuracy"] < 1:
        insights.append(f"La exactitud de campos etiquetados es {real['fieldAccuracy']:.1%}; hay una regresión medible contra ground truth.")
    if free.get("lineCoverage") is not None and free["lineCoverage"] < 80:
        insights.append(f"La cobertura Python está en {free['lineCoverage']:.1f}%; conviene priorizar rutas críticas sin cobertura.")
    for regression in quality.get("regressions") or []:
        insights.append(
            f"La {regression['label']} de la última suite subió {regression['delta']:.0%} "
            f"contra la mediana de corridas comparables ({regression['severity']})."
        )
    if not insights:
        insights.append("La última evidencia gratuita y E2E no muestra regresiones de exactitud.")
    trend_rows = "".join(
        f"<tr><td>{esc(r.get('at', '')[:19])}</td><td>{esc(r.get('suite'))}</td>"
        f"<td>{esc(r.get('status'))}</td><td class='num'>{r.get('tests', 0)}</td>"
        f"<td class='num'>{r.get('failures', 0)}</td>"
        f"<td class='num'>{percentage(r.get('fieldAccuracy'))}</td></tr>"
        for r in reversed((quality.get("trend") or [])[-12:])
    ) or "<tr><td colspan='6'>Sin corridas registradas.</td></tr>"
    insight_html = "".join(f"<li>{esc(item)}</li>" for item in insights)
    return f'''<details class="card quality-card" open>
        <summary>Calidad comprobada e histórico de pruebas</summary>
        <p class="section-sub">Exactitud contra valores humanos, separada de aprobación de reglas y confianza del modelo.</p>
        <div class="kpi-row dense">
          <div class="kpi-tile"><div class="kpi-value">{quality.get('runs', 0)}</div><div class="kpi-label">Corridas conservadas</div></div>
          <div class="kpi-tile"><div class="kpi-value">{esc(latest_free.get('status', '—'))}</div><div class="kpi-label">Suite gratuita</div><div class="kpi-sub">{free.get('tests', 0)} pruebas · {free.get('failures', 0)} fallas</div></div>
          <div class="kpi-tile"><div class="kpi-value">{coverage(free.get('lineCoverage'))} / {coverage(free.get('nodeLineCoverage'))}</div><div class="kpi-label">Cobertura Python / Node</div></div>
          <div class="kpi-tile"><div class="kpi-value">{percentage(real.get('fieldAccuracy'))}</div><div class="kpi-label">Exactitud de campos</div><div class="kpi-sub">sólo campos etiquetados</div></div>
          <div class="kpi-tile"><div class="kpi-value">{percentage(real.get('e2ePassRate'))}</div><div class="kpi-label">Éxito E2E</div></div>
        </div>
        <div class="sub-head">Insights</div><ul class="insight-list">{insight_html}</ul>
        <div class="sub-head">Últimas corridas</div><div class="table-wrap"><table>
          <thead><tr><th>Fecha</th><th>Suite</th><th>Estado</th><th class="num">Pruebas</th><th class="num">Fallas</th><th class="num">Exactitud</th></tr></thead>
          <tbody>{trend_rows}</tbody></table></div>
      </details>'''


def render_html(
    rows: List[Dict[str, Any]], env: str, quality: Optional[Dict[str, Any]] = None,
    infrastructure: Optional[Dict[str, Any]] = None,
    catalog_types: Optional[List[str]] = None,
) -> str:
    from _report_lib import BASE_CSS

    overall = overall_snapshot(rows)
    date_range = f"{overall['first_date'][:10]} a {overall['last_date'][:10]}" if overall["first_date"] else "sin datos"
    report_period = (infrastructure or {}).get("reportPeriod") or {}
    loaded_range = (
        f"{report_period.get('start')} a {report_period.get('end')}"
        if report_period.get("start") and report_period.get("end") else date_range
    )

    quality_html = render_quality_html(quality or {})
    body = f'''<div class="wrap">
      <header class="page-head">
        <h1>Reporte — {esc(env)}</h1>
        <div class="meta">Generado {datetime.now():%Y-%m-%d %H:%M} · periodo cargado: {esc(loaded_range)} · documentos existentes: {esc(date_range)} · {overall['n']} documento(s)</div>
      </header>

      <section class="card filters-card">
        <div class="filters-head">
          <button class="filters-toggle" id="f-toggle" aria-expanded="false" aria-controls="filters-body">
            <span class="ft-caret">▸</span> Filtros <span class="ft-badge" id="f-badge" hidden>0</span>
          </button>
          <span class="filters-summary" id="f-summary"></span>
          <button id="f-reset">Limpiar filtros</button>
        </div>
        <div class="filters" id="filters-body" hidden>
          <div class="filter-group">
            <span class="fg-label">Fecha</span>
            <div class="date-inputs">
              <input type="date" id="f-desde"> <span class="muted">a</span> <input type="date" id="f-hasta">
            </div>
          </div>
          <div class="filter-group"><label class="fg-label" for="f-type-search">Tipo de documento</label><input class="filter-search" type="search" id="f-type-search" placeholder="Buscar tipo o nombre…" autocomplete="off"><div class="chip-row" id="f-types"></div></div>
          <div class="filter-group"><span class="fg-label">Procesamiento</span><div class="chip-row" id="f-modes"></div></div>
          <div class="filter-group"><span class="fg-label">Sistema</span><div class="chip-row" id="f-apps"></div></div>
          <div class="filter-group"><span class="fg-label">Modelo</span><div class="chip-row" id="f-models"></div></div>
          <div class="filter-group"><span class="fg-label">Estado de ejecución</span><div class="chip-row" id="f-workflow"></div></div>
          <div class="filter-group"><span class="fg-label">Estado de negocio</span><div class="chip-row" id="f-validation"></div></div>
          <div class="filter-group"><span class="fg-label">Tráfico</span><div class="chip-row" id="f-traffic"></div></div>
        </div>
      </section>

      <div class="kpi-row" id="kpis"></div>

      <details class="card" open>
        <summary>Procesamiento con y sin IA</summary>
        <p class="section-sub">Clasificación por forma del catálogo (<code>xpath</code>/<code>cell</code>/<code>column</code>) y uso real de Bedrock, nunca por una lista fija de TYPE_IDs. <strong>Sin IA</strong> conserva el costo de Lambda y otros servicios que sí hayan participado.</p>
        <div class="kpi-row dense" id="ai-kpis"></div>
        <div id="ai-table"></div>
      </details>

      <details class="card" open>
        <summary>Infraestructura del stack</summary>
        <p class="section-sub">Inventario AWS en vivo y separado del costo por documento. El piso mensual no cambia con los filtros; S3, DynamoDB, API Gateway y API keys no tienen renta fija y se cobran por consumo.</p>
        <div class="kpi-row dense" id="infra-kpis"></div>
        <div id="infra-env-note"></div>
        <div class="sub-head">Piso fijo desplegado</div>
        <div id="infra-fixed-table"></div>
        <div class="sub-head">Gasto fijo: real vs proyección, mes a mes</div>
        <div class="kpi-row dense" id="infra-history-kpis"></div>
        <div id="infra-fixed-history"></div>
        <div class="sub-head">Gasto real del reporte, mes a mes</div>
        <div class="kpi-row dense" id="monthly-spend-kpis"></div>
        <div id="monthly-spend-chart"></div>
        <div class="sub-head">Servicios sin renta fija</div>
        <div id="infra-usage-table"></div>
        <div class="sub-head">Fallback manual por servicios de la infraestructura</div>
        <div id="infra-manual-table"></div>
        <div class="sub-head">Consumo MTD atribuido por CostStack</div>
        <div id="infra-tagged-table"></div>
        <div id="infra-warnings"></div>
      </details>

      {quality_html}

      <section class="card">
        <h2>Estado de las ejecuciones</h2>
        <p class="section-sub">Completa / en progreso / fallida.</p>
        <div id="wf-stack"></div>
      </section>

      <section class="card">
        <h2>Resultado de negocio</h2>
        <p class="section-sub">Aprobado / requiere revisión / rechazado, solo de ejecuciones completas.</p>
        <div id="validation-stack"></div>
      </section>

      <details class="card" open>
        <summary>Tiempo de proceso</summary>
        <p class="section-sub">De punta a punta (<code>Tracking.ProcessingSeconds</code>): del arranque del enriquecimiento al resultado guardado. Solo los documentos que lo traen registrado entran en estos números.</p>
        <div class="kpi-row dense" id="duration-kpis"></div>
        <div class="sub-head">Distribución</div>
        <div class="bars" id="duration-hist"></div>
        <div class="sub-head">Por tipo de documento</div>
        <div id="duration-by-type"></div>
        <div class="sub-head">Por sistema</div>
        <div id="duration-by-app"></div>
        <div class="sub-head">Promedio por servicio</div>
        <p class="section-sub"><code>durationMs</code> del registro de costo — solo existe para documentos con auditoría de costo (2026-08-05 en adelante).</p>
        <div class="bars" id="duration-by-service"></div>
        <div class="sub-head">Los 10 más lentos</div>
        <div id="duration-slowest"></div>
      </details>

      <details class="card" open>
        <summary>Costo contra las referencias medidas</summary>
        <p class="section-sub">Contraste del corte filtrado contra las mediciones reales de <code>documentacion/COSTOS.md</code> (Bedrock <code>usage</code> real + páginas reales, cruzados contra Cost Explorer). Δ verde ≤ 5 %, ámbar ≤ 30 %, rojo por encima.</p>
        <div id="ref-delta-legend"></div>
        <div class="kpi-row dense" id="ref-kpis"></div>
        <div class="sub-head">Costo por página vs referencias</div>
        <div class="bars" id="ref-bars"></div>
        <div class="sub-head">Costo por documento y por página, por tipo</div>
        <div id="ref-by-type"></div>
      </details>

      <details class="card" id="page-sel-card" hidden open>
        <summary>Ahorro por selección de páginas</summary>
        <p class="section-sub">Solo se mandan al modelo las páginas que traen los campos del tipo (<code>page_selection</code> del catálogo). <strong>Facturado</strong> lo reportó Bedrock; <strong>evitado</strong> estima la llamada que nunca se hizo, así que <em>no</em> está sumado al costo total del corte.</p>
        <div class="kpi-row dense" id="page-sel-kpis"></div>
        <div id="page-sel-table"></div>
      </details>

      <details class="card" open>
        <summary>Volumen diario</summary>
        <div id="volume-chart"></div>
      </details>

      <details class="card" open>
        <summary>Por tipo de documento</summary>
        <div id="by-type-table"></div>
      </details>

      <details class="card" open>
        <summary>Por sistema</summary>
        <p class="section-sub">Sistema/app que subió el documento (<code>AppOrigin</code>, resuelto de la API key del request).</p>
        <div class="bars" id="by-app-bars"></div>
        <div id="by-app-table"></div>
      </details>

      <details class="card" open>
        <summary>Por modelo</summary>
        <div class="bars" id="by-model-bars"></div>
        <div id="by-model-table"></div>
      </details>

      <details class="card">
        <summary>Costo por servicio AWS</summary>
        <p class="section-sub">Click en un servicio para ocultarlo/mostrarlo en la gráfica.</p>
        <div class="chip-row" id="f-services" style="margin-bottom:14px"></div>
        <div class="bars" id="by-service-bars"></div>
      </details>

      <details class="card">
        <summary>Campos que más fallan</summary>
        <p class="section-sub">Cobertura parcial en expedientes multi-documento — el detalle completo de cada documento interno solo vive en <code>results/&#123;id&#125;.json</code>.</p>
        <div id="worst-fields-table"></div>
      </details>

      <details class="card" open>
        <summary>Documentos (<span id="doc-count">0</span>)</summary>
        <p class="section-sub">Click en un renglón para ver el detalle completo (costo + calidad). Click en un encabezado para ordenar.</p>
        <div class="table-wrap" id="documents-table"></div>
        <div id="documents-more"></div>
      </details>

      <dialog id="detail-dialog"><div class="detail-inner" id="detail-content"></div></dialog>
      <div class="report-tooltip" id="report-tooltip" role="tooltip" aria-hidden="true"></div>

      <footer>Generado por <code>pnpm reporte</code> (scripts/reporte.py) · une la tabla de costos y la de tracking de documentos por <code>DocumentId</code> · datos crudos en <code>datos_raw.xlsx</code> · reportes individuales: <code>pnpm reporte:costos</code> / <code>pnpm reporte:documentos</code>.</footer>
    </div>
    <script id="report-data" type="application/json">{json.dumps({"rows": rows, "catalogTypes": sorted(ACTIVE_CATALOG_TYPES) if catalog_types is None else sorted(catalog_types), "costRef": COST_REFERENCE, "infrastructure": infrastructure or {}, "monthlySpend": monthly_spend_snapshot(rows, infrastructure)}, ensure_ascii=False, default=_json_default)}</script>
    <script>{JS_TEMPLATE}</script>'''

    return f'''<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Reporte — {esc(env)}</title>
<style>{BASE_CSS}{EXTRA_CSS}</style>
</head>
<body>{body}</body>
</html>'''


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def write_report_bundle(
    out_dir: Path, rows: List[Dict[str, Any]], quality: Dict[str, Any], env: str,
    infrastructure: Optional[Dict[str, Any]] = None,
    catalog_types: Optional[List[str]] = None,
) -> Dict[str, Any]:
    report_catalog_types = sorted(ACTIVE_CATALOG_TYPES) if catalog_types is None else sorted(catalog_types)
    hydrate_processing_fields(rows)
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "reporte.md").write_text(render_markdown(rows, env, quality, infrastructure), encoding="utf-8")
    (out_dir / "reporte.html").write_text(
        render_html(rows, env, quality, infrastructure, report_catalog_types), encoding="utf-8"
    )
    write_excel(out_dir / "datos_raw.xlsx", rows, quality, infrastructure)
    shutil.copy2(out_dir / "datos_raw.xlsx", out_dir / "reporte_ejecutivo.xlsx")
    snapshot = {"schemaVersion": 4, "environment": env, "generatedAt": datetime.now().astimezone().isoformat(),
                "catalogTypes": report_catalog_types, "rows": rows, "quality": quality,
                "infrastructure": infrastructure or {}}
    snapshot_path = out_dir / "snapshot.json"
    snapshot_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2, default=_json_default), encoding="utf-8")
    os.chmod(snapshot_path, 0o600)
    return overall_snapshot(rows)

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--env", choices=("dev", "prod"), default="dev", help="ambiente a consultar (default: dev)")
    parser.add_argument("--region", default="us-east-1")
    parser.add_argument("--year", type=int, default=date.today().year,
                        help="año calendario por defecto, de enero 1 a diciembre 31")
    parser.add_argument("--stack", help="nombre del stack (default: resuelto de --env)")
    parser.add_argument("--table", help="tabla de tracking de documentos, salta la resolución vía CloudFormation")
    parser.add_argument("--cost-table", help="tabla de costos (default: resuelta de --env)")
    parser.add_argument("--desde", help="filtra ProcessedAt >= YYYY-MM-DD (también filtrable interactivo en el HTML)")
    parser.add_argument("--hasta", help="filtra ProcessedAt <= YYYY-MM-DD")
    parser.add_argument("--sistema", help="filtra por AppOrigin (uno o varios separados por coma; también filtrable interactivo en el HTML)")
    parser.add_argument("--out", help="carpeta de salida (default: reportes/<env>/<timestamp>/)")
    parser.add_argument("--open", action="store_true", help="abrir el HTML en el navegador al terminar")
    parser.add_argument("--history-dir", default=str(DEFAULT_HISTORY), help="historial local de pruebas")
    parser.add_argument("--snapshot", help="regenera el reporte offline desde un snapshot.json anterior")
    args = parser.parse_args([a for a in sys.argv[1:] if a != "--"])

    try:
        report_start, report_end = resolve_report_period(args.year, args.desde, args.hasta)
    except ValueError as exc:
        parser.error(f"fecha inválida: {exc}")

    if args.snapshot:
        saved = json.loads(Path(args.snapshot).read_text(encoding="utf-8"))
        rows, quality = saved.get("rows") or [], saved.get("quality") or quality_snapshot([])
        infrastructure = saved.get("infrastructure") or empty_infrastructure_snapshot(
            args.stack or stack_id(saved.get("environment") or args.env), args.region,
            "El snapshot original no contenía inventario de infraestructura.",
        )
        env = saved.get("environment") or args.env
        out_dir = Path(args.out) if args.out else DEFAULT_OUT_ROOT / env / datetime.now().strftime("%Y%m%d-%H%M%S")
        overall = write_report_bundle(
            out_dir, rows, quality, env, infrastructure, saved.get("catalogTypes")
        )
        print(f"Reporte reproducido desde {args.snapshot}: {overall['n']} documento(s) → {out_dir}/")
        if args.open:
            webbrowser.open((out_dir / "reporte.html").resolve().as_uri())
        return 0

    stack = args.stack or stack_id(args.env)
    cost_table = args.cost_table or cost_table_name(args.env)

    cost_raw: List[Dict[str, Any]] = []
    try:
        cost_raw = scan_all(cost_table, args.region)
    except Exception as e:
        print(f"{YELLOW}No pude leer la tabla de costos {cost_table}: {e} — el reporte sigue sin datos de costo.{RESET}")

    try:
        doc_table = args.table or resolve_processing_table(stack, args.region)
    except Exception as e:
        print(f"{RED}{e}{RESET}")
        return 1

    print(
        f"Leyendo {doc_table} y {cost_table} en {args.region} (stack {stack}) · "
        f"periodo {report_start.isoformat()} a {report_end.isoformat()}..."
    )
    try:
        doc_raw = scan_all(doc_table, args.region)
    except Exception as e:
        print(f"{RED}No pude leer la tabla {doc_table} en {args.region}: {e}{RESET}")
        return 1

    quality_runs = load_runs("local", Path(args.history_dir)) + load_runs(args.env, Path(args.history_dir))
    quality_runs.sort(key=lambda r: r.get("startedAt", ""))
    quality = quality_snapshot(quality_runs)

    if not cost_raw and not doc_raw:
        print(f"{YELLOW}Ambas tablas están vacías — el reporte contendrá únicamente el histórico de pruebas.{RESET}")

    cost_rows = [normalize_cost_item(it) for it in cost_raw]
    doc_rows = [normalize_doc_item(it) for it in doc_raw]
    rows = join_rows(cost_rows, doc_rows)

    report_start_text, report_end_text = report_start.isoformat(), report_end.isoformat()
    rows = [r for r in rows if report_start_text <= r["processedAt"][:10] <= report_end_text]
    if not rows:
        print(f"{YELLOW}Ningún documento cae en {report_start_text} a {report_end_text}; "
              f"el reporte continuará con infraestructura y costos fijos.{RESET}")

    if args.sistema:
        wanted = {s.strip().lower() for s in args.sistema.split(",") if s.strip()}
        rows = [r for r in rows if app_of(r).lower() in wanted]
        if not rows:
            print(f"{YELLOW}Ningún documento del sistema pedido ({args.sistema}).{RESET}")
            return 0

    try:
        print(f"Verificando infraestructura y gasto fijo de {stack}...")
        infrastructure = collect_infrastructure_costs(
            stack, args.region, environment=args.env,
            period_start=report_start, period_end=report_end,
        )
    except Exception as e:
        infrastructure = empty_infrastructure_snapshot(
            stack, args.region, f"No se pudo inventariar/costear la infraestructura del stack: {e}"
        )
        print(f"{YELLOW}{infrastructure['warnings'][0]}{RESET}")
    infrastructure["reportPeriod"] = {"start": report_start_text, "end": report_end_text}

    out_dir = Path(args.out) if args.out else DEFAULT_OUT_ROOT / args.env / datetime.now().strftime("%Y%m%d-%H%M%S")
    overall = write_report_bundle(out_dir, rows, quality, args.env, infrastructure)
    monthly_spend = monthly_spend_snapshot(rows, infrastructure)
    fixed_history = infrastructure.get("fixedMonthlyHistory") or {}
    print(
        f"\n{CYAN}{overall['n']} documento(s) ({overall['n_with_cost']} con costo) · "
        f"{overall['n_no_ai']} sin IA · aprobación {pct(overall['approval_rate'])} · "
        f"periodo {report_start.isoformat()} a {report_end.isoformat()}{RESET}\n"
        f"{CYAN}Gastado real conocido {money(monthly_spend.get('knownActualUsd'))} = "
        f"procesamiento real {money(monthly_spend.get('processingActualUsd'))} + "
        f"fijo real {money(monthly_spend.get('fixedActualUsd'))}{RESET}\n"
        f"{YELLOW}Proyección fija del periodo {money(fixed_history.get('projectedPeriodUsd'))} · "
        f"piso fijo mensual actual {money(infrastructure.get('fixedMonthlyUsd'))}{RESET}"
    )
    print(f"Reporte → {out_dir}/")
    print("  reporte.md        (resumen estático)")
    print("  reporte.html       (dashboard interactivo — ábrelo en el navegador)")
    print("  datos_raw.xlsx     (datos crudos — Excel)")
    print("  reporte_ejecutivo.xlsx (Excel ejecutivo con KPIs e histórico)")
    print("  snapshot.json      (corte reproducible)")

    if args.open:
        webbrowser.open((out_dir / "reporte.html").resolve().as_uri())

    return 0


if __name__ == "__main__":
    sys.exit(main())
