#!/usr/bin/env python3
"""Genera un reporte local de costos reales de AWS a partir de la tabla de
auditoría `document-validation-costs[-dev]` (la que llena `cost_audit.py` en
cada documento procesado): totales, comparación entre modelos, comparación
entre tipos de documento, costo por servicio y los documentos más caros —
tokens de entrada/salida y costo, por fase y en total.

No llama a Bedrock/Textract ni reprocesa nada — solo lee DynamoDB y calcula
todo localmente. Corre 100% en tu máquina; los archivos quedan en
`reportes-costos/<env>/<timestamp>/` para abrir directo:

    reporte.md          resumen en texto/tablas (para pegar en un mensaje o leer en terminal)
    reporte.html         dashboard con KPIs y comparativas (abrir en el navegador)
    costos_raw.xlsx      datos crudos: una fila por documento + hojas agregadas, para pivotear en Excel

Uso:
    pnpm reporte:costos                       # dev, todos los documentos
    pnpm reporte:costos -- --env prod
    pnpm reporte:costos -- --desde 2026-08-01 --hasta 2026-08-17
    pnpm reporte:costos -- --limit 30 --open  # top 30 documentos más caros, abre el HTML al terminar
    pnpm reporte:costos -- --out /ruta/propia
"""

from __future__ import annotations

import argparse
import sys
import webbrowser
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _report_lib import (  # noqa: E402
    BASE_CSS,
    BEDROCK_SERVICES,
    CATEGORICAL,
    CYAN,
    DIM,
    RED,
    RESET,
    SERVICE_LABEL_ES,
    YELLOW,
    cost_table_name as table_name,
    esc,
    h_bar,
    kpi_tile,
    money,
    normalize_cost_item as normalize_item,
    num,
    scan_all,
)

REPO = Path(__file__).resolve().parent.parent
DEFAULT_OUT_ROOT = REPO / "reportes-costos"


def filter_by_date(rows: List[Dict[str, Any]], desde: Optional[str], hasta: Optional[str]) -> List[Dict[str, Any]]:
    if not desde and not hasta:
        return rows
    out = []
    for r in rows:
        day = r["processed_at"][:10]
        if desde and day < desde:
            continue
        if hasta and day > hasta:
            continue
        out.append(r)
    return out


def aggregate_overall(rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    n = len(rows)
    total_cost = sum(r["total_cost"] for r in rows)
    total_pages = sum(r["pages_total"] for r in rows)
    total_in = sum(r["total_input_tokens"] for r in rows)
    total_out = sum(r["total_output_tokens"] for r in rows)
    dates = sorted(r["processed_at"] for r in rows if r["processed_at"])
    return {
        "n": n,
        "total_cost": total_cost,
        "total_pages": total_pages,
        "total_input_tokens": total_in,
        "total_output_tokens": total_out,
        "avg_cost_doc": (total_cost / n) if n else None,
        "avg_cost_page": (total_cost / total_pages) if total_pages else None,
        "first_date": dates[0] if dates else None,
        "last_date": dates[-1] if dates else None,
        "pricing_versions": sorted({r["pricing_version"] for r in rows if r["pricing_version"]}),
    }


def aggregate_page_selection(rows: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """Lo que la selección de páginas evitó gastar en el corte completo.

    Devuelve None si ningún documento la tuvo activa, para no renderizar una
    sección de ceros. `cost_avoided` es una ESTIMACIÓN de la llamada que no se
    hizo — se reporta junto al costo real pero nunca sumado a él.
    """
    active = [r["page_selection"] for r in rows if r.get("page_selection")]
    if not active:
        return None
    cost_avoided = sum(p["estimated_cost_avoided"] for p in active)
    billed_cost = sum(r["total_cost"] for r in rows if r.get("page_selection"))
    in_saved = sum(p["input_tokens_saved"] for p in active)
    out_saved = sum(p["output_tokens_saved"] for p in active)
    in_billed = sum(p["billed_input_tokens"] for p in active)
    out_billed = sum(p["billed_output_tokens"] for p in active)
    return {
        "docs": len(active),
        "docs_total": len(rows),
        "pages_considered": sum(p["pages_considered"] for p in active),
        "pages_extracted": sum(p["pages_extracted"] for p in active),
        "pages_skipped": sum(p["pages_skipped"] for p in active),
        "chunks_sent": sum(p["chunks_sent"] for p in active),
        "chunks_without_selection": sum(p["chunks_without_selection"] for p in active),
        "billed_input_tokens": in_billed,
        "billed_output_tokens": out_billed,
        "input_tokens_saved": in_saved,
        "output_tokens_saved": out_saved,
        "input_saving_pct": (in_saved / (in_billed + in_saved) * 100) if (in_billed + in_saved) else None,
        "cost_avoided": cost_avoided,
        "billed_cost": billed_cost,
        # Lo que habría costado el corte SIN selección, para el titular.
        "cost_without_selection": billed_cost + cost_avoided,
    }


def aggregate_by_type(rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    by_type: Dict[str, Dict[str, Any]] = defaultdict(
        lambda: {"n": 0, "pages": 0, "cost": 0.0, "input_tokens": 0, "output_tokens": 0, "name": ""}
    )
    for r in rows:
        agg = by_type[r["document_type"]]
        agg["n"] += 1
        agg["pages"] += r["pages_total"]
        agg["cost"] += r["total_cost"]
        agg["input_tokens"] += r["total_input_tokens"]
        agg["output_tokens"] += r["total_output_tokens"]
        agg["name"] = r["document_type_name"] or agg["name"]
    return dict(by_type)


def aggregate_by_model(rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    by_model: Dict[str, Dict[str, Any]] = defaultdict(
        lambda: {"calls": 0, "docs": set(), "cost": 0.0, "input_tokens": 0, "output_tokens": 0, "approximate": False}
    )
    for r in rows:
        for svc_name, svc in r["services"].items():
            if svc_name not in BEDROCK_SERVICES:
                continue
            model_id = svc["model_id"]
            if not model_id or not (svc["input_tokens"] or svc["output_tokens"] or svc["cost"]):
                continue
            agg = by_model[model_id]
            agg["calls"] += 1
            agg["docs"].add(r["document_id"])
            agg["cost"] += svc["cost"]
            agg["input_tokens"] += svc["input_tokens"]
            agg["output_tokens"] += svc["output_tokens"]
            agg["approximate"] = agg["approximate"] or svc["approximate_pricing"]
    return dict(by_model)


def aggregate_by_service(rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    by_service: Dict[str, Dict[str, Any]] = defaultdict(
        lambda: {"cost": 0.0, "input_tokens": 0, "output_tokens": 0, "docs": 0}
    )
    for r in rows:
        for name, svc in r["services"].items():
            if not svc["cost"] and not svc["input_tokens"] and not svc["output_tokens"]:
                continue
            agg = by_service[name]
            agg["cost"] += svc["cost"]
            agg["input_tokens"] += svc["input_tokens"]
            agg["output_tokens"] += svc["output_tokens"]
            agg["docs"] += 1
    return dict(by_service)


def top_expensive(rows: List[Dict[str, Any]], limit: int) -> List[Dict[str, Any]]:
    return sorted(rows, key=lambda r: r["total_cost"], reverse=True)[:limit]


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

def render_markdown(
    rows: List[Dict[str, Any]],
    overall: Dict[str, Any],
    by_type: Dict[str, Dict[str, Any]],
    by_model: Dict[str, Dict[str, Any]],
    by_service: Dict[str, Dict[str, Any]],
    top: List[Dict[str, Any]],
    env: str,
    page_sel: Optional[Dict[str, Any]] = None,
) -> str:
    lines = [
        f"# Reporte de costos — {env}",
        "",
        f"Generado {datetime.now():%Y-%m-%d %H:%M} · {overall['n']} documento(s)"
        + (f" · {overall['first_date'][:10]} a {overall['last_date'][:10]}" if overall["first_date"] else ""),
        "",
        "## Totales",
        "",
        "| Métrica | Valor |",
        "| --- | --- |",
        f"| Documentos procesados | {overall['n']} |",
        f"| Páginas procesadas | {num(overall['total_pages'])} |",
        f"| Costo total | {money(overall['total_cost'])} |",
        f"| Costo promedio / documento | {money(overall['avg_cost_doc'])} |",
        f"| Costo promedio / página | {money(overall['avg_cost_page'])} |",
        f"| Tokens de entrada (total) | {num(overall['total_input_tokens'])} |",
        f"| Tokens de salida (total) | {num(overall['total_output_tokens'])} |",
        f"| Versión(es) de tarifa | {', '.join(overall['pricing_versions']) or '—'} |",
        "",
    ]

    if page_sel:
        lines += [
            "## Ahorro por selección de páginas",
            "",
            f"Activa en {page_sel['docs']} de {page_sel['docs_total']} documento(s): se extrajeron "
            f"{num(page_sel['pages_extracted'])} de {num(page_sel['pages_considered'])} páginas.",
            "",
            "| | Facturado | Evitado (est.) | Sin selección |",
            "| --- | --- | --- | --- |",
            f"| Tokens de entrada | {num(page_sel['billed_input_tokens'])} | {num(page_sel['input_tokens_saved'])} | "
            f"{num(page_sel['billed_input_tokens'] + page_sel['input_tokens_saved'])} |",
            f"| Tokens de salida | {num(page_sel['billed_output_tokens'])} | {num(page_sel['output_tokens_saved'])} | "
            f"{num(page_sel['billed_output_tokens'] + page_sel['output_tokens_saved'])} |",
            f"| Llamadas a Bedrock | {num(page_sel['chunks_sent'])} | "
            f"{num(page_sel['chunks_without_selection'] - page_sel['chunks_sent'])} | {num(page_sel['chunks_without_selection'])} |",
            f"| Costo | {money(page_sel['billed_cost'])} | {money(page_sel['cost_avoided'])} | "
            f"{money(page_sel['cost_without_selection'])} |",
            "",
            f"Esos documentos costaron {money(page_sel['billed_cost'])} en vez de "
            f"~{money(page_sel['cost_without_selection'])}.",
            "",
            "> **Facturado** lo reportó Bedrock. **Evitado** estima la llamada que nunca se hizo "
            "(entrada por densidad tokens/carácter de la llamada que sí corrió, salida por chunks "
            "evitados), así que no está incluido en el costo total de arriba — sumarlos daría un "
            "gasto que nadie pagó.",
            "",
        ]

    lines += ["## Por modelo", "", "| Modelo | Docs | Llamadas | Tokens in | Tokens out | Costo | $/doc |", "| --- | --- | --- | --- | --- | --- | --- |"]
    for model, agg in sorted(by_model.items(), key=lambda kv: kv[1]["cost"], reverse=True):
        n_docs = len(agg["docs"])
        per_doc = agg["cost"] / n_docs if n_docs else 0
        mark = " *" if agg["approximate"] else ""
        lines.append(
            f"| {model}{mark} | {n_docs} | {agg['calls']} | {num(agg['input_tokens'])} | "
            f"{num(agg['output_tokens'])} | {money(agg['cost'])} | {money(per_doc)} |"
        )
    if any(a["approximate"] for a in by_model.values()):
        lines.append("")
        lines.append("\\* tarifa aproximada — modelo sin precio verificado en `pipeline/pricing.py`.")
    lines.append("")

    lines += ["## Por tipo de documento", "", "| Tipo | Docs | Páginas | Tokens in | Tokens out | Costo | $/doc | $/página |", "| --- | --- | --- | --- | --- | --- | --- | --- |"]
    for dtype, agg in sorted(by_type.items(), key=lambda kv: kv[1]["cost"], reverse=True):
        per_doc = agg["cost"] / agg["n"] if agg["n"] else 0
        per_page = agg["cost"] / agg["pages"] if agg["pages"] else None
        label = agg["name"] or dtype
        lines.append(
            f"| {label} | {agg['n']} | {num(agg['pages'])} | {num(agg['input_tokens'])} | "
            f"{num(agg['output_tokens'])} | {money(agg['cost'])} | {money(per_doc)} | {money(per_page)} |"
        )
    lines.append("")

    lines += ["## Por servicio AWS", "", "| Servicio | Docs | Tokens in | Tokens out | Costo | % del total |", "| --- | --- | --- | --- | --- | --- |"]
    for svc_name, agg in sorted(by_service.items(), key=lambda kv: kv[1]["cost"], reverse=True):
        pct_total = (agg["cost"] / overall["total_cost"] * 100) if overall["total_cost"] else 0
        lines.append(
            f"| {SERVICE_LABEL_ES.get(svc_name, svc_name)} | {agg['docs']} | {num(agg['input_tokens'])} | "
            f"{num(agg['output_tokens'])} | {money(agg['cost'])} | {pct_total:.1f}% |"
        )
    lines.append("")

    lines += [f"## Documentos más caros (top {len(top)})", "", "| Documento | Tipo | Páginas | Tokens in/out | Costo | Procesado |", "| --- | --- | --- | --- | --- | --- |"]
    for r in top:
        lines.append(
            f"| `{r['document_id']}` | {r['document_type_name'] or r['document_type']} | {r['pages_total']} | "
            f"{num(r['total_input_tokens'])}/{num(r['total_output_tokens'])} | {money(r['total_cost'])} | {r['processed_at'][:19]} |"
        )
    lines.append("")
    lines.append(f"_Generado por `pnpm reporte:costos` (scripts/reporte-costos.py) — datos crudos: `costos_raw.xlsx`._")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

def render_html(
    rows: List[Dict[str, Any]],
    overall: Dict[str, Any],
    by_type: Dict[str, Dict[str, Any]],
    by_model: Dict[str, Dict[str, Any]],
    by_service: Dict[str, Dict[str, Any]],
    top: List[Dict[str, Any]],
    env: str,
    page_sel: Optional[Dict[str, Any]] = None,
) -> str:
    models_ranked = sorted(by_model.items(), key=lambda kv: kv[1]["cost"], reverse=True)
    model_color = {m: CATEGORICAL[i % len(CATEGORICAL)] for i, (m, _) in enumerate(models_ranked)}
    max_model_cost = max((a["cost"] for _, a in models_ranked), default=0)
    model_bars = "".join(
        h_bar(agg["cost"], max_model_cost, model_color[m], f"{m}  ({len(agg['docs'])} docs)", money(agg["cost"]))
        for m, agg in models_ranked
    )
    model_rows = "".join(
        f'''<tr>
          <td><span class="dot" style="background:{model_color[m]}"></span>{esc(m)}{" *" if agg["approximate"] else ""}</td>
          <td class="num">{len(agg['docs'])}</td>
          <td class="num">{agg['calls']}</td>
          <td class="num">{num(agg['input_tokens'])}</td>
          <td class="num">{num(agg['output_tokens'])}</td>
          <td class="num">{money(agg['cost'])}</td>
          <td class="num">{money(agg['cost'] / len(agg['docs']) if agg['docs'] else 0)}</td>
        </tr>'''
        for m, agg in models_ranked
    )
    approx_note = (
        '<p class="note">* costo aproximado — modelo sin tarifa verificada en <code>pipeline/pricing.py</code>.</p>'
        if any(a["approximate"] for _, a in models_ranked) else ""
    )

    type_rows = "".join(
        f'''<tr>
          <td>{esc(agg['name'] or dtype)}</td>
          <td class="num">{agg['n']}</td>
          <td class="num">{num(agg['pages'])}</td>
          <td class="num">{num(agg['input_tokens'])}</td>
          <td class="num">{num(agg['output_tokens'])}</td>
          <td class="num">{money(agg['cost'])}</td>
          <td class="num">{money(agg['cost'] / agg['n'] if agg['n'] else 0)}</td>
          <td class="num">{money(agg['cost'] / agg['pages'] if agg['pages'] else None)}</td>
        </tr>'''
        for dtype, agg in sorted(by_type.items(), key=lambda kv: kv[1]["cost"], reverse=True)
    )

    max_service_cost = max((a["cost"] for a in by_service.values()), default=0)
    service_bars = "".join(
        h_bar(
            agg["cost"], max_service_cost, CATEGORICAL[i % len(CATEGORICAL)],
            SERVICE_LABEL_ES.get(name, name), money(agg["cost"]),
        )
        for i, (name, agg) in enumerate(sorted(by_service.items(), key=lambda kv: kv[1]["cost"], reverse=True))
    )

    top_rows = "".join(
        f'''<tr>
          <td class="mono">{esc(r['document_id'])}</td>
          <td>{esc(r['document_type_name'] or r['document_type'])}</td>
          <td class="num">{r['pages_total']}</td>
          <td class="num">{num(r['total_input_tokens'])} / {num(r['total_output_tokens'])}</td>
          <td class="num">{money(r['total_cost'])}</td>
          <td class="muted">{esc(r['processed_at'][:19])}</td>
        </tr>'''
        for r in top
    )

    kpis = "".join([
        kpi_tile(str(overall["n"]), "Documentos procesados"),
        kpi_tile(num(overall["total_pages"]), "Páginas procesadas"),
        kpi_tile(money(overall["total_cost"]), "Costo total real"),
        kpi_tile(money(overall["avg_cost_doc"]), "Costo promedio / documento"),
        kpi_tile(money(overall["avg_cost_page"]), "Costo promedio / página"),
        kpi_tile(num(overall["total_input_tokens"]), "Tokens de entrada (total)"),
        kpi_tile(num(overall["total_output_tokens"]), "Tokens de salida (total)"),
        kpi_tile(str(len(by_model)), "Modelos usados"),
    ])

    date_range = (
        f"{overall['first_date'][:10]} a {overall['last_date'][:10]}" if overall["first_date"] else "sin datos"
    )

    page_sel_section = ""
    if page_sel:
        chunks_saved = page_sel["chunks_without_selection"] - page_sel["chunks_sent"]
        page_sel_section = f'''
      <section class="card">
        <h2>Ahorro por selección de páginas</h2>
        <p class="section-sub">
          Activa en {page_sel['docs']} de {page_sel['docs_total']} documento(s): se extrajeron
          {num(page_sel['pages_extracted'])} de {num(page_sel['pages_considered'])} páginas
          ({num(page_sel['pages_skipped'])} descartadas sin mandarlas al modelo).
          Esos documentos costaron <strong>{money(page_sel['billed_cost'])}</strong> en vez de
          ~<strong>{money(page_sel['cost_without_selection'])}</strong>.
        </p>
        <div class="table-wrap">
          <table>
            <thead><tr><th></th><th class="num">Facturado</th><th class="num">Evitado (est.)</th><th class="num">Sin selección</th></tr></thead>
            <tbody>
              <tr><td>Tokens de entrada</td><td class="num">{num(page_sel['billed_input_tokens'])}</td><td class="num">{num(page_sel['input_tokens_saved'])}</td><td class="num">{num(page_sel['billed_input_tokens'] + page_sel['input_tokens_saved'])}</td></tr>
              <tr><td>Tokens de salida</td><td class="num">{num(page_sel['billed_output_tokens'])}</td><td class="num">{num(page_sel['output_tokens_saved'])}</td><td class="num">{num(page_sel['billed_output_tokens'] + page_sel['output_tokens_saved'])}</td></tr>
              <tr><td>Llamadas a Bedrock</td><td class="num">{num(page_sel['chunks_sent'])}</td><td class="num">{num(chunks_saved)}</td><td class="num">{num(page_sel['chunks_without_selection'])}</td></tr>
              <tr><td><strong>Costo</strong></td><td class="num"><strong>{money(page_sel['billed_cost'])}</strong></td><td class="num"><strong>{money(page_sel['cost_avoided'])}</strong></td><td class="num"><strong>{money(page_sel['cost_without_selection'])}</strong></td></tr>
            </tbody>
          </table>
        </div>
        <p class="section-sub"><strong>Facturado</strong> lo reportó Bedrock.
          <strong>Evitado</strong> estima la llamada que nunca se hizo (entrada por densidad
          tokens/carácter de la llamada que sí corrió, salida por chunks evitados), así que
          <em>no</em> está incluido en el costo total de arriba — sumarlos daría un gasto que
          nadie pagó.</p>
      </section>
'''

    body = f'''<div class="wrap">
      <header class="page-head">
        <h1>Reporte de costos de procesamiento — {esc(env)}</h1>
        <div class="meta">Generado {datetime.now():%Y-%m-%d %H:%M} · rango de documentos: {esc(date_range)} · tarifa(s): {esc(", ".join(overall["pricing_versions"]) or "—")}</div>
      </header>
      <div class="kpi-row">{kpis}</div>
{page_sel_section}
      <section class="card">
        <h2>Costo por modelo</h2>
        <p class="section-sub">Compara qué modelo procesó más documentos y a qué costo — útil para decidir si vale la pena cambiar de modelo.</p>
        <div class="bars">{model_bars}</div>
        <div class="table-wrap">
          <table>
            <thead><tr><th>Modelo</th><th class="num">Docs</th><th class="num">Llamadas</th><th class="num">Tokens in</th><th class="num">Tokens out</th><th class="num">Costo</th><th class="num">$/doc</th></tr></thead>
            <tbody>{model_rows}</tbody>
          </table>
        </div>
        {approx_note}
      </section>

      <section class="card">
        <h2>Costo por tipo de documento</h2>
        <div class="table-wrap">
          <table>
            <thead><tr><th>Tipo</th><th class="num">Docs</th><th class="num">Páginas</th><th class="num">Tokens in</th><th class="num">Tokens out</th><th class="num">Costo</th><th class="num">$/doc</th><th class="num">$/página</th></tr></thead>
            <tbody>{type_rows}</tbody>
          </table>
        </div>
      </section>

      <section class="card">
        <h2>Costo por servicio AWS</h2>
        <p class="section-sub">A dónde va el dinero: Bedrock (por fase), Textract OCR y Lambda.</p>
        <div class="bars">{service_bars}</div>
      </section>

      <section class="card">
        <h2>Documentos más caros (top {len(top)})</h2>
        <div class="table-wrap">
          <table>
            <thead><tr><th>Documento</th><th>Tipo</th><th class="num">Páginas</th><th class="num">Tokens in/out</th><th class="num">Costo</th><th>Procesado</th></tr></thead>
            <tbody>{top_rows}</tbody>
          </table>
        </div>
      </section>

      <footer>Generado por <code>pnpm reporte:costos</code> (scripts/reporte-costos.py) · datos crudos en <code>costos_raw.xlsx</code> · fuente: tabla <code>document-validation-costs{"" if env == "prod" else "-dev"}</code>.</footer>
    </div>'''

    return f'''<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Reporte de costos — {esc(env)}</title>
<style>{CSS}</style>
</head>
<body>{body}</body>
</html>'''


CSS = BASE_CSS


# ---------------------------------------------------------------------------
# Excel (datos crudos)
# ---------------------------------------------------------------------------

def write_excel(
    out_path: Path,
    rows: List[Dict[str, Any]],
    overall: Dict[str, Any],
    by_type: Dict[str, Dict[str, Any]],
    by_model: Dict[str, Dict[str, Any]],
    by_service: Dict[str, Dict[str, Any]],
    page_sel: Optional[Dict[str, Any]] = None,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    bold = Font(bold=True)

    def write_sheet(ws, headers: List[str], data_rows: List[List[Any]]) -> None:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = bold
        for row in data_rows:
            ws.append(row)
        for col_idx, header in enumerate(headers, start=1):
            width = max(12, min(40, len(header) + 4))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    # Hoja 1: una fila por documento, servicios desplegados en columnas.
    service_names = sorted({name for r in rows for name in r["services"]})
    ws_docs = wb.active
    ws_docs.title = "Documentos"
    headers = [
        "DocumentId", "DocumentType", "DocumentTypeName", "SourceFile", "AppOrigin", "ProcessedAt",
        "PagesTotal", "PagesNative", "PagesOcr", "PagesEmpty",
        "TotalInputTokens", "TotalOutputTokens", "TotalCostUsd", "CostPerPageUsd", "PricingVersion",
        "AiExecuted", "ExtractionStrategies", "ExtractionProfiles", "ExtractionProfileStatuses",
        "Deterministic.InputTokensSaved", "Deterministic.OutputTokensSaved", "Deterministic.CostAvoidedUsd",
        # Selección de páginas: vacío en documentos donde no estuvo activa.
        # CostAvoided es ESTIMADO — no sumarlo a TotalCostUsd.
        "PageSel.PagesExtracted", "PageSel.PagesSkipped",
        "PageSel.InputTokensSaved", "PageSel.OutputTokensSaved", "PageSel.CostAvoidedUsd",
    ]
    for svc in service_names:
        headers += [f"{svc}.CostUsd", f"{svc}.InputTokens", f"{svc}.OutputTokens", f"{svc}.ModelId"]
    data_rows = []
    for r in rows:
        row = [
            r["document_id"], r["document_type"], r["document_type_name"], r["source_file"], r["app_origin"],
            r["processed_at"], r["pages_total"], r["pages_native"], r["pages_ocr"], r["pages_empty"],
            r["total_input_tokens"], r["total_output_tokens"], r["total_cost"], r["cost_per_page"], r["pricing_version"],
        ]
        audit = r.get("extraction_audit")
        row += [
            audit["ai_executed"],
            ", ".join(sorted({d["strategy"] for d in audit["documents"]})),
            ", ".join(sorted({d["profile"] for d in audit["documents"] if d.get("profile")})),
            ", ".join(sorted({d["profile_status"] for d in audit["documents"] if d.get("profile_status")})),
            audit["input_tokens_saved"], audit["output_tokens_saved"], audit["estimated_cost_avoided"],
        ] if audit else [None, None, None, None, None, None, None]
        ps = r.get("page_selection")
        row += [
            ps["pages_extracted"], ps["pages_skipped"],
            ps["input_tokens_saved"], ps["output_tokens_saved"], ps["estimated_cost_avoided"],
        ] if ps else [None, None, None, None, None]
        for svc in service_names:
            s = r["services"].get(svc)
            row += [s["cost"], s["input_tokens"], s["output_tokens"], s["model_id"]] if s else [None, None, None, None]
        data_rows.append(row)
    write_sheet(ws_docs, headers, data_rows)

    ws_model = wb.create_sheet("PorModelo")
    write_sheet(
        ws_model,
        ["ModelId", "Documentos", "Llamadas", "InputTokens", "OutputTokens", "CostoUsd", "CostoPorDocUsd", "TarifaAproximada"],
        [
            [m, len(a["docs"]), a["calls"], a["input_tokens"], a["output_tokens"], a["cost"],
             (a["cost"] / len(a["docs"]) if a["docs"] else 0), a["approximate"]]
            for m, a in sorted(by_model.items(), key=lambda kv: kv[1]["cost"], reverse=True)
        ],
    )

    ws_type = wb.create_sheet("PorTipo")
    write_sheet(
        ws_type,
        ["DocumentType", "DocumentTypeName", "Documentos", "Paginas", "InputTokens", "OutputTokens", "CostoUsd", "CostoPorDocUsd", "CostoPorPaginaUsd"],
        [
            [t, a["name"], a["n"], a["pages"], a["input_tokens"], a["output_tokens"], a["cost"],
             (a["cost"] / a["n"] if a["n"] else 0), (a["cost"] / a["pages"] if a["pages"] else None)]
            for t, a in sorted(by_type.items(), key=lambda kv: kv[1]["cost"], reverse=True)
        ],
    )

    ws_svc = wb.create_sheet("PorServicio")
    write_sheet(
        ws_svc,
        ["Servicio", "Documentos", "InputTokens", "OutputTokens", "CostoUsd", "PorcentajeDelTotal"],
        [
            [SERVICE_LABEL_ES.get(name, name), a["docs"], a["input_tokens"], a["output_tokens"], a["cost"],
             (a["cost"] / overall["total_cost"] if overall["total_cost"] else 0)]
            for name, a in sorted(by_service.items(), key=lambda kv: kv[1]["cost"], reverse=True)
        ],
    )

    ws_tot = wb.create_sheet("Totales")
    write_sheet(
        ws_tot,
        ["Metrica", "Valor"],
        [
            ["Documentos procesados", overall["n"]],
            ["Paginas procesadas", overall["total_pages"]],
            ["Costo total USD", overall["total_cost"]],
            ["Costo promedio por documento USD", overall["avg_cost_doc"]],
            ["Costo promedio por pagina USD", overall["avg_cost_page"]],
            ["Tokens de entrada totales", overall["total_input_tokens"]],
            ["Tokens de salida totales", overall["total_output_tokens"]],
            ["Primer documento procesado", overall["first_date"]],
            ["Ultimo documento procesado", overall["last_date"]],
            ["Version(es) de tarifa", ", ".join(overall["pricing_versions"])],
        ],
    )

    if page_sel:
        ws_sel = wb.create_sheet("SeleccionPaginas")
        write_sheet(
            ws_sel,
            ["Metrica", "Facturado", "EvitadoEstimado", "SinSeleccion"],
            [
                ["Tokens de entrada", page_sel["billed_input_tokens"], page_sel["input_tokens_saved"],
                 page_sel["billed_input_tokens"] + page_sel["input_tokens_saved"]],
                ["Tokens de salida", page_sel["billed_output_tokens"], page_sel["output_tokens_saved"],
                 page_sel["billed_output_tokens"] + page_sel["output_tokens_saved"]],
                ["Llamadas a Bedrock", page_sel["chunks_sent"],
                 page_sel["chunks_without_selection"] - page_sel["chunks_sent"],
                 page_sel["chunks_without_selection"]],
                ["Paginas", page_sel["pages_extracted"], page_sel["pages_skipped"],
                 page_sel["pages_considered"]],
                ["Costo USD", page_sel["billed_cost"], page_sel["cost_avoided"],
                 page_sel["cost_without_selection"]],
                [],
                ["Documentos con seleccion activa", page_sel["docs"], "de", page_sel["docs_total"]],
                ["NOTA", "EvitadoEstimado valora una llamada que nunca se hizo: NO sumarlo al costo real."],
            ],
        )

    wb.save(out_path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--env", choices=("dev", "prod"), default="dev", help="ambiente a consultar (default: dev)")
    parser.add_argument("--region", default="us-east-1")
    parser.add_argument("--desde", help="filtra ProcessedAt >= YYYY-MM-DD")
    parser.add_argument("--hasta", help="filtra ProcessedAt <= YYYY-MM-DD")
    parser.add_argument("--limit", type=int, default=20, help="documentos en 'más caros' (default 20)")
    parser.add_argument("--out", help="carpeta de salida (default: reportes-costos/<env>/<timestamp>/)")
    parser.add_argument("--open", action="store_true", help="abrir el HTML en el navegador al terminar")
    args = parser.parse_args([a for a in sys.argv[1:] if a != "--"])

    table = table_name(args.env)
    print(f"Leyendo {table} en {args.region}...")
    try:
        raw_items = scan_all(table, args.region)
    except Exception as e:
        print(f"{RED}No pude leer la tabla {table} en {args.region}: {e}{RESET}")
        return 1

    if not raw_items:
        print(f"{YELLOW}La tabla {table} está vacía — aún no se ha procesado ningún documento con este stack.{RESET}")
        return 0

    rows = [normalize_item(it) for it in raw_items]
    rows = filter_by_date(rows, args.desde, args.hasta)
    if not rows:
        print(f"{YELLOW}Ningún documento cae en el rango de fechas pedido.{RESET}")
        return 0

    overall = aggregate_overall(rows)
    by_type = aggregate_by_type(rows)
    by_model = aggregate_by_model(rows)
    by_service = aggregate_by_service(rows)
    page_sel = aggregate_page_selection(rows)
    top = top_expensive(rows, args.limit)

    out_dir = Path(args.out) if args.out else DEFAULT_OUT_ROOT / args.env / datetime.now().strftime("%Y%m%d-%H%M%S")
    out_dir.mkdir(parents=True, exist_ok=True)

    md_path = out_dir / "reporte.md"
    md_path.write_text(render_markdown(rows, overall, by_type, by_model, by_service, top, args.env, page_sel), encoding="utf-8")

    html_path = out_dir / "reporte.html"
    html_path.write_text(render_html(rows, overall, by_type, by_model, by_service, top, args.env, page_sel), encoding="utf-8")

    xlsx_path = out_dir / "costos_raw.xlsx"
    write_excel(xlsx_path, rows, overall, by_type, by_model, by_service, page_sel)

    print(f"\n{CYAN}{overall['n']} documento(s) · {len(by_model)} modelo(s) · {len(by_type)} tipo(s) · costo total {money(overall['total_cost'])}{RESET}")
    print(f"Reporte → {out_dir}/")
    print(f"  {md_path.name}  (resumen en texto)")
    print(f"  {html_path.name}  (dashboard — ábrelo en el navegador)")
    print(f"  {xlsx_path.name}  (datos crudos — Excel)")

    if args.open:
        webbrowser.open(html_path.resolve().as_uri())

    return 0


if __name__ == "__main__":
    sys.exit(main())
